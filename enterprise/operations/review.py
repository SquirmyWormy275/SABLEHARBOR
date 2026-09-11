"""Prepare and independently verify the controlled operating-review workbook."""

import json
import shutil
from collections import defaultdict
from decimal import Decimal as D
from pathlib import Path

from enterprise.business.model import money

from .exports import content_hash, file_hash

PUBLICATIONS = Path(__file__).with_name("publications")
BOOK = "operating-review-v1.0.0.xlsx"


def review_inputs(model, result):
    fields = [
        "unit",
        "scenario",
        "year",
        "revenue_usd",
        "expense_usd",
        "net_income_usd",
        "assets_usd",
        "liabilities_usd",
        "equity_usd",
        "ending_cash_usd",
    ]
    annual = [{k: r[k] for k in fields} for r in result["unit_annual_rows"]]
    groups = defaultdict(list)
    for r in model.tables["treasury_reconciliation"]:
        if int(r["year"]) >= 2027:
            groups[r["scenario"], int(r["year"]), r["cash_flow"]].append(r)
    cash = []
    for (scenario, year, flow), rows in sorted(groups.items()):
        rows.sort(key=lambda r: int(r["month"]))
        cash.append(
            {
                "scenario": scenario,
                "year": year,
                "cash_flow": flow,
                "opening_unpaid_usd": rows[0]["opening_unpaid_usd"],
                **{
                    target: money(sum(D(r[source]) for r in rows))
                    for target, source in (
                        ("requested_usd", "requested_usd"),
                        ("funded_usd", "current_funded_usd"),
                        ("arrears_settled_usd", "arrears_paid_usd"),
                    )
                },
                "closing_unpaid_usd": rows[-1]["closing_unpaid_usd"],
            }
        )
    groups = defaultdict(list)
    for r in model.tables["commercial_arr_bridge"]:
        groups[r["scenario"], r["unit"], int(r["period"][:4])].append(r)
    arr = []
    for (scenario, unit, year), rows in sorted(groups.items()):
        rows.sort(key=lambda r: r["period"])
        arr.append(
            {
                "scenario": scenario,
                "unit": unit,
                "year": year,
                "opening_arr_usd": rows[0]["opening_arr_usd"],
                **{
                    target: money(sum(D(r[source]) for r in rows))
                    for target, source in (
                        ("new_arr_usd", "new_arr_usd"),
                        ("expansion_arr_usd", "expansion_arr_usd"),
                        ("price_escalation_arr_usd", "price_arr_usd"),
                        ("contraction_arr_usd", "contraction_arr_usd"),
                        ("churn_arr_usd", "churn_arr_usd"),
                    )
                },
                "closing_arr_usd": rows[-1]["closing_arr_usd"],
            }
        )
    counts = defaultdict(lambda: defaultdict(int))
    for r in model.tables["control_results"]:
        counts[r["scenario"], r["unit"], r["local_control_id"]][r["result"]] += 1
    controls = [
        {
            "scenario": s,
            "unit": u,
            "control": c,
            "scheduled": sum(v.values()),
            "passed": v["PASS"],
            "failed": v["FAIL"],
            "not_run": v["NOT_RUN"],
        }
        for (s, u, c), v in sorted(counts.items())
    ]
    groups = defaultdict(lambda: defaultdict(D))
    for r in model.tables["forecast_variance_contributions"]:
        if r["metric"] not in {"revenue_usd", "net_income_usd"}:
            continue
        key = r["scenario"], r["unit"], int(r["period"][:4]), r["vintage_id"], r["metric"]
        for k in (
            "prior_usd",
            "revised_usd",
            "price_contribution_usd",
            "volume_contribution_usd",
            "timing_contribution_usd",
            "workforce_contribution_usd",
        ):
            groups[key][k] += D(r[k])
    forecasts = [
        {
            "scenario": s,
            "unit": u,
            "year": y,
            "revision": v,
            "metric": m,
            **{k: money(value) for k, value in vals.items()},
        }
        for (s, u, y, v, m), vals in sorted(groups.items())
    ]
    data = {
        "format_version": 1,
        "classification": "PUBLIC_SYNTHETIC_SCENARIOS_NOT_OBSERVED_ACTUALS",
        "business_input_sha256": model.input_hash,
        "tables": {
            "Financial source": annual,
            "Cash obligations": cash,
            "Customer ARR": arr,
            "Control reviews": controls,
            "Forecast review": forecasts,
        },
    }
    if any(not rows for rows in data["tables"].values()):
        raise ValueError("Required workbook source population is empty")
    data["columns"] = {name: list(rows[0]) for name, rows in data["tables"].items()}
    data["content_id"] = content_hash(data)
    return data


def verify_review(data, output):
    from openpyxl import load_workbook

    manifest = json.loads((PUBLICATIONS / "review_manifest.json").read_text())
    source = PUBLICATIONS / BOOK
    if data["content_id"] != manifest["content_id"]:
        raise ValueError("Reviewed workbook is stale; rebuild from current review_inputs.json")
    if file_hash(source) != manifest["workbook_sha256"]:
        raise ValueError("Workbook bytes differ from approved snapshot")
    # Materialize this small workbook: optional XML dimension hints may be omitted.
    # This verification reads only; it never modifies or saves a workbook.
    values = load_workbook(source, data_only=True)
    formulas = load_workbook(source, data_only=False)
    if set(values.sheetnames) != {"Results", *data["tables"]}:
        raise ValueError("Workbook worksheet population mismatch")
    if values._external_links or formulas.vba_archive:
        raise ValueError("External links/macros prohibited")
    for name, rows in data["tables"].items():
        fields = data["columns"][name]
        sheet = values[name]
        if sheet.max_row != len(rows) + 5:
            raise ValueError(f"Workbook source population mismatch: {name}")
        for n, row in enumerate(rows, 6):
            for c, key in enumerate(fields, 1):
                actual, expected = sheet.cell(n, c).value, row[key]
                if isinstance(expected, int):
                    if actual is None or D(str(actual)) != D(expected):
                        raise ValueError(f"Workbook integer mismatch: {name}/{n}/{key}")
                elif key.endswith("_usd"):
                    if actual is None or D(str(actual)).quantize(D(".0001")) != D(
                        str(expected)
                    ).quantize(D(".0001")):
                        raise ValueError(f"Workbook numeric mismatch: {name}/{n}/{key}")
                elif str(actual) != str(expected):
                    raise ValueError(f"Workbook identity mismatch: {name}/{n}/{key}")
            if name != "Financial source":
                check = sheet.cell(n, len(fields) + 1).value
                if check is None or D(str(check)).quantize(D(".0001")) != 0:
                    raise ValueError(f"Workbook reconciliation fails: {name}/{n}")
                expected_formula = {
                    "Cash obligations": f"=D{n}+E{n}-F{n}-G{n}-H{n}",
                    "Customer ARR": f"=D{n}+E{n}+F{n}+G{n}-H{n}-I{n}-J{n}",
                    "Control reviews": f"=D{n}-SUM(E{n}:G{n})",
                    "Forecast review": f"=G{n}-F{n}-SUM(H{n}:K{n})",
                }[name]
                if formulas[name].cell(n, len(fields) + 1).value != expected_formula:
                    raise ValueError(f"Workbook reconciliation formula mismatch: {name}/{n}")
    year = D(str(values["Results"]["D3"].value))
    if year not in {D(y) for y in range(2026, 2032)}:
        raise ValueError("Workbook report year outside reviewed population")
    rows = sorted(
        [r for r in data["tables"]["Financial source"] if int(r["year"]) == year],
        key=lambda r: (r["unit"], r["scenario"]),
    )
    end = len(data["tables"]["Financial source"]) + 5
    for n, row in enumerate(rows, 6):
        if (
            values["Results"].cell(n, 3).value != row["unit"]
            or values["Results"].cell(n, 4).value != row["scenario"]
        ):
            raise ValueError("Workbook summary identity mismatch")
        for c, key, col in zip(
            range(5, 9),
            ("revenue_usd", "expense_usd", "net_income_usd", "ending_cash_usd"),
            ("D", "E", "F", "J"),
        ):
            actual = values["Results"].cell(n, c).value
            if actual is None or D(str(actual)).quantize(D(".0001")) != D(row[key]).quantize(
                D(".0001")
            ):
                raise ValueError("Workbook summary differs from regenerated enterprise")
            formula = (
                f"=SUMIFS('Financial source'!${col}$6:${col}${end},"
                f"'Financial source'!$A$6:$A${end},$C{n},"
                f"'Financial source'!$B$6:$B${end},$D{n},"
                f"'Financial source'!$C$6:$C${end},$D$3)"
            )
            if formulas["Results"].cell(n, c).value != formula:
                raise ValueError("Workbook summary formula/range mismatch")
    for sheet in values:
        for row in sheet:
            if any(c.data_type == "e" for c in row):
                raise ValueError("Workbook formula error")
    values.close()
    formulas.close()
    shutil.copyfile(source, Path(output) / BOOK)
    shutil.copyfile(PUBLICATIONS / "review_manifest.json", Path(output) / "review_manifest.json")
    return {"status": "PASS", "content_id": data["content_id"]}
