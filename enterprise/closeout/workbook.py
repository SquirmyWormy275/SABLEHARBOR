"""Current financial reading workbook using the existing runtime table publication style."""

import csv
import hashlib
import json
from datetime import datetime
from decimal import Decimal

import xlsxwriter
from openpyxl import load_workbook

SHEETS = [
    (
        "Legal balances",
        "enterprise/enterprise_annual_statements.csv",
        [
            "scenario",
            "entity",
            "year",
            "assets_usd",
            "liabilities_usd",
            "equity_usd",
            "net_income_usd",
            "ending_cash_usd",
        ],
    ),
    (
        "Cash flows",
        "enterprise/enterprise_annual_statements.csv",
        [
            "scenario",
            "entity",
            "year",
            "opening_cash_usd",
            "operating_cash_flow_usd",
            "investing_cash_flow_usd",
            "financing_cash_flow_usd",
            "ending_cash_usd",
        ],
    ),
    (
        "Sovereignty",
        "sovereignty.csv",
        [
            "scenario",
            "year",
            "member_cash_usd",
            "sustaining_investment_paid_usd",
            "growth_net_investment_paid_usd",
            "growth_internal_current_year_usd",
            "growth_external_current_year_usd",
            "growth_cash_origin_unattributed_usd",
        ],
    ),
    (
        "Funding constraints",
        "sovereignty.csv",
        [
            "scenario",
            "year",
            "ending_cash_usd",
            "outstanding_due_or_unresolved_settlement_usd",
            "requirement_excess_over_consolidated_cash_usd",
            "peak_core_unpaid_within_year_usd",
            "planned_industrial_project_balance_usd",
            "deferred_industrial_capex_usd",
        ],
    ),
    (
        "Parent provision",
        "parent_tax_provision.csv",
        [
            "scenario",
            "entity",
            "year",
            "book_pretax_usd",
            "current_expense_usd",
            "net_deferred_expense_usd",
            "net_deferred_position_usd",
            "closing_federal_nol_usd",
        ],
    ),
    (
        "Tax jurisdictions",
        "parent_tax_provision.csv",
        [
            "scenario",
            "year",
            "federal_current_usd",
            "california_current_usd",
            "california_unallocated_reserve_usd",
            "interest_carryforward_usd",
            "gross_dta_usd",
            "valuation_allowance_usd",
        ],
    ),
    (
        "Asset tax basis",
        "parent_tax_assets.csv",
        [
            "scenario",
            "asset_id",
            "year",
            "cost",
            "federal_depreciation_usd",
            "california_depreciation_usd",
            "federal_closing_basis_usd",
            "california_closing_basis_usd",
        ],
    ),
]


def build(out):
    path = out / "company-financial-reading-v1.xlsx"
    sources, tables = {}, {}
    for name, source, fields in SHEETS:
        data = out / source
        rows = list(csv.DictReader(data.open(newline="")))
        if not rows or not set(fields) <= rows[0].keys():
            raise ValueError("Workbook source schema/population changed")
        sources[source] = hashlib.sha256(data.read_bytes()).hexdigest()
        tables[name] = (fields, rows)
    book = xlsxwriter.Workbook(path, {"strings_to_formulas": False, "strings_to_urls": False})
    book.set_properties(
        {
            "title": "Sable Harbor company financial reading",
            "author": "Sable Harbor synthetic company archive",
            "created": datetime(2026, 9, 15),
        }
    )
    title = book.add_format(
        {"font_name": "Arial", "font_size": 18, "bold": True, "font_color": "#173247"}
    )
    header = book.add_format(
        {
            "font_name": "Arial",
            "font_size": 9,
            "bold": True,
            "bg_color": "#173247",
            "font_color": "white",
            "text_wrap": True,
        }
    )
    text = book.add_format({"font_name": "Arial", "font_size": 9, "text_wrap": True})
    money = book.add_format(
        {"font_name": "Arial", "font_size": 9, "num_format": "#,##0.0000;[Red](#,##0.0000)"}
    )
    note = book.add_format(
        {"font_name": "Arial", "font_size": 9, "font_color": "#555555", "text_wrap": True}
    )
    start = book.add_worksheet("Read first")
    start.set_column("A:H", 16)
    start.merge_range("A1:H1", "Sable Harbor — company financial reading", title)
    start.merge_range(
        "A3:H5",
        "SYNTHETIC COMPANY EVIDENCE. Current composed source, including tax, goodwill and "
        "funding constraints. Source period/scenario roles control. This workbook is a "
        "selected reading view; full CSV/SQLite populations and workpapers remain in the edition.",
        note,
    )
    start.merge_range(
        "A7:H9",
        "Read CONTRACT.json, MANIFEST.json and the exception register first. Separate legal "
        "entities, consolidated totals, calibration and conditional forecasts. Do not add "
        "these representations together. No bank confirmation or audit opinion is asserted.",
        note,
    )
    start.merge_range(
        "A11:H13",
        "This successor reuses the existing runtime workbook typography, colors, tables and "
        "print conventions. All prior approved workbooks retain their original bytes. "
        "Four-decimal source values are retained; each sheet identifies its exact source.",
        note,
    )
    for index, (name, source, _) in enumerate(SHEETS, 15):
        start.write_url(index, 0, f"internal:'{name}'!A1", text, name)
        start.merge_range(index, 2, index, 7, source, note)
    start.set_landscape()
    start.set_paper(9)
    start.fit_to_pages(1, 1)
    start.print_area(0, 0, 23, 7)
    checked = 0
    for name, source, fields in SHEETS:
        _, rows = tables[name]
        sheet = book.add_worksheet(name)
        sheet.set_column(0, len(fields) - 1, 18)
        sheet.merge_range(0, 0, 0, len(fields) - 1, name, title)
        sheet.merge_range(
            1, 0, 2, len(fields) - 1, "Source: " + source + " | SHA-256 " + sources[source], note
        )
        sheet.set_row(4, 42)
        for col, field in enumerate(fields):
            sheet.write_string(4, col, field.replace("_usd", " USD").replace("_", " "), header)
        for index, row in enumerate(rows, 5):
            sheet.set_row(index, 27)
            for col, field in enumerate(fields):
                value = row[field]
                if field.endswith("_usd") or field == "cost":
                    number = Decimal(value)
                    if not number.is_finite():
                        raise ValueError("Nonfinite workbook amount")
                    sheet.write_number(index, col, float(number), money)
                else:
                    sheet.write_string(index, col, value, text)
                checked += 1
        sheet.autofilter(4, 0, len(rows) + 4, len(fields) - 1)
        sheet.freeze_panes(5, 3)
        sheet.repeat_rows(0, 4)
        sheet.set_landscape()
        sheet.set_paper(9)
        sheet.fit_to_pages(1, 0)
        sheet.print_area(0, 0, len(rows) + 4, len(fields) - 1)
    book.close()
    # Independent library readback, against original CSV cells, not another writer output.
    readback = load_workbook(path, read_only=True, data_only=False)
    for name, _, fields in SHEETS:
        _, rows = tables[name]
        actual = readback[name].iter_rows(min_row=6, values_only=True)
        for row in rows:
            values = next(actual)
            for field, value in zip(fields, values, strict=True):
                if field.endswith("_usd") or field == "cost":
                    if abs(Decimal(str(value)) - Decimal(row[field])) > Decimal(".0001"):
                        raise ValueError("Workbook monetary readback differs")
                elif value != row[field]:
                    raise ValueError("Workbook source identity differs")
        if next(actual, None) is not None:
            raise ValueError("Extra workbook row")
    readback.close()
    receipt = {
        "workbook": path.name,
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "source_hashes": sources,
        "sheets": {n: len(rows) for n, (_, rows) in tables.items()},
        "independent_cell_readback_count": checked,
        "result": "PASS",
        "scope": "Selected reading; complete source CSV/SQLite populations remain in the edition",
    }
    (out / "workbook-receipt.json").write_text(json.dumps(receipt, indent=2) + "\n")
    return receipt
