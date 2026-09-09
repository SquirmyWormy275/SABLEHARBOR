"""Independent evidence reconciliation and export verification."""

from __future__ import annotations

import json
import sqlite3
from collections import defaultdict
from decimal import Decimal as D
from pathlib import Path

from industrial.planning.integrity import journals
from .exports import file_hash, encoded


def total(rows, keys, value="signed_usd"):
    answer = defaultdict(D)
    for row in rows:
        answer[tuple(str(row[k]) for k in keys)] += D(str(row[value]))
    return {k: v for k, v in answer.items() if v}


def business(model):
    rows = model.tables["journal"]
    result = journals(rows)
    events = model.tables["events"]
    ids = [r["event_id"] for r in events]
    if len(ids) != len(set(ids)):
        raise ValueError("Duplicate business event identity")
    if not {r["source_id"] for r in rows} <= set(ids):
        raise ValueError("Financial entry lacks an originating business event")
    assignments = defaultdict(D)
    for row in model.tables["workforce_assignments"]:
        assignments[row["scenario"], row["period"], row["person_id"]] += D(row["assignment_fte"])
        if row["home_group"] in {"j2", "atlas-meridian"} and row["unit"] == "advisory":
            raise ValueError("J2/Atlas is being used as an Advisory staffing reserve")
    if any(v != 1 for v in assignments.values()):
        raise ValueError("An employed position is counted more or less than once")
    for r in model.tables["capacity"]:
        if D(r["used_hours"]) < 0 or D(r["used_hours"]) > D(r["capacity_hours"]):
            raise ValueError("Delivered effort exceeds staffed capacity")
    for r in events:
        if r["kind"] == "RECOVERY_RUN" and r["bypass"] and D(r["recovered_kg"]):
            raise ValueError("Hard bypass still produces recovered material")
        if r["kind"] == "RECOVERY_RUN" and D(r["recovered_kg"]) > D(r["contained_kg"]):
            raise ValueError("Recovered mass exceeds contained material")
    accounts = defaultdict(lambda: defaultdict(D))
    by_period = defaultdict(list)
    for row in rows:
        by_period[row["scenario"], row["period"]].append(row)
    roll = {
        (r["scenario"], r["period"], r["unit"]): r for r in model.tables["subledger_rollforward"]
    }
    mappings = {
        "BIZ_AR": ("gross_ar_usd", 1),
        "BIZ_AP": ("vendor_ap_usd", -1),
        "BIZ_DEFERRED": ("deferred_revenue_usd", -1),
        "BIZ_INVENTORY": ("inventory_usd", 1),
        "BIZ_ALLOWANCE": ("allowance_usd", -1),
        "BIZ_HOST_AP": ("host_payable_usd", -1),
        "BIZ_PPE": ("gross_ppe_usd", 1),
        "BIZ_ACCUM": ("accumulated_depreciation_usd", -1),
    }
    for scenario, period in sorted(by_period):
        for row in by_period[scenario, period]:
            accounts[scenario, row["unit"]][row["account"]] += D(row["signed_usd"])
        for unit in model.payroll:
            r = roll[scenario, period, unit]
            for account, (field, sign) in mappings.items():
                if accounts[scenario, unit][account] * sign != D(r[field]):
                    raise ValueError(
                        f"Subledger rollforward does not independently tie: {unit}/{field}"
                    )
    return {
        "business_journals": result["journal_count"],
        "source_events": len(events),
        "occupied_person_months": len(assignments),
        "subledger_period_units": len(roll),
    }


def enterprise(result, baseline, model):
    journals(result["journal_rows"])
    keys = ["scenario", "entity", "year", "month", "account", "segment"]
    old2026 = [r for r in baseline["journal_rows"] if int(r["year"]) == 2026]
    new2026 = [r for r in result["journal_rows"] if int(r["year"]) == 2026]
    if total(old2026, keys) != total(new2026, keys):
        raise ValueError("Preserved 2026 enterprise reconstruction changed")
    for row in result["journal_rows"]:
        if int(row["year"]) >= 2027 and row["account"] in {"CORE_REVENUE", "CORE_OPEX"}:
            raise ValueError("Old projected Core operating envelope survived replacement")
    imported = [r for r in result["journal_rows"] if r["source_type"] == "BUSINESS_DRIVEN_FORECAST"]
    if total(imported, ["scenario", "year", "month", "account", "segment", "source_id"]) != total(
        model.tables["journal"], ["scenario", "year", "month", "account", "segment", "source_id"]
    ):
        raise ValueError("Business postings were omitted, duplicated or changed on import")
    groupkeys = ["scenario", "year", "month", "account"]
    if total(result["unit_rows"], groupkeys) != total(
        result["legal_trial_balance_rows"], groupkeys
    ):
        raise ValueError("Seven units plus corporate/eliminations do not bridge to legal books")
    if total(result["legal_trial_balance_rows"], ["scenario", "entity", "year", "month"]):
        raise ValueError("A legal monthly trial balance is unbalanced")
    for row in result["monthly_rows"]:
        if D(row["assets_usd"]) != D(row["liabilities_usd"]) + D(row["equity_usd"]):
            raise ValueError("Financial position does not balance")
        cash = sum(
            D(row[k])
            for k in (
                "opening_cash_usd",
                "operating_cash_flow_usd",
                "investing_cash_flow_usd",
                "financing_cash_flow_usd",
                "opening_or_noncash_cash_bridge_usd",
            )
        )
        if cash != D(row["ending_cash_usd"]):
            raise ValueError("Classified cash-flow statement does not roll")
    return {
        "enterprise_journal_lines": len(result["journal_rows"]),
        "monthly_statements": len(result["monthly_rows"]),
        "preserved_2026": "EXACT_BY_SOURCE_ACCOUNT_SEGMENT",
        "core_replacement": "NO_OLD_OPERATING_ENVELOPE",
        "unit_to_legal_reconciliation": "EXACT_BY_SCENARIO_PERIOD_ACCOUNT",
    }


def unit_exports(output, packages):
    from openpyxl import load_workbook

    counts = {}
    for unit, tables in packages.items():
        path = Path(output) / unit
        with sqlite3.connect(path / "evidence.sqlite3") as db:
            names = {r[0] for r in db.execute("SELECT name FROM sqlite_master WHERE type='table'")}
            if names != set(tables):
                raise ValueError("Unexpected or omitted SQLite table")
            for name, rows in tables.items():
                columns = [r[1] for r in db.execute(f'PRAGMA table_info("{name}")')]
                actual = db.execute(f'SELECT * FROM "{name}"').fetchall()
                expected = [tuple(encoded(row.get(c)) for c in columns) for row in rows]
                if actual != expected:
                    raise ValueError("Unit export population differs from allowed source rows")

        values = load_workbook(path / "audit.xlsx", data_only=True, read_only=True)
        formulas = load_workbook(path / "audit.xlsx", data_only=False, read_only=True)
        if values._external_links or formulas.vba_archive:
            raise ValueError("Workbook external link or macro is forbidden")
        for table, sheet_name in [
            ("annual_statements", "Annual"),
            ("monthly_statements", "Monthly"),
            ("trial_balance", "Trial balance"),
        ]:
            sheet = values[sheet_name]
            headers = [c.value for c in next(sheet.rows)]
            actual = list(sheet.iter_rows(min_row=2, values_only=True))
            if len(actual) != len(tables[table]):
                raise ValueError("Workbook financial population count mismatch")
            for cells, expected in zip(actual, tables[table]):
                row = dict(zip(headers, cells))
                for key, value in expected.items():
                    if key.endswith("_usd"):
                        if D(str(row[key])).quantize(D("0.0001")) != D(str(value)):
                            raise ValueError(
                                f"Workbook amount differs from verified SQLite: {unit}/{table}/{key}"
                            )
                    elif key in {"scenario", "year", "month", "unit", "account"} and str(
                        row[key]
                    ) != str(value):
                        raise ValueError("Workbook financial row identity differs from SQLite")
        for row in values["Checks"].iter_rows(min_row=2, values_only=True):
            if row[3] != 0:
                raise ValueError("Workbook cached reconciliation is not zero")
        values.close()
        formulas.close()
        verify_files(path)
        counts[unit] = {name: len(rows) for name, rows in tables.items()}
    return counts


def verify_files(root):
    root = Path(root)
    manifest = json.loads((root / "manifest.json").read_text())
    actual = {str(p.relative_to(root)) for p in root.rglob("*") if p.is_file()}
    if actual != set(manifest["artifacts"]) | {"manifest.json", "SHA256SUMS.txt"}:
        raise ValueError("Missing, stale or unmanifested distribution artifact")
    for relative, expected in manifest["artifacts"].items():
        if file_hash(root / relative) != expected:
            raise ValueError(f"Artifact hash mismatch: {relative}")
    lines = (root / "SHA256SUMS.txt").read_text().splitlines()
    checks = {line.split("  ", 1)[1]: line.split("  ", 1)[0] for line in lines}
    if set(checks) != set(manifest["artifacts"]) | {"manifest.json"}:
        raise ValueError("Checksum inventory mismatch")
    if any(file_hash(root / p) != h for p, h in checks.items()):
        raise ValueError("Distribution checksum mismatch")
    return manifest
