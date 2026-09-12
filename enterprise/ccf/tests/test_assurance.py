"""Adversarial assessment cases and complete workbench re-performance."""

import copy
import hashlib
import json
from datetime import date

import openpyxl
import pytest

from enterprise.ccf.assurance.__main__ import build, customer_statements, files, read_json, verify
from enterprise.ccf.assurance.catalog import starter
from enterprise.ccf.assurance.engine import data, plan
from enterprise.ccf.assurance.examples import (
    bind_fixture_reviews,
    blank_assessment,
    example,
    review,
)
from enterprise.ccf.assurance.models import Assessment, Catalog, Source
from enterprise.ccf.assurance.reporting import csv_report, explorer
from enterprise.ccf.registry import compile_registry, digest


@pytest.fixture(scope="module")
def native():
    return compile_registry()


@pytest.fixture
def fixture(native):
    c, a = example(native)
    return data(c), data(a)


def run(c, a, native):
    catalog = Catalog.model_validate(bind_fixture_reviews(data(Catalog.model_validate(c))))
    a = copy.deepcopy(a)
    a["catalog_digest"] = digest(data(catalog))
    assessment = Assessment.model_validate(bind_fixture_reviews(data(Assessment.model_validate(a))))
    return plan(catalog, assessment, native)


def first(report):
    return next(r for r in report["rows"] if r["requirement_id"] == "SOC2:EXAMPLE-ACCESS")


def test_complete_example_and_separate_reuse(fixture, native):
    c, a = fixture
    r = run(c, a, native)
    assert r["summary"] == {"GAP": 6, "SUPPORTED": 2}
    assert len({eid for t in a["tests"][:2] for eid in t["evidence_ids"]}) == 1
    assert a["tests"][0]["requirement_id"] != a["tests"][1]["requirement_id"]
    extension = next(x for x in r["rows"] if x["requirement_id"] == "ISO27001:EXAMPLE-ACCESS")
    assert extension["candidate_reuse_from_baseline"] == ["SH-IAM-004"]
    assert extension["attributes"][0]["disposition"] == "EVIDENCE_GAP"
    assert first(r)["attributes"][0]["disposition"] == "DEMONSTRATED_SYNTHETIC"
    assert set(x["attributes"][0]["disposition"] for x in r["rows"]) >= {
        "NEW_CONTROL_OR_MAPPING_NEEDED",
        "IMPLEMENT_CONTROL",
        "REVIEW_OR_ENHANCE_CONTROL",
        "ASSESSMENT_GAP",
    }


@pytest.mark.parametrize(
    "mutation,reason",
    [
        (lambda c, a: a["tests"][0].update(result="FAIL", findings=["Failure"]), "TEST_FAIL"),
        (
            lambda c, a: a["tests"][0].update(result="NOT_RUN", findings=["Missing population"]),
            "TEST_NOT_RUN",
        ),
        (lambda c, a: a["tests"][0].update(review=None), "TEST_REVIEW_PENDING"),
        (lambda c, a: a["tests"][0].update(mode="DESIGN"), "TEST_MODE_MISMATCH"),
        (lambda c, a: a["tests"][0].update(period_start="2026-09-03"), "TEST_PERIOD_GAP"),
        (
            lambda c, a: a["tests"][0].update(expected_population_digest="0" * 64),
            "EVIDENCE_POPULATION_MISMATCH",
        ),
        (
            lambda c, a: a["evidence"][0].update(boundary_id="advisory"),
            "EVIDENCE_BOUNDARY_MISMATCH",
        ),
        (lambda c, a: a["evidence"][0].update(period_start="2026-09-03"), "EVIDENCE_PERIOD_GAP"),
        (
            lambda c, a: a["evidence"][0].update(origin="OPERATING_RECORDS"),
            "EVIDENCE_ORIGIN_MISMATCH",
        ),
        (lambda c, a: a["scope"].update(known_on="2026-10-02"), "EVIDENCE_EXPIRED"),
        (lambda c, a: a["implementations"][0].update(state="PROPOSED"), "IMPLEMENTATION_GAP"),
        (
            lambda c, a: a["implementations"][0].update(effective_from="2026-09-03"),
            "IMPLEMENTATION_GAP",
        ),
        (lambda c, a: a["implementations"][0].update(review=None), "IMPLEMENTATION_GAP"),
        (lambda c, a: c["mappings"][0].update(review=None), "MAPPING_REVIEW_PENDING"),
        (
            lambda c, a: c["frameworks"][0].update(inventory_complete=False),
            "FRAMEWORK_INVENTORY_INCOMPLETE",
        ),
        (lambda c, a: c["sources"][0].update(access="METADATA_ONLY"), "SOURCE_ACCESS_PENDING"),
        (
            lambda c, a: c["sources"][0].update(effective_to="2026-09-09"),
            "SOURCE_EDITION_PERIOD_UNRESOLVED",
        ),
        (lambda c, a: c["sources"][0].update(review=None), "SOURCE_REVIEW_PENDING"),
        (lambda c, a: c["requirements"][0].update(review=None), "REQUIREMENT_REVIEW_PENDING"),
        (lambda c, a: a["scope"].update(review=None), "SCOPE_REVIEW_PENDING"),
        (
            lambda c, a: a["applicability"][0].update(disposition="UNRESOLVED"),
            "APPLICABILITY_UNRESOLVED",
        ),
        (
            lambda c, a: a["applicability"][0].update(disposition="EXCLUDED", review=None),
            "APPLICABILITY_UNRESOLVED",
        ),
    ],
)
def test_invalid_support_stays_visible(fixture, native, mutation, reason):
    c, a = fixture
    mutation(c, a)
    row = first(run(c, a, native))
    assert row["status"] != "SUPPORTED"
    assert reason in row["blockers"] + row["attributes"][0]["issues"]


@pytest.mark.parametrize(
    "mutation",
    [
        lambda c, a: c["requirements"].pop(),
        lambda c, a: c["mappings"].append(copy.deepcopy(c["mappings"][0])),
        lambda c, a: c["mappings"][0].update(control_id="SH-IAM-999"),
        lambda c, a: c["mappings"][0].update(control_digest="0" * 64),
        lambda c, a: c["mappings"][0].update(attribute_ids=["UNKNOWN"]),
        lambda c, a: a["scope"].update(boundaries=["unknown"]),
        lambda c, a: a["scope"]["baseline"][0].update(categories=["UNKNOWN"]),
        lambda c, a: a["scope"]["baseline"][0].update(categories=[]),
        lambda c, a: a["tests"][0].update(implementation_version="2"),
        lambda c, a: a["tests"][0].update(evidence_ids=["missing"]),
        lambda c, a: a["tests"][0].update(evidence_ids=[]),
        lambda c, a: a["tests"][0].update(attribute_ids=["unknown"]),
        lambda c, a: a["tests"][0].update(period_end="2026-08-01"),
        lambda c, a: a["tests"][0]["review"].update(reviewer="SYNTHETIC-PREPARER"),
        lambda c, a: a["tests"][0]["review"].update(reviewed_on="2026-09-09"),
        lambda c, a: a["evidence"][0]["payload"].update(forged=True),
        lambda c, a: a["implementations"][0].update(owner_role_id="ROLE-UNKNOWN"),
        lambda c, a: a["implementations"][0].update(effective_to="2025-01-01"),
        lambda c, a: c["frameworks"][0].update(inventory_complete=True, inventory_review=None),
        lambda c, a: c["frameworks"][0].update(inventory_complete="true"),
        lambda c, a: a.update(unknown_claim="PASS"),
    ],
)
def test_structural_or_identity_corruption_rejected(fixture, native, mutation):
    c, a = fixture
    mutation(c, a)
    with pytest.raises(ValueError):
        run(c, a, native)


def test_assessment_cannot_silently_follow_catalog_change(fixture, native):
    c, a = fixture
    c["version"] = "2"
    with pytest.raises(ValueError, match="different catalog"):
        plan(Catalog.model_validate(c), Assessment.model_validate(a), native)
    c["native_snapshot_id"] = "0" * 64
    with pytest.raises(ValueError, match="native snapshot"):
        run(c, a, native)


def test_requirement_needs_all_attributes(fixture, native):
    c, a = fixture
    attribute = copy.deepcopy(c["requirements"][0]["attributes"][0])
    attribute["id"] = "A2"
    c["requirements"][0]["attributes"].append(attribute)
    row = first(run(c, a, native))
    assert row["status"] == "GAP"
    assert row["attributes"][0]["disposition"] == "DEMONSTRATED_SYNTHETIC"
    assert row["attributes"][1]["disposition"] == "NEW_CONTROL_OR_MAPPING_NEEDED"


def test_failed_test_cannot_be_erased_by_new_pass(fixture, native):
    c, a = fixture
    failed = copy.deepcopy(a["tests"][0])
    failed.update(id="UNRESOLVED-FAILURE", result="FAIL", findings=["Unremediated failure"])
    a["tests"].append(failed)
    assert first(run(c, a, native))["status"] == "GAP"


def test_two_boundaries_do_not_inherit_evidence(fixture, native):
    c, a = fixture
    a["scope"]["boundaries"].append("advisory")
    r = run(c, a, native)
    assert len(r["rows"]) == 16
    assert all(x["status"] == "UNRESOLVED" for x in r["rows"] if x["boundary_id"] == "advisory")


def test_reviewed_exclusion_is_not_coverage(fixture, native):
    c, a = fixture
    a["applicability"][0].update(disposition="EXCLUDED")
    row = first(run(c, a, native))
    assert row["status"] == "EXCLUDED"
    assert row["applicability_rationale"]
    assert row["attributes"][0]["disposition"] != "DEMONSTRATED_SYNTHETIC"


def test_starter_never_claims_a_complete_framework(native):
    c = starter(native)
    a = blank_assessment(c)
    r = plan(c, a, native)
    assert len(c.requirements) == 105
    assert len(c.frameworks) == 5
    assert r["summary"] == {"UNRESOLVED": 77}
    assert {x.category for x in c.requirements if x.framework_id == "HIPAA"} == {
        "GENERAL",
        "SECURITY",
        "PRIVACY",
        "BREACH",
    }
    assert all("FRAMEWORK_INVENTORY_INCOMPLETE" in x["blockers"] for x in r["rows"])
    a.scope.baseline[0].categories.extend(
        ["AVAILABILITY", "CONFIDENTIALITY", "PRIVACY", "PROCESSING_INTEGRITY"]
    )
    assert len(plan(c, a, native)["rows"]) == 105


def test_source_requires_hash_and_valid_dates():
    with pytest.raises(ValueError):
        Source(
            id="S",
            publisher="P",
            edition="V",
            url="https://example.invalid",
            retrieved_on=date(2026, 9, 11),
            effective_from=date(2026, 9, 1),
            access="AVAILABLE",
            rights_note="Test",
        )


def test_customer_export_requires_operating_review_and_does_not_leak(fixture, native, tmp_path):
    c, a = fixture
    with pytest.raises(ValueError, match="operating scope"):
        customer_statements(
            Catalog.model_validate(c), Assessment.model_validate(a), native, tmp_path
        )
    a["scope"]["origin"] = "OPERATING_RECORDS"
    for source in c["sources"]:
        source["access"] = "AVAILABLE"
        content = source["id"].encode()
        source["content_sha256"] = hashlib.sha256(content).hexdigest()
        (tmp_path / source["content_sha256"]).write_bytes(content)
    for evidence in a["evidence"]:
        evidence["origin"] = "OPERATING_RECORDS"
        evidence["classification"] = "RESTRICTED"
        evidence["payload"]["private"] = "DO-NOT-DISCLOSE"
        evidence["sha256"] = digest(evidence["payload"])
    a["catalog_digest"] = digest(c)
    a["disclosures"] = [
        dict(
            audience="Example customer",
            statement="Approved scoped statement",
            requirement_ids=["SOC2:EXAMPLE-ACCESS"],
            review=review(),
        )
    ]
    c = data(Catalog.model_validate(bind_fixture_reviews(data(Catalog.model_validate(c)))))
    a["catalog_digest"] = digest(c)
    a = data(Assessment.model_validate(bind_fixture_reviews(data(Assessment.model_validate(a)))))
    output = customer_statements(
        Catalog.model_validate(c), Assessment.model_validate(a), native, tmp_path
    )
    assert "DO-NOT-DISCLOSE" not in json.dumps(output)
    assert "findings" not in json.dumps(output)
    assert output["statements"] == [
        {"audience": "Example customer", "statement": "Approved scoped statement"}
    ]
    a["disclosures"][0]["requirement_ids"] = ["ISO27001:EXAMPLE-ACCESS"]
    a = data(Assessment.model_validate(bind_fixture_reviews(data(Assessment.model_validate(a)))))
    with pytest.raises(ValueError, match="unsupported"):
        customer_statements(
            Catalog.model_validate(c), Assessment.model_validate(a), native, tmp_path
        )


def test_html_and_csv_escape_untrusted_text(fixture, native):
    c, a = fixture
    c["requirements"][0]["attributes"][0]["objective"] = '<script>alert("x")</script>'
    r = run(c, a, native)
    assert '<script>alert("x")</script>' not in explorer(r)
    assert "&lt;script&gt;" in explorer(r)
    c["requirements"][0]["attributes"][0]["objective"] = "=HYPERLINK(1)"
    assert "'=HYPERLINK(1)" in csv_report(run(c, a, native))


def test_complete_packages_reproduce_and_rehashed_false_report_fails(tmp_path, native):
    c, a = example(native)
    one, two = tmp_path / "one", tmp_path / "two"
    build(one, c, a, native)
    build(two, c, a, native)
    assert all(p.read_bytes() == (two / p.name).read_bytes() for p in one.iterdir())
    assert verify(one)["verified"]
    assert one.stat().st_mode & 0o777 == 0o700
    with pytest.raises(ValueError, match="new directory"):
        build(one, c, a, native)
    book = openpyxl.load_workbook(one / "workbench.xlsx", read_only=True)
    assert book.sheetnames == [
        "Read me",
        "Delta",
        "Native controls",
        "Sources",
        "Mappings",
        "Evidence requests",
        "Evidence index",
        "Tests",
    ]
    assert book["Native controls"].max_row == 167
    assert all(cell.data_type != "f" for sheet in book for row in sheet for cell in row)
    book.close()
    report = read_json(one / "delta.json")
    report["summary"] = {"SUPPORTED": 8}
    (one / "delta.json").write_text(json.dumps(report))
    manifest = read_json(one / "MANIFEST.json")
    manifest["files"] = files(one)
    (one / "MANIFEST.json").write_text(json.dumps(manifest))
    with pytest.raises(ValueError, match="re-performance"):
        verify(one)


@pytest.mark.parametrize("text", ['{"a":1,"a":2}', '{"a":NaN}', '{"a":Infinity}'])
def test_ambiguous_json_rejected(tmp_path, text):
    path = tmp_path / "invalid.json"
    path.write_text(text)
    with pytest.raises(ValueError):
        read_json(path)


def test_original_source_required_and_checked(fixture, native, tmp_path):
    from enterprise.ccf.assurance.__main__ import verify_sources

    c, a = fixture
    c["sources"][0]["access"] = "AVAILABLE"
    catalog = Catalog.model_validate(c)
    with pytest.raises(ValueError, match="original document"):
        verify_sources(catalog)
    with pytest.raises(ValueError, match="Missing original"):
        verify_sources(catalog, tmp_path)
    (tmp_path / c["sources"][0]["content_sha256"]).write_bytes(b"wrong document")
    with pytest.raises(ValueError, match="hash mismatch"):
        verify_sources(catalog, tmp_path)


@pytest.mark.parametrize(
    "mutate",
    [
        lambda c, a: a["scope"].update(period_start="2026-09-03"),
        lambda c, a: a["tests"][0].update(result="FAIL", findings=["Changed"]),
        lambda c, a: c["mappings"][0].update(rationale="Changed mapping"),
        lambda c, a: c["sources"][0].update(edition="Changed edition"),
    ],
)
def test_record_changes_invalidate_existing_reviews(fixture, native, mutate):
    c, a = fixture
    mutate(c, a)
    with pytest.raises(ValueError, match="Review is not bound"):
        plan(Catalog.model_validate(c), Assessment.model_validate(a), native)


def test_hipaa_paragraphs_preserve_addressable_status(native):
    c = starter(native)
    r = next(r for r in c.requirements if r.id == "HIPAA:164.308")
    assert next(a for a in r.attributes if a.id == "a(3)(ii)(C)").specification == "ADDRESSABLE"
    mapping = next(
        m for m in c.mappings if m.requirement_id == r.id and m.attribute_ids == ["a(3)(ii)(C)"]
    )
    assert mapping.control_id == "SH-IAM-004"
    assert mapping.review is None
    assert all(r.review is None for r in c.requirements if r.framework_id == "HIPAA")


def test_nested_collection_mutation_cannot_skip_schema(native):
    c, a = example(native)
    a.scope.baseline.clear()
    with pytest.raises(ValueError):
        plan(c, a, native)


def test_native_control_cannot_support_a_period_before_its_version(fixture, native):
    c, a = fixture
    a["scope"]["period_start"] = "2026-09-01"
    row = first(run(c, a, native))
    assert row["status"] != "SUPPORTED"
    assert "NATIVE_CONTROL_PERIOD_GAP" in row["attributes"][0]["issues"]
