# ruff: noqa: E402
# Apply hard limits before importing document parser dependencies.
"""Standalone parser worker. Invoked only inside parser_sandbox namespaces."""

import resource
import sys

if __name__ != "__main__":
    raise RuntimeError("Parser worker must run only as an isolated process")

resource.setrlimit(resource.RLIMIT_CPU, (10, 10))
resource.setrlimit(resource.RLIMIT_AS, (768 * 1024 * 1024, 768 * 1024 * 1024))
resource.setrlimit(resource.RLIMIT_FSIZE, (8 * 1024 * 1024, 8 * 1024 * 1024))
resource.setrlimit(resource.RLIMIT_NOFILE, (32, 32))

sys.path.insert(0, "/packages")
import csv
import hashlib
import io
import json
import re
import zipfile
from html.parser import HTMLParser
from pathlib import Path, PurePosixPath

DomainError = ValueError
MAX_BYTES = 25 * 1024 * 1024
MAX_EXPANDED = 100 * 1024 * 1024
PARSER_VERSION = "audit-review-isolated-1"
MIMES = {
    "csv": "text/csv",
    "json": "application/json",
    "html": "text/html",
    "log": "text/plain",
    "txt": "text/plain",
    "pdf": "application/pdf",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "png": "image/png",
}


def digest(value):
    return hashlib.sha256(
        json.dumps(value, sort_keys=True, separators=(",", ":"), allow_nan=False).encode()
    ).hexdigest()


def safe_name(name: str) -> str:
    if (
        not isinstance(name, str)
        or not name
        or len(name) > 180
        or name in {".", ".."}
        or any(ord(c) < 32 for c in name)
        or re.search(r"[/\\:]", name)
    ):
        raise DomainError("Invalid artifact filename")
    return name


class _Text(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts, self.hidden = [], 0

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style"}:
            self.hidden += 1

    def handle_endtag(self, tag):
        if tag in {"script", "style"} and self.hidden:
            self.hidden -= 1

    def handle_data(self, data):
        if not self.hidden and data.strip():
            self.parts.append(data.strip())


def extract(manifest: dict, data: bytes) -> dict:
    """Return inert text with exact original locations; never execute attachments."""
    if manifest.get("status") != "AVAILABLE":
        return {"artifact_id": manifest["id"], "status": "QUARANTINED", "sources": []}
    sources = []

    def add(locator: str, text: str):
        if text.strip() and len(sources) < 5000:
            sources.append(
                {
                    "id": manifest["id"] + ":" + locator,
                    "artifact_id": manifest["id"],
                    "locator": locator,
                    "text": text[:12000],
                    "text_truncated": len(text) > 12000,
                    "full_text_characters": len(text),
                    "full_text_digest": digest(text),
                    "original_sha256": manifest["sha256"],
                }
            )

    extension = Path(manifest["name"]).suffix.lower()
    if extension == ".pdf":
        import fitz

        with fitz.open(stream=data, filetype="pdf") as document:
            for i, page in enumerate(document):
                add(f"page:{i + 1}", page.get_text())
    elif extension == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(
            io.BytesIO(data), read_only=True, data_only=False, keep_links=False
        )
        try:
            for sheet in workbook:
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.value is not None:
                            add(f"sheet:{sheet.title}!{cell.coordinate}", str(cell.value))
                    if len(sources) >= 5000:
                        break
                if len(sources) >= 5000:
                    break
        finally:
            workbook.close()
    elif extension == ".csv":
        for line, row in enumerate(csv.reader(io.StringIO(data.decode("utf-8-sig"))), 1):
            add(f"record:{line}", " | ".join(row))
            if len(sources) >= 5000:
                break
    elif extension == ".json":
        document = json.loads(data)

        def walk(value, pointer="", depth=0):
            if depth > 20 or len(sources) >= 5000:
                return
            if isinstance(value, dict):
                for key, item in value.items():
                    walk(
                        item,
                        pointer + "/" + str(key).replace("~", "~0").replace("/", "~1"),
                        depth + 1,
                    )
            elif isinstance(value, list):
                for index, item in enumerate(value):
                    walk(item, pointer + "/" + str(index), depth + 1)
            else:
                add("json:" + pointer, str(value))

        walk(document)
    elif extension in {".txt", ".log", ".html"}:
        text = data.decode("utf-8-sig")
        if extension == ".html":
            parser = _Text()
            parser.feed(text)
            add("html:text", "\n".join(parser.parts))
        else:
            for line, content in enumerate(text.splitlines(), 1):
                add(f"line:{line}", content)
    else:
        return {
            "artifact_id": manifest["id"],
            "status": "HUMAN_REVIEW_REQUIRED",
            "sources": [],
            "limitation": "This file has no validated text extraction path",
        }
    result = {
        "artifact_id": manifest["id"],
        "status": "EXTRACTED",
        "parser": PARSER_VERSION,
        "sources": sources,
        "limits": {"locations": 5000, "characters_per_location": 12000},
    }
    result["extraction_digest"] = digest(result)
    return result


def inspect_upload(name: str, data: bytes) -> dict:
    """Originals can be retained quarantined; untrusted content is never executed."""
    safe_name(name)
    if not data or len(data) > MAX_BYTES:
        raise DomainError("Empty or oversized upload")
    suffix = Path(name).suffix.lower().lstrip(".")
    reason = None
    if suffix not in MIMES:
        reason = "Unsupported file type; retained without preview"
    elif suffix == "xlsx":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                infos = archive.infolist()
                if len(infos) > 2000 or sum(i.file_size for i in infos) > MAX_EXPANDED:
                    raise ValueError("Archive expansion limit")
                names = set()
                for item in infos:
                    p = PurePosixPath(item.filename)
                    if (
                        p.is_absolute()
                        or ".." in p.parts
                        or "\\" in item.filename
                        or item.filename in names
                        or item.flag_bits & 1
                        or (item.external_attr >> 16) & 0o170000 == 0o120000
                    ):
                        raise ValueError("Unsafe archive entry")
                    names.add(item.filename)
                    if any(
                        part.lower() in item.filename.lower()
                        for part in ("vbaproject", "externallinks", "embeddings")
                    ):
                        raise ValueError("Active or external workbook content")
                    if item.filename.endswith((".xml", ".rels")):
                        content = archive.read(item)
                        if re.search(
                            rb"<!DOCTYPE|<!ENTITY|TargetMode\s*=\s*['\"]External", content, re.I
                        ):
                            raise ValueError("External references or XML entities")
                if not {"[Content_Types].xml", "xl/workbook.xml"} <= names:
                    raise ValueError("Not an XLSX workbook")
        except (ValueError, zipfile.BadZipFile, RuntimeError) as exc:
            reason = f"Workbook quarantined: {exc}"
    elif suffix == "pdf":
        if not data.startswith(b"%PDF-"):
            reason = "Invalid PDF signature"
        elif re.search(rb"/(JavaScript|JS|Launch|EmbeddedFile|OpenAction|AA)\b", data):
            reason = "Active or embedded PDF content"
        else:
            try:
                import fitz

                with fitz.open(stream=data, filetype="pdf") as doc:
                    if doc.is_encrypted or doc.page_count > 2000:
                        raise ValueError("Encrypted or excessive PDF")
                    # Check decompressed objects, not just visible raw bytes.
                    for i in range(1, doc.xref_length()):
                        if re.search(
                            r"/(JavaScript|JS|Launch|EmbeddedFile|OpenAction|AA)\b",
                            doc.xref_object(i),
                        ):
                            raise ValueError("Active PDF object")
            except Exception as exc:
                reason = f"PDF quarantined: {type(exc).__name__}"
    elif suffix == "png" and not data.startswith(b"\x89PNG\r\n\x1a\n"):
        reason = "Invalid PNG signature"
    elif suffix == "html":
        reason = "HTML upload retained as attachment; active preview prohibited"
    elif suffix in {"csv", "json", "txt", "log"}:
        try:
            text = data.decode("utf-8-sig")
            if "\x00" in text:
                raise ValueError("Binary content")
            if suffix == "json":
                json.loads(text)
        except (ValueError, UnicodeError):
            reason = "Invalid text content"
    return {
        "status": "QUARANTINED" if reason else "AVAILABLE",
        "reason": reason,
        "mime": MIMES.get(suffix, "application/octet-stream"),
    }


if __name__ == "__main__":
    try:
        mode, name, identity, sha = sys.argv[1:]
        data = Path("/input").read_bytes()
        if len(data) > MAX_BYTES:
            raise ValueError("Input limit")
        result = (
            inspect_upload(name, data)
            if mode == "inspect"
            else extract({"name": name, "id": identity, "sha256": sha, "status": "AVAILABLE"}, data)
        )
        encoded = json.dumps(result, allow_nan=False).encode()
        if len(encoded) > 8 * 1024 * 1024:
            raise ValueError("Output limit")
        sys.stdout.buffer.write(encoded)
    except BaseException:
        sys.stdout.write('{"sandbox_error":"PARSER_FAILED_OR_LIMIT_EXCEEDED"}')
        sys.exit(2)
