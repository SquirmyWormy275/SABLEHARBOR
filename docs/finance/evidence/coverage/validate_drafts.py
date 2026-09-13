"""Check every draft workbook cell against complete source populations and frozen source hashes."""

import csv
import hashlib
import json
from decimal import Decimal
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[4]
BASE = ROOT / "docs/finance/evidence"


def validate():
    counts = {}
    for family in ("customer", "treasury", "close", "supporting-schedules", "tax-transaction"):
        folder = BASE / family
        draft = folder / "draft"
        manifest = json.loads((draft / "manifest.json").read_text())
        assert (
            manifest["source_register_sha256"]
            == hashlib.sha256((folder / "evidence-register.json").read_bytes()).hexdigest()
        ), family
        for relative, digest in manifest.get("source_dependencies", {}).items():
            assert hashlib.sha256((ROOT / relative).read_bytes()).hexdigest() == digest, relative
        for relative, digest in manifest["artifacts"].items():
            assert hashlib.sha256((ROOT / relative).read_bytes()).hexdigest() == digest, relative
        workbook = openpyxl.load_workbook(
            draft / "working-papers.xlsx", data_only=True, read_only=True
        )
        population = 0
        for item in json.loads((draft / "workbook-map.json").read_text()):
            path = ROOT / item["source"]
            assert hashlib.sha256(path.read_bytes()).hexdigest() == item["source_sha256"], item[
                "source"
            ]
            rows = list(csv.DictReader(path.open()))
            assert len(rows) == item["row_count"], item["sheet"]
            sheet = workbook[item["sheet"]]
            cells = sheet.iter_rows(
                min_row=item["first_data_row"],
                max_row=item["first_data_row"] + len(rows) - 1,
                max_col=len(item["columns"]),
                values_only=True,
            )
            for source, values in zip(rows, cells, strict=True):
                for column, value in zip(item["columns"], values, strict=True):
                    expected = source[column]
                    if (
                        column.endswith("_usd")
                        and expected != ""
                        and len(Decimal(expected).normalize().as_tuple().digits) <= 15
                    ):
                        assert value == float(expected), (item["sheet"], column, expected, value)
                    else:
                        assert (value or "") == expected, (item["sheet"], column, expected, value)
            population += len(rows)
        for row in workbook["Checks"].iter_rows(min_row=5, values_only=True):
            assert row[3] == "PASS", (family, row)
        assert not workbook._external_links
        workbook.close()
        counts[family] = population
    print("PASS complete draft workbook cells and source/artifact hashes:", counts)


if __name__ == "__main__":
    validate()
