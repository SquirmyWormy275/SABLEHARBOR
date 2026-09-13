"""Public worked calculations must retain source lineage and blank response cells."""

import importlib.util
import json
import shutil
import sqlite3
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
SPEC = importlib.util.spec_from_file_location(
    "legal_practice", ROOT / "tools/legal_gaps/practice.py"
)
assert SPEC and SPEC.loader
practice = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(practice)


@pytest.fixture
def practice_copy(tmp_path):
    paths = {Path("tools/legal_gaps/practice.py")}
    for packet in practice.make_packets(ROOT):
        paths.update(Path(e["path"]) for e in packet["evidence"])
    for source in paths:
        target = tmp_path / source
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(ROOT / source, target)
    shutil.copytree(ROOT / practice.DEST, tmp_path / practice.DEST)
    assert practice.validate(tmp_path) == []
    return tmp_path


def test_current_public_packets():
    assert practice.validate(ROOT) == []
    packets = practice.make_packets(ROOT)
    assert len(packets) == 4
    assert all(p["classification"] == "PUBLIC_WORKED_EXAMPLE" for p in packets)
    results = {p["slug"]: {c["id"]: c["expected_usd"] for c in p["calculations"]} for p in packets}
    assert results["acquisition"]["uses"] == 61_500_000
    assert results["acquisition"]["sources"] == 61_500_000
    assert 14_762_500 in results["acquisition"].values()
    assert 13_000_000 in results["acquisition"].values()
    assert 23_383_219 in results["debt"].values()
    assert 609_000 in results["revenue-dispute"].values()
    assert 971_500 in results["revenue-dispute"].values()
    assert 3_010 in results["legal-due-diligence"].values()


def test_changed_source_rejected(practice_copy):
    path = practice_copy / "industrial/source/finance.json"
    path.write_bytes(path.read_bytes() + b"\n")
    assert any("source/packet drift" in e for e in practice.validate(practice_copy))


def test_blank_answer_rejected(practice_copy):
    path = practice_copy / practice.DEST / "acquisition/blank.xlsx"
    wb = load_workbook(path)
    wb["Working paper"]["B4"] = 123
    wb.save(path)
    assert any("blank answer is populated" in e for e in practice.validate(practice_copy))


def test_worked_formula_rejected(practice_copy):
    path = practice_copy / practice.DEST / "debt/worked.xlsx"
    wb = load_workbook(path)
    wb["Working paper"]["B4"] = "=1+1"
    wb.save(path)
    assert any("formula differs" in e for e in practice.validate(practice_copy))


def test_input_amount_rejected(practice_copy):
    path = practice_copy / practice.DEST / "revenue-dispute/blank.xlsx"
    wb = load_workbook(path)
    wb["Source inputs"]["B4"] = 1
    wb.save(path)
    assert any("source input drift" in e for e in practice.validate(practice_copy))


def test_database_finding_rejected(practice_copy):
    db = sqlite3.connect(practice_copy / practice.DEST / "practice.sqlite3")
    db.execute(
        "UPDATE finding SET record_json=? WHERE id=?",
        (json.dumps({"unsupported": True}), "LEGAL-F01"),
    )
    db.commit()
    db.close()
    assert any("SQLite finding drift" in e for e in practice.validate(practice_copy))


def test_missing_workbook_rejected(practice_copy):
    (practice_copy / practice.DEST / "acquisition/blank.xlsx").unlink()
    assert practice.validate(practice_copy)
