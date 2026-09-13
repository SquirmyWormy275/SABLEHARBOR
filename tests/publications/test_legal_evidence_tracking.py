"""Exercise real evidence proposals and reject unsafe or stale changes."""

import copy
from pathlib import Path

import pytest
from openpyxl import load_workbook

from tools.legal_gaps import evidence_tracking as et


@pytest.fixture
def completed(tmp_path):
    path = tmp_path / "completed.xlsx"
    et.workbook(et.model(), path)
    wb = load_workbook(path)
    evidence = Path("industrial/transaction/05_ARU_CLOSING_AND_TAX_DELIVERY.md")
    citation = next(
        line for line in (et.ROOT / evidence).read_text().splitlines() if len(line) >= 30
    )
    values = [
        "incomplete",
        str(evidence),
        et.sha(et.ROOT / evidence),
        citation,
        "Existing model delivery memo supplied; external bank confirmation remains missing.",
        "TEST-EVENT-001",
        "2026-09-13",
    ]
    for col, value in enumerate(values, 3):
        wb["Updates"].cell(2, col, value)
    wb.save(path)
    return path


def test_populated_import_and_append_preserve_sources(completed, tmp_path):
    pins = et.model()["source_pins"]
    out = tmp_path / "review1"
    state = et.import_workbook(completed, out)
    assert state["current"]["WALK-REQ-ARU-01"] == "incomplete"
    assert state["history"][0]["lines"][0] > 0
    assert et.model()["source_pins"] == pins
    wb = load_workbook(out / "tracker.xlsx")
    for col, value in enumerate(
        [
            "disputed",
            state["history"][0]["evidence_path"],
            state["history"][0]["evidence_sha256"],
            state["history"][0]["citation"],
            "Memo still lacks independent creditor discharge; sufficiency is disputed.",
            "TEST-EVENT-002",
            "2026-09-14",
        ],
        3,
    ):
        wb["Updates"].cell(2, col, value)
    copy_path = tmp_path / "second.xlsx"
    wb.save(copy_path)
    updated = et.import_workbook(copy_path, tmp_path / "review2", state_path=out / "tracker.json")
    assert updated["history"][:1] == state["history"]
    assert updated["current"]["WALK-REQ-ARU-01"] == "disputed"
    assert len(updated["history"]) == 2


@pytest.mark.parametrize(
    ("sheet", "row", "col", "value"),
    [
        ("Updates", 2, 1, "UNKNOWN-ID"),
        ("Updates", 2, 2, "received"),
        ("Updates", 2, 3, "closed"),
        ("Updates", 2, 3, "unresolved"),
        ("Updates", 2, 4, "../secret.md"),
        ("Updates", 2, 4, "/etc/passwd"),
        ("Updates", 2, 5, "0" * 64),
        ("Updates", 2, 6, "Not a citation from this record"),
        ("Updates", 2, 7, "=1+1"),
        ("Updates", 2, 8, "bad event id"),
        ("Updates", 2, 9, "2026-02-30"),
        ("Context", 2, 2, "changed"),
        ("Requests", 2, 3, "Altered scope"),
    ],
)
def test_reject_bad_input(completed, tmp_path, sheet, row, col, value):
    wb = load_workbook(completed)
    wb[sheet].cell(row, col, value)
    wb.save(completed)
    with pytest.raises(ValueError):
        et.import_workbook(completed, tmp_path / "rejected")
    assert not (tmp_path / "rejected").exists()


def test_missing_evidence_for_received(completed, tmp_path):
    wb = load_workbook(completed)
    wb["Updates"].cell(2, 3, "received")
    for col in [4, 5, 6]:
        wb["Updates"].cell(2, col).value = None
    wb.save(completed)
    with pytest.raises(ValueError):
        et.import_workbook(completed, tmp_path / "rejected")


def test_duplicate_events(completed, tmp_path):
    state = et.import_workbook(completed, tmp_path / "review")
    wb = load_workbook(tmp_path / "review/tracker.xlsx")
    original = load_workbook(completed)
    for col in range(3, 10):
        wb["Updates"].cell(3, col, original["Updates"].cell(2, col).value)
    p = tmp_path / "duplicate.xlsx"
    wb.save(p)
    with pytest.raises(ValueError, match="Duplicate"):
        et.import_workbook(p, tmp_path / "bad", state_path=tmp_path / "review/tracker.json")
    assert len(state["history"]) == 1


def test_tampered_history(completed, tmp_path):
    state = et.import_workbook(completed, tmp_path / "review")
    state["history"][0]["note"] = "Altered historic assertion"
    with pytest.raises(ValueError, match="History chain"):
        et.validate_state(state)


def test_output_boundaries(completed, tmp_path):
    with pytest.raises(ValueError):
        et.import_workbook(completed, et.ROOT / "forbidden-output")
    with pytest.raises(ValueError):
        et.import_workbook(completed, tmp_path)


def test_empty_workbook(tmp_path):
    p = tmp_path / "empty.xlsx"
    et.workbook(et.model(), p)
    with pytest.raises(ValueError, match="No proposed"):
        et.import_workbook(p, tmp_path / "bad")


def test_baseline_derivatives():
    et.validate()


def test_state_cannot_change_context():
    state = copy.deepcopy(et.model())
    state["requests"][0]["request"] = "Changed request"
    with pytest.raises(ValueError, match="context"):
        et.validate_state(state)


def test_symlink_evidence_rejected(tmp_path):
    (tmp_path / "docs").mkdir()
    target = tmp_path / "real.md"
    target.write_text("This is an exact sufficiently long citation.")
    (tmp_path / "docs/link.md").symlink_to(target)
    with pytest.raises(ValueError, match="Symlink"):
        et.evidence(tmp_path, "docs/link.md", et.sha(target), target.read_text())


def test_deleted_history_event_rejected(completed, tmp_path):
    state = et.import_workbook(completed, tmp_path / "review")
    state["history"] = []
    with pytest.raises(ValueError, match="reconcile"):
        et.validate_state(state)


def test_received_retains_blocked_conclusion(completed, tmp_path):
    wb = load_workbook(completed)
    wb["Updates"].cell(2, 3, "received")
    wb.save(completed)
    state = et.import_workbook(completed, tmp_path / "review")
    assert state["current"]["WALK-REQ-ARU-01"] == "received"
    assert (
        state["requests"][0]["blocked_conclusion"]
        == et.model()["requests"][0]["blocked_conclusion"]
    )
    assert state["status"] == et.STATUS


def test_aru_close_requests_are_separate_from_original_population():
    data = et.model()
    assert len(data["requests"]) == 14
    close = data["requests"][11:]
    assert [r["id"] for r in close] == [
        f"SH-CLOSE-ARU-2027-01-{area}-REQ" for area in ["AR", "AP", "PPE"]
    ]
    assert all(
        r["scope"] == "ARU_GROUP industrial reporting book; base scenario; January 2027"
        for r in close
    )
    assert [r["source_locator"] for r in close] == [
        "/limitations/2",
        "/limitations/3",
        "/limitations/4",
    ]
    assert all(value == "unresolved" for value in data["current"].values())
    old = {r["id"]: r for r in data["requests"][:11]}
    assert old["SH-RECON-AR-REQ"]["scope"].startswith("Atlas Meridian")
    assert old["SH-RECON-FIXED-ASSETS-REQ"]["scope"].startswith("Willow")
