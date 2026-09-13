"""Protection against missing evidence, stale derivatives and false exercise blanks."""

import importlib.util
import json
import shutil
import sqlite3
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[2]
SPEC = importlib.util.spec_from_file_location(
    "legal_reconciliations", ROOT / "tools/legal_gaps/reconciliations.py"
)
M = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(M)


def fixture_root(tmp_path):
    data = M.model(ROOT)
    for e in data["evidence"]:
        p = tmp_path / e["path"]
        p.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(ROOT / e["path"], p)
    M.generate(tmp_path, tmp_path / M.DEST)
    return tmp_path


def test_current_package():
    M.validate(ROOT)
    data = M.model(ROOT)
    assert len(data["workpapers"]) == 5
    assert sum(len(p["checks"]) for p in data["workpapers"]) == 9
    assert all(c["result"] in ["0", "0.0000"] for p in data["workpapers"] for c in p["checks"])


def test_workbooks_separate_answers():
    blank = openpyxl.load_workbook(ROOT / M.DEST / "blank.xlsx", data_only=False)
    worked = openpyxl.load_workbook(ROOT / M.DEST / "worked.xlsx", data_only=False)
    cached = openpyxl.load_workbook(ROOT / M.DEST / "worked.xlsx", data_only=True)
    assert len(blank.sheetnames) == len(worked.sheetnames) == 7
    count = 0
    for ws in worked:
        for row in ws:
            for cell in row:
                if cell.data_type == "f":
                    count += 1
                    assert blank[ws.title][cell.coordinate].value is None
                    assert abs(cached[ws.title][cell.coordinate].value) < 0.02
    assert count == 9


def test_source_mutation_fails(tmp_path):
    root = fixture_root(tmp_path)
    path = root / M.model(root)["evidence"][0]["path"]
    path.write_text(path.read_text().replace("4198440", "4198441", 1))
    with pytest.raises(ValueError):
        M.validate(root)


def test_deleted_row_fails(tmp_path):
    root = fixture_root(tmp_path)
    path = root / M.DEST / "source.json"
    obj = json.loads(path.read_text())
    obj["evidence"][0]["rows"] = []
    M.dump(path, obj)
    with pytest.raises(ValueError, match="hash drift"):
        M.validate(root)


def test_sqlite_rehashed_row_mutation_fails(tmp_path):
    root = fixture_root(tmp_path)
    db = root / M.DEST / "reconciliations.sqlite3"
    with sqlite3.connect(db) as cx:
        cx.execute("UPDATE records SET payload='{}' WHERE kind='package'")
    path = root / M.DEST / "manifest.json"
    obj = json.loads(path.read_text())
    obj["files"][db.name] = M.sha(db)
    M.dump(path, obj)
    with pytest.raises(ValueError, match="contents drift"):
        M.validate(root)


def test_sqlite_rehashed_extra_schema_fails(tmp_path):
    root = fixture_root(tmp_path)
    db = root / M.DEST / "reconciliations.sqlite3"
    with sqlite3.connect(db) as cx:
        cx.execute("CREATE TABLE extra(x)")
    path = root / M.DEST / "manifest.json"
    obj = json.loads(path.read_text())
    obj["files"][db.name] = M.sha(db)
    M.dump(path, obj)
    with pytest.raises(ValueError, match="contents drift"):
        M.validate(root)


def test_deterministic_outputs(tmp_path):
    M.generate(ROOT, tmp_path / "one")
    M.generate(ROOT, tmp_path / "two")
    assert {p.name: p.read_bytes() for p in (tmp_path / "one").iterdir()} == {
        p.name: p.read_bytes() for p in (tmp_path / "two").iterdir()
    }


def test_duplicate_snapshot_rejected(tmp_path):
    root = fixture_root(tmp_path)
    path = root / M.model(root)["evidence"][0]["path"]
    lines = path.read_text().splitlines()
    path.write_text("\n".join(lines + [lines[1]]) + "\n")
    with pytest.raises(ValueError, match="Expected one snapshot"):
        M.validate(root)
