"""Fail closed on scoped source, arithmetic and derived artifact mutations."""

import csv
import json
import shutil
import sqlite3

import openpyxl
import pytest

from tools.legal_gaps import period_close as p


@pytest.fixture
def root(tmp_path):
    shutil.copytree(p.ROOT / p.DEST, tmp_path / p.DEST, ignore=shutil.ignore_patterns("qa"))
    pins = json.loads((tmp_path / p.DEST / "inputs.json").read_text())
    for s in pins["sources"]:
        if "upstream_path" in s:
            dest = tmp_path / s["upstream_path"]
            dest.parent.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(p.ROOT / s["upstream_path"], dest)
    return tmp_path


def mutate(root, name, edit, repin=False):
    path = root / p.DEST / "source" / (name + ".csv")
    rows = list(csv.DictReader(path.open()))
    fields = list(rows[0])
    edit(rows)
    with path.open("w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)
    if repin:
        loc = root / p.DEST / "inputs.json"
        pins = json.loads(loc.read_text())
        source = next(s for s in pins["sources"] if s["id"] == name)
        source["sha256"] = p.sha(path)
        p.dump(loc, pins)


def test_complete_single_scope_and_checks():
    data = p.model()
    assert len(data["tables"]["journal"]) == 274
    assert len(data["tables"]["opening_balances"]) == 22
    assert len(data["trial_balance"]) == 34
    assert len(data["checks"]) == 18
    assert data["proposed_adjustments"] == []
    assert all(p.D(c["difference"]) == 0 for c in data["checks"])
    assert all(r["month"] == "1" for r in data["tables"]["journal"])


def test_source_edit_rejected(root):
    mutate(root, "journal", lambda rows: rows[0].update(signed_usd="1"))
    with pytest.raises(ValueError, match="Source hash"):
        p.model(root)


def test_scope_rejected_even_if_rehashed(root):
    mutate(root, "journal", lambda rows: rows[0].update(entity="ATLAS"), True)
    with pytest.raises(ValueError, match="scope/coverage"):
        p.model(root)


def test_missing_row_rejected_even_if_rehashed(root):
    mutate(root, "journal", lambda rows: rows.pop(), True)
    with pytest.raises(ValueError, match="scope/coverage"):
        p.model(root)


def test_unbalanced_posting_rejected(root):
    mutate(root, "journal", lambda rows: rows[0].update(signed_usd="1", debit_usd="1"), True)
    with pytest.raises(ValueError, match="Unbalanced"):
        p.model(root)


def test_balanced_account_reclassification_rejected(root):
    mutate(root, "journal", lambda rows: rows[0].update(account="1200"), True)
    with pytest.raises(ValueError, match="rollforward"):
        p.model(root)


def test_bank_discrepancy_rejected(root):
    mutate(
        root,
        "industrial_bank_reconciliations",
        lambda rows: rows[0].update(adjusted_bank_cash_usd="1"),
        True,
    )
    with pytest.raises(ValueError, match="Bank adjusted"):
        p.model(root)


def test_asset_discrepancy_rejected(root):
    mutate(root, "assets", lambda rows: rows[0].update(gross_usd="1"), True)
    with pytest.raises(ValueError, match="PPE gross"):
        p.model(root)


def test_blank_cells_and_worked_formulas():
    blank = openpyxl.load_workbook(p.ROOT / p.DEST / "blank.xlsx")
    worked = openpyxl.load_workbook(p.ROOT / p.DEST / "worked.xlsx")
    assert len(blank.sheetnames) == len(worked.sheetnames) == 17
    for name, cells in {
        "Closing TB": ["C5", "D38", "F38", "H38"],
        "Five reconciliations": ["B5", "D22"],
        "PPE": ["F12"],
        "Debt": ["F7"],
        "Statements": ["B13", "D13"],
    }.items():
        for cell in cells:
            assert blank[name][cell].value is None
            assert worked[name][cell].data_type == "f"
    assert worked["Activity"].max_row == 278
    assert worked["Receivables"].max_row == 34
    assert worked["Payables"].max_row == 77


def test_reproducible_and_saved_integrity(root):
    p.validate(root)
    target = root / p.DEST / "worked.xlsx"
    target.write_bytes(target.read_bytes() + b"changed")
    with pytest.raises(ValueError, match="artifact drift"):
        p.validate(root)


def test_database_rehashed_mutation_rejected(root):
    target = root / p.DEST / "close.sqlite3"
    with sqlite3.connect(target) as db:
        db.execute("UPDATE journal SET signed_usd='1' WHERE rowid=1")
    manifest_path = root / p.DEST / "manifest.json"
    manifest = json.loads(manifest_path.read_text())
    manifest["files"]["close.sqlite3"] = p.sha(target)
    p.dump(manifest_path, manifest)
    with pytest.raises(ValueError, match="Database contents"):
        p.validate(root)


def test_deleted_manifest_artifact_rejected(root):
    target = root / p.DEST / "manifest.json"
    manifest = json.loads(target.read_text())
    del manifest["files"]["blank.xlsx"]
    p.dump(target, manifest)
    with pytest.raises(ValueError, match="Manifest coverage"):
        p.validate(root)


def test_bank_identity_rejected_despite_equal_total(root):
    mutate(
        root, "industrial_bank_transactions", lambda rows: rows[0].update(journal_id="wrong"), True
    )
    with pytest.raises(ValueError, match="Bank event identity"):
        p.model(root)


def test_database_metadata_rejected(root):
    target = root / p.DEST / "close.sqlite3"
    with sqlite3.connect(target) as db:
        db.execute("PRAGMA user_version=99")
    loc = root / p.DEST / "manifest.json"
    manifest = json.loads(loc.read_text())
    manifest["files"]["close.sqlite3"] = p.sha(target)
    p.dump(loc, manifest)
    with pytest.raises(ValueError, match="Database contents"):
        p.validate(root)


def test_qa_workbook_change_rejected(root):
    qa = root / p.DEST / "qa"
    qa.mkdir()
    shutil.copyfile(p.ROOT / p.DEST / "qa/REVIEW.json", qa / "REVIEW.json")
    target = root / p.DEST / "blank.xlsx"
    target.write_bytes(target.read_bytes() + b"changed")
    with pytest.raises(ValueError, match="Reviewed workbook changed"):
        p.validate_qa(root)


def test_qa_omitted_page_rejected(root):
    qa = root / p.DEST / "qa"
    qa.mkdir()
    receipt = json.loads((p.ROOT / p.DEST / "qa/REVIEW.json").read_text())
    receipt["records"][0]["pages"].pop()
    p.dump(qa / "REVIEW.json", receipt)
    with pytest.raises(ValueError, match="Page QA coverage"):
        p.validate_qa(root)
