"""A filled response copy cannot mutate or adopt its source documents."""

import difflib
import importlib.util
import json
import shutil
import sqlite3
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
spec = importlib.util.spec_from_file_location(
    "decision_import", ROOT / "tools/legal_gaps/import_decisions.py"
)
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)


@pytest.fixture
def workspace(tmp_path):
    root = tmp_path / "repository"
    base = root / module.BASE
    shutil.copytree(ROOT / module.BASE / "source", base / "source")
    shutil.copytree(
        ROOT / module.BASE / "review-support",
        base / "review-support",
        ignore=shutil.ignore_patterns("qa"),
    )
    workbook = tmp_path / "completed.xlsx"
    shutil.copyfile(base / "review-support/decisions.xlsx", workbook)
    return root, workbook, tmp_path / "output"


def edit(workbook, changes):
    wb = load_workbook(workbook)
    for sheet, cell, value in changes:
        wb[sheet][cell] = value
    wb.save(workbook)
    wb.close()


def test_explicit_replacement_exact_diff_database_and_no_writeback(workspace):
    root, workbook, output = workspace
    before = {p: p.read_bytes() for p in root.rglob("*") if p.is_file()}
    edit(
        workbook,
        [
            ("Review", "E2", 'SET {"term":"Proposed revised text for review."}'),
            ("Review", "E3", "approve"),
            ("Review", "E4", "KEEP"),
            ("Source details", "I5", "DISCUSS: Need evidence first."),
        ],
    )
    result = module.import_workbook(workbook, output, root)
    assert len(result["items"]) == 125 and len(result["changes"]) == 1
    change = result["changes"][0]
    assert change["pointer"] == "/proposed_terms/0/term"
    assert change["id"] == "DEBT-P01"
    assert change["after"] == "Proposed revised text for review."
    proposed = json.loads((output / "proposed" / change["source_path"]).read_text())
    original = json.loads((root / change["source_path"]).read_text())
    original["proposed_terms"][0]["term"] = change["after"]
    assert proposed == original
    assert result["items"][1]["disposition"] == "COMMENT_ONLY_NOT_APPROVAL"
    assert not result["acceptance_created"] and not result["source_writeback"]
    assert '-      "term": "ARU' in (output / "debt-liens.diff").read_text()
    assert '+      "term": "Proposed revised' in (output / "debt-liens.diff").read_text()
    expected_diff = "".join(
        difflib.unified_diff(
            (root / change["source_path"]).read_text().splitlines(keepends=True),
            (output / "proposed" / change["source_path"]).read_text().splitlines(keepends=True),
            fromfile=change["source_path"],
            tofile="proposed/" + change["source_path"],
        )
    )
    assert (output / "debt-liens.diff").read_text() == expected_diff
    with sqlite3.connect(output / "review.sqlite3") as db:
        assert db.execute("SELECT count(*) FROM response").fetchone()[0] == 125
        assert db.execute("SELECT count(*) FROM proposed_change").fetchone()[0] == 1
    assert all(p.read_bytes() == content for p, content in before.items())


@pytest.mark.parametrize(
    "sheet,cell,value",
    [
        ("Review", "A2", "UNKNOWN-ID"),
        ("Review", "A3", "DEBT-P01"),
        ("Review", "C2", "Changed question"),
        ("Source details", "H2", "../../escape.json#/proposed_terms/0"),
        ("Review", "E127", "extra item"),
        ("Review", "E2", "=1+1"),
        ("Source details", "I2", '=HYPERLINK("https://example.invalid")'),
        ("Review", "E2", 'SET {"status":"LOCKED"}'),
        ("Review", "E2", 'SET {"term":null}'),
        ("Review", "E2", 'SET {"term":"one","term":"two"}'),
        ("Review", "E2", 'SET {"../term":"escape"}'),
        ("Review", "E2", "set the term to something"),
        ("Review", "E2", "KEEP but change the amount"),
        ("Review", "E2", "@SUM(A1:A5)"),
    ],
)
def test_rejects_tampering_and_ambiguous_controls(workspace, sheet, cell, value):
    root, workbook, output = workspace
    edit(workbook, [(sheet, cell, value)])
    with pytest.raises(ValueError):
        module.import_workbook(workbook, output, root)
    assert not output.exists()


def test_conflicting_response_columns(workspace):
    root, workbook, output = workspace
    edit(workbook, [("Review", "E2", "KEEP"), ("Source details", "I2", "KEEP")])
    with pytest.raises(ValueError, match="Two responses"):
        module.import_workbook(workbook, output, root)
    assert not output.exists()


@pytest.mark.parametrize("mutation", ["missing_row", "missing_sheet", "link"])
def test_rejects_missing_context_or_changed_link(workspace, mutation):
    root, workbook, output = workspace
    wb = load_workbook(workbook)
    if mutation == "missing_row":
        wb["Review"].delete_rows(2)
    elif mutation == "missing_sheet":
        del wb["Source details"]
    else:
        wb["Review"]["D2"].hyperlink = "https://example.invalid"
    wb.save(workbook)
    with pytest.raises(ValueError):
        module.import_workbook(workbook, output, root)
    assert not output.exists()


@pytest.mark.parametrize(
    "target",
    [
        "source/debt-liens.json",
        "source/debt-liens.md",
        "review-support/decisions.json",
        "review-support/decisions.xlsx",
    ],
)
def test_stale_source_or_baseline_fails(workspace, target):
    root, workbook, output = workspace
    path = root / module.BASE / target
    path.write_bytes(path.read_bytes() + b" ")
    with pytest.raises(ValueError, match="Stale|baseline"):
        module.import_workbook(workbook, output, root)
    assert not output.exists()


def test_output_must_be_new_and_outside_repository(workspace):
    root, workbook, output = workspace
    output.mkdir()
    with pytest.raises(ValueError, match="new directory"):
        module.import_workbook(workbook, output, root)
    with pytest.raises(ValueError, match="outside"):
        module.import_workbook(workbook, root / "new-output", root)
    assert not (root / "new-output").exists()


def test_blank_and_prose_do_not_create_changes(workspace):
    root, workbook, output = workspace
    edit(workbook, [("Review", "E2", "Reject this; the lender is unverified.")])
    result = module.import_workbook(workbook, output, root)
    assert result["changes"] == []
    assert result["items"][0]["disposition"] == "COMMENT_ONLY_NOT_APPROVAL"
    assert result["items"][1]["disposition"] == "UNANSWERED"
    assert not (output / "proposed").exists()


def test_source_and_input_symlinks_are_rejected(workspace, tmp_path):
    root, workbook, output = workspace
    link = tmp_path / "linked.xlsx"
    link.symlink_to(workbook)
    with pytest.raises(ValueError, match="symlink"):
        module.import_workbook(link, output, root)
    target = root / module.BASE / "source/debt-liens.json"
    external = tmp_path / "external.json"
    target.rename(external)
    target.symlink_to(external)
    with pytest.raises(ValueError, match="symlink"):
        module.import_workbook(workbook, output, root)


def test_unresolved_field_replacement_keeps_it_unresolved(workspace):
    root, workbook, output = workspace
    items = json.loads((root / module.BASE / "review-support/decisions.json").read_text())["items"]
    number, item = next(
        (n, r)
        for n, r in enumerate(items, 2)
        if r["json_pointer"].startswith("/unresolved_fields/")
    )
    edit(
        workbook,
        [
            (
                "Review",
                f"E{number}",
                'SET {"next_action":"Request the missing evidence for separate review."}',
            )
        ],
    )
    result = module.import_workbook(workbook, output, root)
    assert result["changes"][0]["pointer"].endswith("/next_action")
    assert result["changes"][0]["id"] == item["id"]
    proposed = json.loads((output / "proposed" / item["source_path"]).read_text())
    assert proposed["status"] == "DRAFT_FOR_REVIEW"
