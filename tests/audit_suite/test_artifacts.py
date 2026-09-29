import hashlib
import io
import json
import zipfile

import pytest

from enterprise.audit_suite.artifacts import Artifacts, inspect_upload, render
from enterprise.audit_suite.store import DomainError


@pytest.mark.parametrize("format_name", ["csv", "json", "html", "log", "xlsx", "pdf", "png"])
def test_native_outputs_and_redaction(format_name):
    data, mime = render(
        {
            "format": format_name,
            "title": "Fictional inspection record",
            "columns": ["id", "private"],
            "rows": [{"id": "X-1", "private": "SECRET-NEVER-DISCLOSE"}],
            "redact_columns": ["private"],
        }
    )
    assert data and b"SECRET-NEVER-DISCLOSE" not in data
    if format_name == "pdf":
        import fitz

        with fitz.open(stream=data, filetype="pdf") as doc:
            assert "SECRET-NEVER-DISCLOSE" not in "".join(p.get_text() for p in doc)
    if format_name == "xlsx":
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            assert all(b"SECRET-NEVER-DISCLOSE" not in z.read(n) for n in z.namelist())


def test_excel_safe_literals_and_real_arithmetic():
    from openpyxl import load_workbook

    data, _ = render(
        {
            "format": "xlsx",
            "columns": ["label", "amount"],
            "rows": [
                {"label": '=HYPERLINK("https://invalid.example")', "amount": 5},
                {"label": "Other", "amount": 7},
            ],
            "total_columns": ["amount"],
        }
    )
    wb = load_workbook(io.BytesIO(data), data_only=False)
    assert wb.active["A2"].data_type == "s"
    assert wb.active["B4"].value == "=SUM(B2:B3)"
    wb = load_workbook(io.BytesIO(data), data_only=True)
    assert wb.active["B4"].value == 12
    assert inspect_upload("ledger.xlsx", data)["status"] == "AVAILABLE"


def test_quarantine_active_pdf_and_archive_traversal():
    assert (
        inspect_upload("x.pdf", b"%PDF-1.7 /OpenAction /JavaScript evil")["status"] == "QUARANTINED"
    )
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("../escape.xml", "text")
    assert inspect_upload("x.xlsx", buf.getvalue())["status"] == "QUARANTINED"
    with pytest.raises(DomainError, match="filename"):
        inspect_upload("../escape.csv", b"id\n1")


def test_immutable_originals_safe_duplicate_names_and_integrity(tmp_path):
    store = Artifacts(tmp_path)
    a = store.retain("E-1", "record.csv", b"id\n1\n", source={"kind": "upload"}, coverage={})
    b = store.retain("E-1", "record.csv", b"id\n2\n", source={"kind": "upload"}, coverage={})
    with zipfile.ZipFile(io.BytesIO(store.bundle([a, b]))) as z:
        assert z.read(f"files/{a['id']}/record.csv") == b"id\n1\n"
        assert z.read(f"files/{b['id']}/record.csv") == b"id\n2\n"
        manifest = json.loads(z.read("manifest.json"))
        assert manifest["artifacts"][0]["sha256"] == hashlib.sha256(b"id\n1\n").hexdigest()
    (store.root / a["sha256"]).write_bytes(b"changed")
    with pytest.raises(DomainError, match="integrity"):
        store.read(a)
