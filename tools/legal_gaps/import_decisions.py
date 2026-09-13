"""Read a completed COPY of the frozen review workbook; emit proposals only."""

import argparse
import copy
import difflib
import hashlib
import json
import re
import sqlite3
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
BASE = Path("docs/legal/gap-instruments")
PINS = {
    "decisions.json": "42cd29aa93d64f1f57eecd85d92afe418eefc13f729a4d5468b37b95585b9cd8",
    "decisions.xlsx": "36394a1a6b2c13f1e4a721d904f6b5a956015610302b284b3121d72e2dd6823a",
}
STATUS = "PROPOSAL_ONLY_NOT_APPROVED_OR_APPLIED"


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def require(condition, message):
    if not condition:
        raise ValueError(message)


def safe_source(root, relative):
    path = Path(relative)
    require(not path.is_absolute() and ".." not in path.parts, "Unsafe source path")
    require(path.parent == BASE / "source", "Source outside instrument sources")
    target = root / path
    require(
        not target.is_symlink() and target.resolve().is_relative_to(root.resolve()),
        "Source symlink escape",
    )
    require(target.is_file(), f"Missing source: {relative}")
    return target


def strict_object(pairs):
    result = {}
    for key, value in pairs:
        require(key not in result, f"Duplicate JSON key: {key}")
        result[key] = value
    return result


def response(value, allowed):
    if value is None or value == "":
        return "UNANSWERED", {}
    require(isinstance(value, str), "Responses must be text")
    text = value.strip()
    if not text:
        return "UNANSWERED", {}
    require(not text.startswith(("=", "+", "-", "@")), "Formula-like response is not allowed")
    if text == "KEEP":
        return "KEEP_NO_CHANGE_NOT_APPROVAL", {}
    if text.startswith("DISCUSS:"):
        require(bool(text[8:].strip()), "DISCUSS requires a note")
        return "DISCUSSION_ONLY", {}
    if text.startswith("SET "):
        data = json.loads(text[4:], object_pairs_hook=strict_object)
        require(isinstance(data, dict) and bool(data), "SET requires a nonempty JSON object")
        require(set(data) <= allowed, "SET contains an unknown or forbidden field")
        for value in data.values():
            require(
                isinstance(value, str) and bool(value.strip()),
                "Replacement values must be nonempty strings",
            )
            require(
                not any(ord(c) < 32 and c not in "\n\t" for c in value),
                "Control character in replacement",
            )
        return "EXPLICIT_REPLACEMENT_PROPOSAL", data
    require(
        not re.match(r"(?i)^(set|keep|discuss)\b", text),
        "Ambiguous control syntax; see import guide",
    )
    return "COMMENT_ONLY_NOT_APPROVAL", {}


def hyperlink(cell):
    link = cell.hyperlink
    return None if link is None else (link.target, link.location, link.tooltip, link.display)


def inspect_workbook(path, template, count):
    require(path.suffix.lower() == ".xlsx", "Input must be an XLSX copy")
    with ZipFile(path) as archive:
        names = archive.namelist()
        require(len(names) == len(set(names)), "Duplicate workbook archive member")
        require(
            not any("vbaproject" in n.lower() or "externallinks/" in n.lower() for n in names),
            "Macros/external workbook links forbidden",
        )
        require(
            sum(i.file_size for i in archive.infolist()) < 20_000_000, "Workbook archive too large"
        )
    actual = load_workbook(path, data_only=False, keep_links=True)
    expected = load_workbook(template, data_only=False)
    require(actual.sheetnames == expected.sheetnames, "Sheets changed or missing")
    responses = {}
    for ws in actual:
        original = expected[ws.title]
        require(
            ws.max_row == original.max_row and ws.max_column == original.max_column,
            f"Cell dimensions changed: {ws.title}",
        )
        require(str(ws.merged_cells) == str(original.merged_cells), "Merged cells changed")
        require(ws.sheet_state == original.sheet_state, "Sheet visibility changed")
        for row in ws:
            for cell in row:
                old = original[cell.coordinate]
                require(cell.data_type != "f", f"Formula forbidden: {ws.title}!{cell.coordinate}")
                editable = 2 <= cell.row <= count + 1 and (
                    (ws.title == "Review" and cell.column == 5)
                    or (ws.title == "Source details" and cell.column == 9)
                )
                require(
                    hyperlink(cell) == hyperlink(old), f"Link changed: {ws.title}!{cell.coordinate}"
                )
                require(
                    cell.comment is None,
                    f"Comments are not response cells: {ws.title}!{cell.coordinate}",
                )
                if editable:
                    responses[(ws.title, cell.row)] = cell.value
                else:
                    require(
                        cell.value == old.value and cell.data_type == old.data_type,
                        f"Frozen source context changed: {ws.title}!{cell.coordinate}",
                    )
    actual.close()
    expected.close()
    return responses


def import_workbook(workbook, output, root=ROOT):
    root, workbook, output = Path(root).resolve(), Path(workbook), Path(output)
    require(not output.exists() and not output.is_symlink(), "Output must be a new directory")
    require(not output.resolve().is_relative_to(root), "Output must be outside the repository")
    require(not workbook.is_symlink(), "Input symlink is not allowed")
    review = root / BASE / "review-support"
    for name, expected in PINS.items():
        require(
            sha(review / name) == expected,
            f"Frozen baseline changed: {name}; requires separately reviewed importer revision",
        )
    payload = json.loads((review / "decisions.json").read_text())
    require(
        sha(review / "organization.json") == payload["organization_sha256"],
        "Decision organization baseline changed",
    )
    items = payload["items"]
    require(len({r["id"] for r in items}) == len(items), "Duplicate decision IDs")
    docs, raw, inputs = {}, {}, {}
    for item in items:
        path = safe_source(root, item["source_path"])
        md = safe_source(root, item["clause_path"].split("#")[0])
        require(
            sha(path) == item["source_sha256"] and sha(md) == item["markdown_sha256"],
            f"Stale source for {item['id']}",
        )
        relative = item["source_path"]
        raw[relative] = path.read_text()
        docs[relative] = json.loads(raw[relative])
        inputs[relative] = sha(path)
        inputs[str(md.relative_to(root))] = sha(md)
        match = re.fullmatch(
            r"/(proposed_terms|unresolved_fields)/(0|[1-9][0-9]*)", item["json_pointer"]
        )
        require(match is not None, "Invalid source pointer")
        group, index = match[1], int(match[2])
        record = docs[relative][group][index]
        require(record["id"] == item["id"], "Source pointer ID mismatch")
        require(record.get("term", record.get("field")) == item["text"], "Source text mismatch")
        require(
            record.get("basis", record.get("next_action")) == item["action_or_basis"],
            "Source basis mismatch",
        )
    cells = inspect_workbook(workbook, review / "decisions.xlsx", len(items))
    proposals = copy.deepcopy(docs)
    results, changes = [], []
    for row, item in enumerate(items, 2):
        a, b = cells["Review", row], cells["Source details", row]
        require(
            not (a is not None and str(a).strip() and b is not None and str(b).strip()),
            f"Two responses for {item['id']}; use one sheet only",
        )
        value = a if a is not None and str(a).strip() else b
        allowed = (
            {"term", "basis"}
            if item["json_pointer"].startswith("/proposed_terms/")
            else {"field", "next_action"}
        )
        disposition, replacements = response(value, allowed)
        result = {
            "id": item["id"],
            "response": value,
            "disposition": disposition,
            "source_path": item["source_path"],
            "record_pointer": item["json_pointer"],
        }
        results.append(result)
        _, group, index = item["json_pointer"].split("/")
        record = proposals[item["source_path"]][group][int(index)]
        for key, after in sorted(replacements.items()):
            before = record[key]
            if before != after:
                changes.append(
                    {
                        "id": item["id"],
                        "source_path": item["source_path"],
                        "source_sha256": item["source_sha256"],
                        "pointer": item["json_pointer"] + "/" + key,
                        "before": before,
                        "after": after,
                    }
                )
                record[key] = after
    # No writes occur until the entire workbook and every source have passed.
    report = {
        "status": STATUS,
        "baseline_revision": payload["baseline_revision"],
        "template_sha256": PINS["decisions.xlsx"],
        "input_workbook_sha256": sha(workbook),
        "source_hashes": inputs,
        "items": results,
        "changes": changes,
        "publications_regenerated": False,
        "source_writeback": False,
        "acceptance_created": False,
    }
    output.mkdir(parents=True, exist_ok=False)
    (output / "report.json").write_text(json.dumps(report, indent=2) + "\n")
    lines = [
        "# Decision import proposal",
        "",
        "**PROPOSAL ONLY — not approved or applied.**",
        "",
        f"Source items checked: {len(items)}. Proposed field replacements: {len(changes)}.",
        "",
        (
            "The original sources and all PDF, HTML and workbook editions remain unchanged. "
            "Proposed JSON files are review copies, not canon. Markdown clauses, financial "
            "implications and regenerated editions require separate reconciliation and "
            "exact-file review before adoption."
        ),
        "",
        (
            "See `report.json` for every response and source hash; `review.sqlite3` contains "
            "the same response and change records."
        ),
        "",
    ]
    for relative in sorted({c["source_path"] for c in changes}):
        before = raw[relative].splitlines(keepends=True)
        after_text = json.dumps(proposals[relative], indent=2, ensure_ascii=False) + "\n"
        target = output / "proposed" / relative
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(after_text)
        difference = "".join(
            difflib.unified_diff(
                before,
                after_text.splitlines(keepends=True),
                fromfile=relative,
                tofile="proposed/" + relative,
            )
        )
        (output / (Path(relative).stem + ".diff")).write_text(difference)
        lines += [
            f"## {relative}",
            "",
            (
                f"[Proposed JSON](proposed/{relative}) · "
                f"[Before/after diff]({Path(relative).stem}.diff)"
            ),
            "",
        ]
        for c in [x for x in changes if x["source_path"] == relative]:
            lines += [
                f"- `{c['id']}`: `{c['pointer']}` "
                "(exact before/after text in the diff and JSON report)."
            ]
        lines += [""]
    (output / "REPORT.md").write_text("\n".join(lines).rstrip() + "\n")
    with sqlite3.connect(output / "review.sqlite3") as db:
        db.execute("CREATE TABLE metadata (key TEXT PRIMARY KEY, value TEXT NOT NULL)")
        db.executemany(
            "INSERT INTO metadata VALUES (?,?)",
            [("status", STATUS), ("input_workbook_sha256", report["input_workbook_sha256"])],
        )
        db.execute(
            "CREATE TABLE response (id TEXT PRIMARY KEY, "
            "disposition TEXT NOT NULL, record_json TEXT NOT NULL)"
        )
        db.executemany(
            "INSERT INTO response VALUES (?,?,?)",
            [(r["id"], r["disposition"], json.dumps(r, sort_keys=True)) for r in results],
        )
        db.execute(
            "CREATE TABLE proposed_change (source_path TEXT, pointer TEXT, "
            "record_json TEXT, PRIMARY KEY(source_path,pointer))"
        )
        db.executemany(
            "INSERT INTO proposed_change VALUES (?,?,?)",
            [(r["source_path"], r["pointer"], json.dumps(r, sort_keys=True)) for r in changes],
        )
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    try:
        report = import_workbook(args.workbook, args.output)
    except (ValueError, KeyError, IndexError, OSError, BadZipFile) as error:
        parser.exit(1, f"FAIL: {error}\n")
    print(
        f"PASS: {len(report['items'])} source items; "
        f"{len(report['changes'])} proposed replacements; no source writes or approval"
    )


if __name__ == "__main__":
    main()
