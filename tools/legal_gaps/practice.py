"""Build and check four public worked examples from accepted repository evidence."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import sqlite3
import subprocess
import tempfile
from collections import defaultdict
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import xlsxwriter
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
DEST = Path("docs/legal/gap-instruments/practice")
CLASSIFICATION = "PUBLIC_WORKED_EXAMPLE"
PREPARED = "2026-09-12"
SOURCE_REVISION = "659a56747fe76522d18645ff115888c13fa8d2b0"


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_json(root, path):
    return json.loads((root / path).read_text())


def read_csv(root, path):
    with (root / path).open() as stream:
        return list(csv.DictReader(stream))


def write_json(path, value):
    path.write_text(json.dumps(value, indent=2, ensure_ascii=False) + "\n")


def decimal(value):
    return Decimal(str(value))


def make_packets(root=ROOT):
    finance_path = "industrial/source/finance.json"
    closing_path = "industrial/transaction/05_ARU_CLOSING_AND_TAX_DELIVERY.md"
    accounting_path = "industrial/finance/TRANSACTION_ACCOUNTING.md"
    bridge_path = "docs/finance/evidence/tax-transaction/CURRENT_SOURCE_BRIDGE.json"
    tb_path = "docs/finance/evidence/tax-transaction/aru_acquisition_opening_trial_balance.csv"
    debt_path = "docs/finance/evidence/tax-transaction/aru_2026_debt.csv"
    invoice_path = "docs/finance/evidence/SH-FIN-HUMAN-001/source.json"
    agreement_path = "industrial/transaction/08_INTERCOMPANY_LOGISTICS_AGREEMENT.md"
    finance = read_json(root, finance_path)
    tx = finance["transaction"]
    bridge = read_json(root, bridge_path)
    ppa = bridge["ppa"]
    tb = read_csv(root, tb_path)
    debt = read_csv(root, debt_path)
    invoice = read_json(root, invoice_path)
    packets = []

    def packet(slug, title, period, scenario, revision, route, sources, task, limits):
        p = dict(
            id="SH-LEGAL-PRACTICE-" + slug.upper(),
            slug=slug,
            title=title,
            classification=CLASSIFICATION,
            design_status="DRAFT_FOR_REVIEW",
            prepared=PREPARED,
            period=period,
            scenario=scenario,
            source_revision=revision,
            route=route,
            evidence=[dict(path=path, sha256=sha(root / path)) for path in sources],
            task=task,
            limits=limits,
            inputs=[],
            calculations=[],
            findings=[],
            selected_rows={},
        )
        packets.append(p)
        return p

    def inp(p, key, label, amount, source, locator):
        p["inputs"].append(
            dict(
                id=key,
                label=label,
                amount_usd=float(decimal(amount)),
                source_path=source,
                locator=locator,
            )
        )
        return key

    def calc(p, key, label, terms, note):
        inputs = {r["id"]: decimal(r["amount_usd"]) for r in p["inputs"]}
        amount = sum((decimal(sign) * inputs[ref] for sign, ref in terms), Decimal(0))
        p["calculations"].append(
            dict(
                id=key,
                label=label,
                terms=[list(t) for t in terms],
                expected_usd=float(amount),
                explanation=note,
            )
        )

    def finding(p, key, question, result, evidence, request):
        p["findings"].append(
            dict(
                id=key,
                question=question,
                worked_finding=result,
                evidence=evidence,
                additional_evidence=request,
            )
        )

    acq = packet(
        "acquisition",
        "ARU acquisition accounting",
        "January 7, 2026 acquisition opening",
        None,
        bridge["base_revision"],
        "docs/reader/exercises/ACQUISITION.md",
        [
            finance_path,
            closing_path,
            accounting_path,
            bridge_path,
            tb_path,
            "industrial/publications/SH-IND-ARU-CLOSE-001_v1.0.0.pdf",
        ],
        [
            "Reconcile buyer uses before fees to new debt and parent equity.",
            "Separate parent cash including fees from stock consideration and retained finance leases.",
            "Recompute book goodwill and independent tax goodwill; explain the difference.",
            "Sum all 16 opening accounts and record the debit-credit difference.",
            "Write an evidence conclusion distinguishing model reconciliation from signed closing and filing evidence.",
        ],
        [
            "Selected industrial reconstruction; do not add it to the separate operations release.",
            "PPA and tax basis retain conditional model assumptions; no valuation or tax opinion is supplied.",
            "An opening group trial balance is not independently reconstructed statutory books for each subsidiary.",
        ],
    )
    for key, label, field in [
        ("price", "Buyer stock consideration", "buyer_consideration"),
        ("payoff", "Old term and revolver payoff", "existing_term_revolver_refinance"),
        ("debt", "New term debt", "new_debt"),
        ("equity", "Parent equity before fees", "parent_equity_before_fees"),
        ("transaction_fee", "Transaction services expense", "transaction_expense"),
        ("debt_fee", "Debt issuance cost", "debt_issuance_cost"),
    ]:
        inp(acq, key, label, tx[field], finance_path, "transaction." + field)
    inp(
        acq,
        "net_assets",
        "Fair identifiable book net assets",
        ppa["identifiable_net_assets_before_refinancing_usd"],
        bridge_path,
        "ppa.identifiable_net_assets_before_refinancing_usd",
    )
    inp(
        acq,
        "tax_agub",
        "Modeled adjusted grossed-up tax basis",
        ppa["tax_allocation"]["modeled_agub_usd"],
        bridge_path,
        "ppa.tax_allocation.modeled_agub_usd",
    )
    inp(
        acq,
        "tax_other",
        "Other tax asset bases (sum)",
        sum(ppa["tax_allocation"]["other_tax_asset_bases_usd"].values()),
        bridge_path,
        "sum ppa.tax_allocation.other_tax_asset_bases_usd",
    )
    for r in tb:
        for side in ["debit", "credit"]:
            inp(
                acq,
                f"tb_{r['account']}_{side}",
                f"{r['account']} {r['account_name']} — {side}",
                r[side + "_balance_usd"],
                tb_path,
                f"entity={r['entity']}; year={r['year']}; account={r['account']}; {side}_balance_usd",
            )
    for args in [
        (
            "uses",
            "Closing uses before fees",
            [(1, "price"), (1, "payoff")],
            "Stock consideration includes escrow; retained leases are not paid off again.",
        ),
        (
            "sources",
            "Closing sources before fees",
            [(1, "debt"), (1, "equity")],
            "New term debt and parent equity finance the same uses.",
        ),
        (
            "difference",
            "Funding reconciliation difference",
            [(1, "price"), (1, "payoff"), (-1, "debt"), (-1, "equity")],
            "Zero means this selected funding bridge reconciles, not that cash settled.",
        ),
        (
            "parent_cash",
            "Parent cash including fees",
            [(1, "equity"), (1, "transaction_fee"), (1, "debt_fee")],
            "Fees are separate from stock consideration.",
        ),
        (
            "book_goodwill",
            "Book goodwill residual",
            [(1, "price"), (-1, "net_assets")],
            "Source-model residual; no invented operating-income plug.",
        ),
        (
            "tax_goodwill",
            "Tax goodwill residual",
            [(1, "tax_agub"), (-1, "tax_other")],
            "Independent conditional tax allocation, not book goodwill copied into tax.",
        ),
    ]:
        calc(acq, *args)
    calc(
        acq,
        "tb_debits",
        "Opening trial-balance debit total",
        [(1, "tb_" + r["account"] + "_debit") for r in tb],
        "All 16 source accounts included.",
    )
    calc(
        acq,
        "tb_credits",
        "Opening trial-balance credit total",
        [(1, "tb_" + r["account"] + "_credit") for r in tb],
        "All 16 source accounts included.",
    )
    calc(
        acq,
        "tb_difference",
        "Opening trial-balance difference",
        [
            (sign, "tb_" + r["account"] + "_" + side)
            for r in tb
            for sign, side in [(1, "debit"), (-1, "credit")]
        ],
        "Balanced group opening books do not certify external evidence.",
    )
    finding(
        acq,
        "ACQ-F01",
        "What does ARU-CL-07 establish?",
        "The March 2 source records filing-ready review, not IRS submission or acceptance.",
        closing_path + "#delivery-register",
        "Actual submission/status evidence and qualified eligibility review.",
    )
    finding(
        acq,
        "ACQ-F02",
        "Why do the goodwill amounts differ?",
        "Book goodwill uses fair identifiable net assets including the conditional reserve DTA; tax goodwill uses separately modeled recognized liabilities and other tax bases.",
        accounting_path + "#purchase-price-allocation",
        "Independent valuation and support for conditional tax classifications before external reliance.",
    )
    acq["selected_rows"]["opening_trial_balance"] = tb

    dp = packet(
        "debt",
        "ARU debt reconciliation",
        "January–December 2026 selected industrial debt model",
        None,
        bridge["base_revision"],
        "docs/reader/exercises/ACQUISITION.md",
        [finance_path, accounting_path, closing_path, debt_path, bridge_path],
        [
            "Reconcile the opening term advance to old debt payoff and upstream acquisition financing.",
            "Reperform every monthly term principal rollforward and the full-year principal change.",
            "Separate year-end funded principal from term carrying amount and from interest expense.",
            "State which assertions require lender, payoff and lien-release evidence.",
        ],
        [
            "ACT/365 and the 6.75% rate are source model conventions, not a verified lender agreement.",
            "The twelve monthly rows are conditional modeled schedules; no bank settlement is proved.",
            "Do not add a revolver limit to drawn principal or expense debt issuance costs twice.",
        ],
    )
    for key, label, value in [
        ("advance", "Acquisition term advance", tx["new_debt"]),
        ("payoff", "Old term/revolver refinance", tx["existing_term_revolver_refinance"]),
        (
            "upstream",
            "Upstream acquisition financing",
            ppa["new_debt_upstream_acquisition_distribution_usd"],
        ),
    ]:
        inp(
            dp,
            key,
            label,
            value,
            finance_path if key != "upstream" else bridge_path,
            "transaction funding bridge / " + key,
        )
    calc(
        dp,
        "advance_use",
        "Advance allocation difference",
        [(1, "advance"), (-1, "payoff"), (-1, "upstream")],
        "New borrowing is used once; upstream funds are financing, not revenue.",
    )
    for r in debt:
        month = r["month"]
        for field, label in [
            ("opening_term_usd", "Opening term principal"),
            ("term_principal_usd", "Term principal reduction"),
            ("closing_term_usd", "Closing term principal"),
        ]:
            inp(
                dp,
                month + "_" + field,
                "Month " + month + ": " + label,
                r[field],
                debt_path,
                "month=" + month + "; " + field,
            )
        calc(
            dp,
            "monthly_" + month,
            "Month " + month + " principal difference",
            [
                (1, month + "_opening_term_usd"),
                (-1, month + "_term_principal_usd"),
                (-1, month + "_closing_term_usd"),
            ],
            "Opening less principal reduction equals closing; interest is not principal.",
        )
    for field, label in [
        ("closing_lease_usd", "December retained lease principal"),
        ("deferred_debt_issue_cost_usd", "December unamortized issuance cost"),
        ("revolver_draw_usd", "December revolver draw"),
    ]:
        inp(dp, field, label, debt[-1][field], debt_path, "month=12; " + field)
    calc(
        dp,
        "annual_principal",
        "Full-year term principal reductions",
        [(1, r["month"] + "_term_principal_usd") for r in debt],
        "Three modeled quarterly reductions, not twelve installments.",
    )
    calc(
        dp,
        "closing_funded",
        "December funded principal including leases",
        [(1, "12_closing_term_usd"), (1, "closing_lease_usd"), (1, "revolver_draw_usd")],
        "Principal amount; do not net unamortized issuance costs here.",
    )
    calc(
        dp,
        "term_carrying",
        "December term debt carrying amount",
        [(1, "12_closing_term_usd"), (-1, "deferred_debt_issue_cost_usd")],
        "Term principal less unamortized term-debt issuance cost; excludes leases.",
    )
    finding(
        dp,
        "DEBT-F01",
        "Does a zero model payoff balance prove liens were released?",
        "No. The source records internal payoff instructions and modeled refinancing; it supplies no independent executed release.",
        closing_path + "#delivery-register",
        "Lender payoff statement, verified settlement and item-specific release evidence.",
    )
    dp["selected_rows"]["debt_schedule"] = debt

    rev = packet(
        "revenue-dispute",
        "Foundry Field revenue and credit dispute",
        "Base scenario, January–October 2027; FF-003 term 0",
        "base",
        invoice["release_source_revision"],
        "docs/reader/exercises/INVOICE.md",
        [
            invoice_path,
            "docs/finance/evidence/SH-FIN-HUMAN-001/PACKET.md",
            "docs/finance/evidence/SH-FIN-HUMAN-001/reconciliation.xlsx",
            "docs/finance/evidence/SH-FIN-HUMAN-001/ACCEPTANCE.json",
        ],
        [
            "Evaluate the review assertion: zero ledger receivable means the invoice was fully paid. This is an exercise prompt, not a newly attributed customer statement.",
            "Reconcile billed amount, receipt, writeoff, two credits and recovery using the five movement IDs.",
            "Explain why the July and August credits affect different debit accounts and do not create a cash refund.",
            "Trace all six included event groups to balanced journal lines; distinguish this subset from a full revenue-recognition population.",
        ],
        [
            "The accepted packet is reconstructed evidence, not an original invoice, signed SLA or bank confirmation.",
            "No revenue dispute outcome, customer consent or proposed billing identity is invented.",
            "Contract-wide recognition and pooled allowance entries are excluded; this is not a complete customer P&L.",
        ],
    )
    native_invoice = invoice["rows"]["invoices"][0]
    inp(
        rev,
        "billed",
        "Original invoice",
        native_invoice["amount_usd"],
        invoice_path,
        "rows.invoices; invoice_id=INV-base-FF-003-TERM-0; amount_usd",
    )
    movements = invoice["rows"]["credit_history"]
    for n, r in enumerate(movements):
        inp(
            rev,
            "movement_" + str(n),
            r["period"] + " " + r["action"],
            r["amount_usd"],
            invoice_path,
            r["source_id"],
        )
    calc(
        rev,
        "ar",
        "Closing ledger receivable",
        [(1, "billed"), (-1, "movement_0"), (-1, "movement_1")],
        "Receivable is removed by writeoff, not full payment.",
    )
    calc(
        rev,
        "collected",
        "Cash collections including recovery",
        [(1, "movement_0"), (1, "movement_4")],
        "Recovery is already included in the released collected total.",
    )
    calc(
        rev,
        "credits",
        "Credits against written-off claim",
        [(1, "movement_2"), (1, "movement_3")],
        "Credits reduce the claim, with no refund in this selected packet.",
    )
    calc(
        rev,
        "claim",
        "Surviving written-off claim",
        [(1, "movement_1"), (-1, "movement_2"), (-1, "movement_3"), (-1, "movement_4")],
        "Derived claim measure; not an asset newly recognized here.",
    )
    grouped = defaultdict(list)
    for r in invoice["rows"]["journal"]:
        grouped[r["source_id"]].append(r)
    for n, (event, lines) in enumerate(grouped.items()):
        for side in ["debit", "credit"]:
            inp(
                rev,
                f"journal_{n}_{side}",
                event + ": " + side,
                sum(decimal(r[side + "_usd"]) for r in lines),
                invoice_path,
                "rows.journal; source_id=" + event,
            )
        calc(
            rev,
            "journal_" + str(n),
            "Event " + str(n + 1) + " journal difference",
            [(1, f"journal_{n}_debit"), (-1, f"journal_{n}_credit")],
            event + "; two selected journal lines balance.",
        )
    for n, credit in enumerate(invoice["rows"]["credit_notes"]):
        finding(
            rev,
            f"REV-F{n + 1:02}",
            credit["reason"] + ": what is supported?",
            "The source credit debits "
            + credit["debit_account"]
            + ", applies to a written-off claim and has no refund amount in this selected record.",
            credit["source_id"],
            "Underlying signed SLA or amendment and customer acceptance before asserting contractual validity.",
        )
    finding(
        rev,
        "REV-F03",
        "Was the invoice paid in full?",
        "No such conclusion follows: modeled receipts plus recovery are less than billing, and writeoff explains zero ledger AR.",
        "INV-base-FF-003-TERM-0",
        "Independent collection confirmation would still be required to establish actual cash.",
    )
    rev["selected_rows"] = invoice["rows"]

    due = packet(
        "legal-due-diligence",
        "Taylor–Red Wash legal due diligence",
        "Agreement source effective July 7, 2026; publication cutoff September 5, 2026",
        None,
        SOURCE_REVISION,
        "docs/reader/exercises/CONTRACTS.md",
        [
            agreement_path,
            "industrial/publications/SH-IND-IC-001_v1.0.0.pdf",
            "industrial/source/entities.json",
        ],
        [
            "Prepare a clause-based due-diligence schedule for the stated Taylor–Red Wash agreement.",
            "Separate stated contractual terms, performance claims and independent execution evidence.",
            "Reperform one source rate calculation without inventing a shipment or invoice.",
            "For each issue below, identify the relevant source heading, supported conclusion, missing evidence and appropriate next request.",
        ],
        [
            "The accepted source is a reconstructed internal agreement summary, not an authentic executed agreement.",
            "An absent document in this selected packet does not prove it is absent everywhere or that a party breached.",
            "No legal enforceability opinion, shipment, payment deadline breach or uranium qualification is asserted.",
        ],
    )
    # Values are extracted from the accepted table; the task does not invent a delivery.
    text = (root / agreement_path).read_text()
    acid = next(line for line in text.splitlines() if line.startswith("| Sulfuric acid |"))
    cells = [x.strip() for x in acid.strip("|").split("|")]
    vals = [Decimal(x.replace("$", "").replace(",", "")) for x in cells[1:]]
    for key, label, value in [
        ("rail", "Acid rail rate per car", vals[0]),
        ("terminal", "Acid terminal rate per car", vals[1]),
        ("truck_total", "Four truck trips at stated rate (extended)", vals[2] * vals[3]),
        ("combined", "Published source combined rate per car", vals[4]),
    ]:
        inp(due, key, label, value, agreement_path, "Schedule A; Sulfuric acid row")
    calc(
        due,
        "rate",
        "Combined rate per car (not an invoice)",
        [(1, "rail"), (1, "terminal"), (1, "truck_total")],
        "Four road trips are priced, not one; no physical quantity ordered.",
    )
    calc(
        due,
        "rate_difference",
        "Difference from source combined rate",
        [(1, "rail"), (1, "terminal"), (1, "truck_total"), (-1, "combined")],
        "Arithmetic agreement does not establish delivered service.",
    )
    for key, question, conclusion, heading, request in [
        (
            "LEGAL-F01",
            "Who performs and bills?",
            "ARU supplies terminal/trucking; BS&T supplies rail; RWH is customer. Billing agency does not merge the sellers.",
            "Parties, authority and term",
            "Verified signatures and invoice naming the actual performing seller.",
        ),
        (
            "LEGAL-F02",
            "Does planned capacity guarantee revenue?",
            "No minimum-volume guarantee or take-or-pay covenant is stated; unused slots produce no invoice.",
            "Schedule A — service and rates",
            "Actual quantities, source events and accepted rate version for a selected invoice.",
        ),
        (
            "LEGAL-F03",
            "Can you test a late-service breach?",
            "The source states 72 elapsed hours, 96 for steel, with a separate 36–48-hour operating goal. No selected shipment timestamps are supplied here.",
            "Schedule B — clock, free time and service remedies",
            "Interchange/empty-release timestamps, exclusions and responsibility evidence.",
        ),
        (
            "LEGAL-F04",
            "Is storage automatically chargeable?",
            "Only supported customer-caused occupancy after free time may receive the stated prorated charge; baseline has no chargeable dwell.",
            "Schedule B — clock, free time and service remedies",
            "Car, custody, start/end time, expiry, cause and charging entity; any third-party invoice.",
        ),
        (
            "LEGAL-F05",
            "Is uranium custody authorized?",
            "Finished uranium concentrate is excluded and direct product custody remains OPEN_GATED.",
            "Schedule C — custody, insurance and qualification",
            "Separate agreement and provider-specific qualification before any future activation.",
        ),
        (
            "LEGAL-F06",
            "Are insurance policies evidenced?",
            "The source states synthetic coverage assumptions, not insurer names, policies or certificates.",
            "Schedule C — custody, insurance and qualification",
            "Applicable policy/certificate and exclusions; no inference from a budget.",
        ),
        (
            "LEGAL-F07",
            "Who owns the Phase 1 assets?",
            "RWH owns the mine programme; ARU owns reusable Taylor infrastructure. Ownership and depreciation remain separate.",
            "Schedule D — assets, billing and termination",
            "Asset register and funding/acceptance evidence for any installed asset assertion.",
        ),
        (
            "LEGAL-F08",
            "Does next-month settlement prove timely payment?",
            "No. The source due period is 30 calendar days, while the model uses following-month settlement; exact invoice and payment dates are needed.",
            "Schedule D — assets, billing and termination",
            "Dated invoice and independent settlement reference for the selected obligation.",
        ),
    ]:
        finding(due, key, question, conclusion, agreement_path + " — " + heading, request)
    return packets


def formula(calculation, inputs):
    positions = {r["id"]: n + 4 for n, r in enumerate(inputs)}
    return "=" + "".join(
        ("+" if sign > 0 and n else "-" if sign < 0 else "") + f"'Source inputs'!B{positions[key]}"
        for n, (sign, key) in enumerate(calculation["terms"])
    )


def workbook(packet, output, worked):
    wb = xlsxwriter.Workbook(output)
    wb.set_properties(
        {
            "title": packet["title"],
            "created": datetime(2026, 9, 12),
            "comments": CLASSIFICATION + "; draft design",
        }
    )
    title = wb.add_format(
        {
            "bold": True,
            "font_color": "white",
            "bg_color": "#243238",
            "text_wrap": True,
            "valign": "top",
            "font_size": 12,
        }
    )
    body = wb.add_format({"text_wrap": True, "valign": "top", "font_size": 11, "indent": 1})
    num = wb.add_format(
        {
            "num_format": "#,##0.00;[Red](#,##0.00);0.00",
            "valign": "top",
            "font_size": 11,
            "indent": 1,
        }
    )
    entry = wb.add_format(
        {
            "bg_color": "#FFF2CC",
            "border": 1,
            "border_color": "#D7D1C3",
            "text_wrap": True,
            "valign": "top",
        }
    )

    def sheet(name, heading, headers, widths):
        ws = wb.add_worksheet(name)
        ws.hide_gridlines(2)
        ws.merge_range(0, 0, 0, len(headers) - 1, heading, title)
        ws.set_row(0, 34)
        ws.merge_range(
            1,
            0,
            1,
            len(headers) - 1,
            ("PUBLIC WORKED EXAMPLE" if worked else "PUBLIC PRACTICE — BLANK RESPONSES")
            + " | Draft design for review",
            body,
        )
        ws.set_row(1, 28)
        for c, (h, width) in enumerate(zip(headers, widths, strict=True)):
            ws.write(2, c, h, title)
            ws.set_column(c, c, width)
        ws.set_row(2, 32)
        ws.freeze_panes(3, 0)
        ws.set_landscape()
        ws.set_paper(8)
        ws.fit_to_pages(1, 0)
        ws.repeat_rows(0, 2)
        ws.set_margins(0.3, 0.3, 0.4, 0.6)
        ws.set_footer("&C" + packet["id"] + " | &P of &N", {"margin": 0.2})
        return ws

    def row(ws, row_index, values, formats=None):
        max_lines = 1
        for c, value in enumerate(values):
            fmt = formats[c] if formats else body
            ws.write(row_index, c, value, fmt)
            width = 40 if c != 0 else 34
            max_lines = max(max_lines, len(str(value or "")) // width + 1)
        ws.set_row(row_index, max(30, max_lines * 15 + 8))

    ws = sheet("Read first", packet["title"], ["Field", "Instruction"], [28, 95])
    info = [
        (
            "Classification",
            CLASSIFICATION + "; public learning material, not a private assessment key.",
        ),
        ("Scope", packet["period"]),
        (
            "Scenario",
            packet["scenario"] or "Selected source reconstruction; no operations scenario combined",
        ),
        ("Source revision", packet["source_revision"]),
        (
            "Mode",
            "Worked calculations and findings supplied."
            if worked
            else "Yellow response cells are intentionally blank. Source inputs remain supplied facts.",
        ),
        (
            "Method",
            "Read Evidence, then Source inputs. Complete Working paper and Findings. Compare the separate worked example only when ready.",
        ),
        ("Limitations", " ".join(packet["limits"])),
        (
            "Approval",
            "Source-derived public answers are authorized; new workbook designs remain held for exact-file review.",
        ),
    ]
    for n, values in enumerate(info, 3):
        row(ws, n, values)
    ws.print_area(0, 0, 2 + len(info), 1)
    ws = sheet(
        "Source inputs",
        "Selected source inputs — USD unless stated otherwise",
        ["Measure", "Amount USD", "Source and native locator"],
        [43, 21, 89],
    )
    for n, r in enumerate(packet["inputs"], 3):
        row(
            ws,
            n,
            [r["label"], r["amount_usd"], r["source_path"] + "\n" + r["locator"]],
            [body, num, body],
        )
    ws.print_area(0, 0, 2 + len(packet["inputs"]), 2)
    ws = sheet(
        "Working paper",
        "Reperform the selected calculations",
        ["Calculation", "Result USD", "Basis / explanation"],
        [44, 22, 85],
    )
    for n, c in enumerate(packet["calculations"], 3):
        row(
            ws,
            n,
            [
                c["label"],
                None,
                c["explanation"] if worked else "Enter calculation and explain its evidence basis.",
            ],
            [body, num if worked else entry, body if worked else entry],
        )
        if worked:
            ws.write_formula(n, 1, formula(c, packet["inputs"]), num, c["expected_usd"])
    ws.print_area(0, 0, 2 + len(packet["calculations"]), 2)
    ws = sheet(
        "Findings",
        "Evidence conclusions and further requests",
        ["Issue / question", "Supported finding", "Source / native reference", "Further evidence"],
        [32, 49, 47, 48],
    )
    for n, f in enumerate(packet["findings"], 3):
        row(
            ws,
            n,
            [
                f["id"] + ": " + f["question"],
                f["worked_finding"] if worked else None,
                f["evidence"],
                f["additional_evidence"] if worked else None,
            ],
            [body, body if worked else entry, body, body if worked else entry],
        )
    ws.print_area(0, 0, 2 + len(packet["findings"]), 3)
    ws = sheet(
        "Evidence",
        "Exact evidence selected for this example",
        ["Repository path", "SHA-256"],
        [88, 72],
    )
    for n, e in enumerate(packet["evidence"], 3):
        row(ws, n, [e["path"], e["sha256"]])
    ws.print_area(0, 0, 2 + len(packet["evidence"]), 1)
    wb.close()


def link(base, target):
    return os.path.relpath(target, base).replace(os.sep, "/")


def markdown(packet, directory, root):
    evidence = "\n".join(
        f"- [{e['path']}]({link(directory, root / e['path'])}) — SHA-256 `{e['sha256']}`"
        for e in packet["evidence"]
    )
    head = f"# {packet['title']}\n\n**{packet['id']} · PUBLIC_WORKED_EXAMPLE · prepared {PREPARED}**\n\nScope: {packet['period']}. Source revision: `{packet['source_revision']}`. These are public learning materials with source-derived worked answers, not private assessment keys. New workbook designs remain draft for exact-file review.\n\n"
    task = (
        head + "## Your task\n\n" + "\n".join(f"{n}. {s}" for n, s in enumerate(packet["task"], 1))
    )
    task += "\n\nOpen [the blank working paper](blank.xlsx). The source inputs and evidence are supplied; yellow answer cells are blank. Save a separate copy, record your calculation and distinguish supported findings from further evidence requests. Then compare with the separate [worked explanation](WORKED.md) and [worked workbook](worked.xlsx).\n\n"
    task += (
        f"Earlier route: [{packet['route']}]({link(directory, root / packet['route'])}).\n\n## Exact selected evidence\n\n"
        + evidence
    )
    task += "\n\n## Scope limits\n\n" + "\n".join("- " + s for s in packet["limits"]) + "\n"
    (directory / "TASK.md").write_text(task)
    worked = (
        head + "## Worked calculations\n\n| Calculation | USD | Explanation |\n|---|---:|---|\n"
    )
    worked += "\n".join(
        f"| {c['label']} | {c['expected_usd']:,.2f} | {c['explanation']} |"
        for c in packet["calculations"]
    )
    worked += "\n\nThe [worked workbook](worked.xlsx) contains actual Excel formulas referencing the supplied Source inputs sheet, with cached results for readers that do not recalculate. Source data and derived totals are distinct. Do not sum the rows as one financial total.\n\n## Supported findings and missing evidence\n\n"
    for f in packet["findings"]:
        worked += f"### {f['id']} — {f['question']}\n\n{f['worked_finding']}\n\nSource: `{f['evidence']}`. Further evidence: {f['additional_evidence']}\n\n"
    worked += (
        "## Evidence and limitations\n\n"
        + evidence
        + "\n\n"
        + "\n".join("- " + s for s in packet["limits"])
        + "\n\n[Return to the task](TASK.md). Public worked conclusions do not establish execution, payment, audit assurance or legal enforceability.\n"
    )
    (directory / "WORKED.md").write_text(worked)


def build(root=ROOT):
    target = root / DEST
    target.mkdir(parents=True, exist_ok=True)
    packets = make_packets(root)
    artifacts = []
    dbpath = target / "practice.sqlite3"
    if dbpath.exists():
        dbpath.unlink()
    db = sqlite3.connect(dbpath)
    db.executescript(
        "CREATE TABLE packet(id TEXT PRIMARY KEY, slug TEXT UNIQUE, record_json TEXT NOT NULL); CREATE TABLE evidence(packet_id TEXT,path TEXT,sha256 TEXT,PRIMARY KEY(packet_id,path)); CREATE TABLE calculation(packet_id TEXT,id TEXT,expected_usd REAL,record_json TEXT,PRIMARY KEY(packet_id,id)); CREATE TABLE finding(packet_id TEXT,id TEXT,record_json TEXT,PRIMARY KEY(packet_id,id));"
    )
    for p in packets:
        directory = target / p["slug"]
        directory.mkdir(exist_ok=True)
        write_json(directory / "packet.json", p)
        markdown(p, directory, root)
        workbook(p, directory / "blank.xlsx", False)
        workbook(p, directory / "worked.xlsx", True)
        db.execute(
            "INSERT INTO packet VALUES(?,?,?)", (p["id"], p["slug"], json.dumps(p, sort_keys=True))
        )
        db.executemany(
            "INSERT INTO evidence VALUES(?,?,?)",
            [(p["id"], e["path"], e["sha256"]) for e in p["evidence"]],
        )
        db.executemany(
            "INSERT INTO calculation VALUES(?,?,?,?)",
            [
                (p["id"], c["id"], c["expected_usd"], json.dumps(c, sort_keys=True))
                for c in p["calculations"]
            ],
        )
        db.executemany(
            "INSERT INTO finding VALUES(?,?,?)",
            [(p["id"], f["id"], json.dumps(f, sort_keys=True)) for f in p["findings"]],
        )
        for name in ["TASK.md", "WORKED.md", "packet.json", "blank.xlsx", "worked.xlsx"]:
            path = directory / name
            artifacts.append(dict(path=str(path.relative_to(root)), sha256=sha(path)))
    db.commit()
    db.execute("VACUUM")
    db.close()
    readme = "# Public worked legal and accounting practice\n\nFour source-backed learning packets. Start with a task and its blank workbook; open the worked explanation separately. These public examples are expressly shared learning material, not private assessment keys. All amounts retain their source scenario and period. The new workbook designs remain draft for exact-file review.\n\n| Packet | Start | Separate worked example |\n|---|---|---|\n"
    readme += "\n".join(
        f"| {p['title']} | [Task]({p['slug']}/TASK.md) · [Blank workbook]({p['slug']}/blank.xlsx) | [Explanation]({p['slug']}/WORKED.md) · [Worked workbook]({p['slug']}/worked.xlsx) |"
        for p in packets
    )
    readme += "\n\n[Structured practice database](practice.sqlite3) stores complete packet, evidence, calculation and finding records. Each packet JSON retains native source identifiers and selected rows. [Build manifest](manifest.json) records source and derivative hashes.\n\nBuild: `python tools/legal_gaps/practice.py build`. Validate: `python tools/legal_gaps/practice.py validate`. Optional native workbook review renders: `python tools/legal_gaps/practice.py render`.\n"
    (target / "README.md").write_text(readme)
    for path in [dbpath, target / "README.md"]:
        artifacts.append(dict(path=str(path.relative_to(root)), sha256=sha(path)))
    manifest = dict(
        classification=CLASSIFICATION,
        design_status="DRAFT_FOR_REVIEW",
        prepared=PREPARED,
        packets=[p["id"] for p in packets],
        workbooks=8,
        worksheets=40,
        artifacts=artifacts,
        generator_sha256=sha(root / "tools/legal_gaps/practice.py"),
    )
    write_json(target / "manifest.json", manifest)
    print(
        "Built 4 public packets, 8 workbooks / 40 sheets, 8 Markdown task/worked documents and SQLite"
    )


def validate(root=ROOT):
    errors = []
    target = root / DEST
    try:
        manifest = read_json(root, DEST / "manifest.json")
        if manifest["generator_sha256"] != sha(root / "tools/legal_gaps/practice.py"):
            errors.append("generator drift")
        for a in manifest["artifacts"]:
            if sha(root / a["path"]) != a["sha256"]:
                errors.append("artifact drift: " + a["path"])
        db = sqlite3.connect(target / "practice.sqlite3")
        expected = make_packets(root)
        for p in expected:
            actual = read_json(root, DEST / p["slug"] / "packet.json")
            if actual != p:
                errors.append(p["slug"] + ": source/packet drift")
            row = db.execute("SELECT record_json FROM packet WHERE id=?", (p["id"],)).fetchone()
            if row is None or json.loads(row[0]) != p:
                errors.append(p["slug"] + ": SQLite packet drift")
            for table, key, fields in [
                ("evidence", "path", [(e["path"], e["sha256"]) for e in p["evidence"]]),
                (
                    "calculation",
                    "id",
                    [
                        (c["id"], c["expected_usd"], json.dumps(c, sort_keys=True))
                        for c in p["calculations"]
                    ],
                ),
                (
                    "finding",
                    "id",
                    [(f["id"], json.dumps(f, sort_keys=True)) for f in p["findings"]],
                ),
            ]:
                columns = (
                    "path,sha256"
                    if table == "evidence"
                    else "id,expected_usd,record_json"
                    if table == "calculation"
                    else "id,record_json"
                )
                rows = db.execute(
                    f"SELECT {columns} FROM {table} WHERE packet_id=? ORDER BY {key}", (p["id"],)
                ).fetchall()
                if rows != sorted(fields):
                    errors.append(p["slug"] + ": SQLite " + table + " drift")
            for mode in ["blank", "worked"]:
                path = target / p["slug"] / (mode + ".xlsx")
                wb = load_workbook(path, data_only=False)
                cached = load_workbook(path, data_only=True)
                if wb.sheetnames != [
                    "Read first",
                    "Source inputs",
                    "Working paper",
                    "Findings",
                    "Evidence",
                ]:
                    errors.append(p["slug"] + ": sheets differ")
                for n, item in enumerate(p["inputs"], 4):
                    if wb["Source inputs"].cell(n, 2).value != item["amount_usd"]:
                        errors.append(p["slug"] + ": workbook source input drift")
                for n, c in enumerate(p["calculations"], 4):
                    v = wb["Working paper"].cell(n, 2).value
                    if mode == "blank":
                        if v is not None:
                            errors.append(p["slug"] + ": blank answer is populated")
                    else:
                        if v != formula(c, p["inputs"]):
                            errors.append(p["slug"] + ": formula differs")
                        if cached["Working paper"].cell(n, 2).value != c["expected_usd"]:
                            errors.append(p["slug"] + ": cached result differs")
                for n, f in enumerate(p["findings"], 4):
                    for column, key in [(2, "worked_finding"), (4, "additional_evidence")]:
                        if wb["Findings"].cell(n, column).value != (
                            f[key] if mode == "worked" else None
                        ):
                            errors.append(p["slug"] + ": findings mode differs")
                wb.close()
                cached.close()
        db.close()
    except (OSError, KeyError, ValueError, sqlite3.Error) as exc:
        errors.append(str(exc))
    if errors:
        print("\n".join("FAIL: " + e for e in errors))
    else:
        print("PASS: four packets, source hashes, database, blank cells and worked formulas/caches")
    return errors


def render(root=ROOT):
    import fitz
    from PIL import Image, ImageDraw

    target = root / DEST
    qa = target / "qa"
    qa.mkdir(exist_ok=True)
    records = []
    with tempfile.TemporaryDirectory(prefix="sh-practice-lo-") as td:
        temp = Path(td)
        for p in make_packets(root):
            for mode in ["blank", "worked"]:
                path = target / p["slug"] / (mode + ".xlsx")
                out = temp / (p["slug"] + "-" + mode)
                out.mkdir()
                subprocess.run(
                    [
                        "libreoffice",
                        "-env:UserInstallation=" + (temp / "profile").as_uri(),
                        "--headless",
                        "--convert-to",
                        "pdf",
                        "--outdir",
                        str(out),
                        str(path),
                    ],
                    check=True,
                    stdout=subprocess.DEVNULL,
                    stderr=subprocess.PIPE,
                )
                pdf = out / (mode + ".pdf")
                document = fitz.open(pdf)
                images = []
                for page in document:
                    pix = page.get_pixmap(matrix=fitz.Matrix(1.4, 1.4))
                    images.append(Image.frombytes("RGB", [pix.width, pix.height], pix.samples))
                count = len(images)
                width = max(im.width for im in images)
                height = max(im.height for im in images)
                contact = Image.new(
                    "RGB", (width * 2 + 18, (height + 28) * ((count + 1) // 2)), "#dce1e1"
                )
                draw = ImageDraw.Draw(contact)
                for n, im in enumerate(images):
                    x = (n % 2) * (width + 9)
                    y = (n // 2) * (height + 28)
                    draw.text(
                        (x + 5, y + 4), f"{p['slug']} {mode}: native page {n + 1}", fill="black"
                    )
                    contact.paste(im, (x, y + 22))
                image = qa / (p["slug"] + "-" + mode + ".png")
                contact.save(image)
                records.append(
                    dict(
                        packet=p["id"],
                        mode=mode,
                        workbook_sha256=sha(path),
                        pages=count,
                        image=str(image.relative_to(root)),
                        image_sha256=sha(image),
                    )
                )
                document.close()
    write_json(qa / "renders.json", dict(status="RENDERED_REQUIRES_MANUAL_REVIEW", records=records))
    print(
        f"Rendered {len(records)} workbooks / {sum(r['pages'] for r in records)} native pages; manual inspection required"
    )


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", choices=["build", "validate", "render"])
    parser.add_argument("--root", type=Path, default=ROOT)
    args = parser.parse_args()
    if args.command == "build":
        build(args.root)
    elif args.command == "render":
        render(args.root)
    else:
        return bool(validate(args.root))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
