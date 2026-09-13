"""Consolidate every existing draft term and unresolved field without adopting it."""
import argparse
import hashlib
import json
import re
import sqlite3
import textwrap
from datetime import datetime
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parents[2]
HERE = ROOT / 'docs/legal/gap-instruments'
OUT = HERE / 'review-support'


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def anchor(heading):
    return re.sub(r'[^\w\- ]', '', heading.lower()).replace(' ', '-')


def records():
    config = json.loads((OUT / 'organization.json').read_text())
    mappings = {r['id']: r for r in config['items']}
    assert len(mappings) == len(config['items']), 'Duplicate decision mapping'
    rows = []
    for path in sorted((HERE / 'source').glob('*.json')):
        doc = json.loads(path.read_text())
        text = path.with_suffix('.md').read_text()
        for kind in ('unresolved_fields', 'proposed_terms'):
            for i, item in enumerate(doc[kind]):
                mapping = mappings.pop(item['id'])
                heading = mapping['clause_heading']
                assert '\n## ' + heading + '\n' in text, (path, heading)
                rows.append({
                    'id': item['id'], 'package': path.stem, 'group': mapping['group'],
                    'kind': 'Unresolved field' if kind == 'unresolved_fields' else 'Draft term and basis',
                    'text': item.get('field', item.get('term')),
                    'action_or_basis': item.get('next_action', item.get('basis')),
                    'status': 'PENDING_REVIEW_NOT_ACCEPTED',
                    'source_path': str(path.relative_to(ROOT)), 'source_sha256': digest(path),
                    'json_pointer': f'/{kind}/{i}', 'clause_heading': heading,
                    'clause_path': str(path.with_suffix('.md').relative_to(ROOT)) + '#' + anchor(heading),
                    'markdown_sha256': digest(path.with_suffix('.md')),
                })
    assert not mappings, 'Stale or extra decision mapping'
    return sorted(rows, key=lambda r: (r['group'], r['kind'], r['package'], r['id']))


def build():
    rows = records()
    payload = {'status': 'DRAFT_REVIEW_NAVIGATION', 'baseline_revision': '076485ccd954d3292fb5ad1940129027b6590284',
               'organization_sha256': digest(OUT / 'organization.json'), 'items': rows}
    (OUT / 'decisions.json').write_text(json.dumps(payload, indent=2) + '\n')
    lines = ['# Consolidated review decisions', '',
             ('Review related questions together, then record each disposition against its own field ID. '
             'A group is not blanket approval. Some draft terms repeat accepted source facts; read their basis before proposing a change.'), '',
             f'{len(rows)} source items across 17 packages. **Every item remains pending review.**', '',
             '[Filterable workbook](decisions.xlsx) · [Structured records](decisions.json) · [Database](review.sqlite3)', '',
             'The clause link opens the relevant draft reading context. The JSON pointer identifies the exact term or unresolved-field record.', '']
    groups = sorted({r['group'] for r in rows})
    for group in groups:
        lines += [f'## {group}', '']
        for r in [x for x in rows if x['group'] == group]:
            clause = '../source/' + r['clause_path'].split('/')[-1]
            source = '../source/' + Path(r['source_path']).name
            lines += [f"### {r['id']}", '', f"**{r['kind']} · {r['package']}**", '', r['text'], '',
                      '**Next action / source basis:** ' + r['action_or_basis'], '',
                      f"[Read clause {r['clause_heading']}]({clause}) · [Exact source field]({source}) `{r['json_pointer']}`", '']
    (OUT / 'DECISIONS.md').write_text('\n'.join(lines).rstrip() + '\n')
    wb = Workbook()
    intro = wb.active
    intro.title = 'Read first'
    for row in [
        ['SABLE HARBOR — consolidated review decisions'],
        ['DRAFT REVIEW NAVIGATION — no terms accepted'],
        ['Use Review filters for one topic. Item cells identify package and record type.'],
        ['Source details preserves every exact source pointer and original field.'],
        ['Reviewer response cells are intentionally blank. This workbook cannot confer acceptance.'],
        ['JSON pointers locate exact records; clause links provide reading context.'],
        ['Some draft terms restate accepted facts. Read the basis before changing them.'],
    ]:
        intro.append(row)
    intro.column_dimensions['A'].width = 100
    for row in intro:
        row[0].alignment = Alignment(wrap_text=True, vertical='top')
        intro.row_dimensions[row[0].row].height = 34
    review = wb.create_sheet('Review')
    review.append(['Item', 'Review topic', 'Question / term and basis', 'Draft clause', 'Reviewer response'])
    for r in rows:
        review.append([r['id'] + '\n' + r['package'] + '\n' + r['kind'], r['group'],
                       r['text'] + '\n\nNext action / source basis: ' + r['action_or_basis'], 'Read clause', None])
        review.cell(review.max_row, 4).hyperlink = '../source/' + Path(r['source_path']).with_suffix('.md').name + '#' + anchor(r['clause_heading'])
        review.cell(review.max_row, 4).style = 'Hyperlink'
    review_widths = [40, 25, 70, 18, 40]
    for i, width in enumerate(review_widths, 1):
        review.column_dimensions[chr(64 + i)].width = width
    for row in review:
        for c in row:
            c.alignment = Alignment(horizontal='left', wrap_text=True, vertical='top', indent=1 if c.column == 4 else 0)
            c.border = Border(bottom=Side(style='hair', color='D7DEE3'))
            c.font = Font(name='Aptos', size=11, color='174D70' if c.column == 4 and c.row > 1 else '20292C')
        line_counts = []
        for c in row:
            width = max(5, int(review_widths[c.column - 1] * .85))
            line_counts.append(sum(max(1, len(textwrap.wrap(line, width=width))) for line in str(c.value or '').split('\n')))
        wrapped = max(line_counts)
        review.row_dimensions[row[0].row].height = max(64, wrapped * 15 + 14) if row[0].row > 1 else 30
    for c in review[1]:
        c.fill = PatternFill('solid', fgColor='182C42')
        c.font = Font(name='Aptos', size=11, bold=True, color='FFFFFF')
    review.freeze_panes = 'A2'
    review.auto_filter.ref = review.dimensions
    review.sheet_view.zoomScale = 80
    review.sheet_view.showGridLines = False
    review.sheet_properties.pageSetUpPr.fitToPage = True
    review.page_setup.orientation = 'landscape'
    review.page_setup.paperSize = review.PAPERSIZE_A3
    review.page_setup.fitToWidth = 1
    review.page_setup.fitToHeight = 0
    review.print_title_rows = '1:1'
    review.print_options.horizontalCentered = True
    wb.active = 1
    ws = wb.create_sheet('Source details')
    headers = ['Field ID', 'Review topic', 'Package', 'Record type', 'Question or term', 'Next action or basis', 'Draft clause', 'Source pointer', 'Reviewer response']
    ws.append(headers)
    for r in rows:
        ws.append([r['id'], r['group'], r['package'], r['kind'], r['text'], r['action_or_basis'], r['clause_heading'], r['source_path'] + '#' + r['json_pointer'], None])
        ws.cell(ws.max_row, 7).hyperlink = '../source/' + Path(r['source_path']).with_suffix('.md').name + '#' + anchor(r['clause_heading'])
    widths = [36, 30, 23, 25, 75, 80, 55, 65, 55]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width
    for row in ws:
        for c in row:
            c.alignment = Alignment(horizontal='left', wrap_text=True, vertical='top', indent=1 if c.column in (6, 8) else 0)
            c.font = Font(name='Aptos', size=11)
        ws.row_dimensions[row[0].row].height = 75 if row[0].row > 1 else 30
    for c in ws[1]:
        c.fill = PatternFill('solid', fgColor='182C42')
        c.font = Font(name='Aptos', size=11, bold=True, color='FFFFFF')
    ws.freeze_panes = 'E2'
    ws.auto_filter.ref = ws.dimensions
    # This is an on-screen review register, not a compressed wall-sized print table.
    ws.sheet_view.zoomScale = 80
    wb.properties.created = wb.properties.modified = datetime(2026, 9, 13)  # noqa: DTZ001 - Excel metadata uses a fixed naive date.
    destination = OUT / 'decisions.xlsx'
    wb.save(destination)
    normalized = BytesIO()
    with ZipFile(destination) as source, ZipFile(normalized, 'w', ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            data = source.read(name)
            if name == 'docProps/core.xml':
                data = re.sub(rb'(<dcterms:(?:created|modified)[^>]*>)[^<]+', rb'\g<1>2026-09-13T00:00:00Z', data)
            entry = ZipInfo(name, (2026, 9, 13, 0, 0, 0))
            entry.compress_type = ZIP_DEFLATED
            target.writestr(entry, data)
    destination.write_bytes(normalized.getvalue())
    dbpath = OUT / 'review.sqlite3'
    dbpath.unlink(missing_ok=True)
    with sqlite3.connect(dbpath) as db:
        db.execute('CREATE TABLE decision (id TEXT PRIMARY KEY, package TEXT, topic TEXT, kind TEXT, text TEXT, source_path TEXT, json_pointer TEXT, clause_path TEXT, record_json TEXT)')
        db.executemany('INSERT INTO decision VALUES (?,?,?,?,?,?,?,?,?)', [(r['id'], r['package'], r['group'], r['kind'], r['text'], r['source_path'], r['json_pointer'], r['clause_path'], json.dumps(r, sort_keys=True)) for r in rows])
        db.execute('CREATE VIRTUAL TABLE decision_search USING fts5(id UNINDEXED, text, basis)')
        db.executemany('INSERT INTO decision_search VALUES (?,?,?)', [(r['id'], r['text'], r['action_or_basis']) for r in rows])
    print(f'{len(rows)} source items; {len(groups)} review topics; all 17 packages')


def validate():
    expected = records()
    saved = json.loads((OUT / 'decisions.json').read_text())
    assert saved['items'] == expected, 'Decision source drift'
    assert saved['organization_sha256'] == digest(OUT / 'organization.json')
    wb = load_workbook(OUT / 'decisions.xlsx')
    ws = wb['Source details']
    review = wb['Review']
    assert review.max_row == len(expected) + 1
    assert review.freeze_panes == 'A2'
    for n, r in enumerate(expected, 2):
        assert [review.cell(n, i).value for i in range(1, 6)] == [r['id'] + '\n' + r['package'] + '\n' + r['kind'], r['group'], r['text'] + '\n\nNext action / source basis: ' + r['action_or_basis'], 'Read clause', None]
        assert review.cell(n, 4).hyperlink.target == '../source/' + Path(r['source_path']).with_suffix('.md').name + '#' + anchor(r['clause_heading'])
    assert ws.max_row == len(expected) + 1
    for n, r in enumerate(expected, 2):
        assert [ws.cell(n, i).value for i in range(1, 9)] == [r['id'], r['group'], r['package'], r['kind'], r['text'], r['action_or_basis'], r['clause_heading'], r['source_path'] + '#' + r['json_pointer']]
        assert ws.cell(n, 9).value is None, 'Template must not claim reviewer responses'
    with sqlite3.connect(OUT / 'review.sqlite3') as db:
        actual = {r[0]: json.loads(r[1]) for r in db.execute('SELECT id,record_json FROM decision')}
        assert actual == {r['id']: r for r in expected}
        assert db.execute('SELECT count(*) FROM decision_search').fetchone()[0] == len(expected)
    print(f'PASS: {len(expected)} exact fields, workbook cells, clause anchors and database records')


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--check', action='store_true')
    args = ap.parse_args()
    validate() if args.check else build()
