"""Read-only native accounting linkage; build derivatives, never post journals."""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import sqlite3
import tempfile
import zipfile
from datetime import datetime
from decimal import Decimal as D
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[2]
HERE = Path("docs/legal/gap-instruments/accounting")


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def select(root, link):
    path = root / link["source_path"]
    sel = link["selector"]
    if "pointer" in sel:
        value = json.loads(path.read_text())
        for token in sel["pointer"].strip("/").split("/"):
            value = value[int(token)] if isinstance(value, list) else value[token]
        return value
    if sel.get("csv"):
        rows = list(csv.DictReader(io.StringIO(path.read_text())))
        return [r for r in rows if all(r[k] in v for k, v in sel.get("where", {}).items())]
    text = path.read_text()
    if "start" in sel:
        start = text.index(sel["start"])
        end = text.index(sel["end"], start)
        return text[start:end]
    if sel.get("whole_document"):
        return {"evidence_path": link["source_path"], "posting": None}
    raise ValueError("Unsupported selector")


def validate(root, data):
    gaps = json.loads((root / "docs/reader/transactions/gaps.json").read_text())
    expected = {g["id"]: g["status"] for g in gaps}
    actual = {g["gap_id"]: g["preserved_gap_status"] for g in data["coverage"]}
    if actual != expected or len(data["coverage"]) != len(expected):
        raise ValueError("Coverage or preserved disposition differs from controlling register")
    for path, sha in data["sources"].items():
        if digest(root / path) != sha:
            raise ValueError(f"Source hash changed: {path}")
    ids = set()
    resolved = []
    kinds = {"SOURCE_LEDGER", "MODELED_SCHEDULE", "MODEL_RULE", "NO_ENTRY"}
    for link in data["links"]:
        if link["id"] in ids or link["gap_id"] not in expected:
            raise ValueError("Duplicate link or unknown gap")
        ids.add(link["id"])
        if link["new_posting_authorized"] or link["evidence_kind"] not in kinds:
            raise ValueError("Invalid posting authority or evidence kind")
        for field in ("clause_path", "source_path"):
            if link[field] not in data["sources"]:
                raise ValueError("Unpinned source")
        if (
            "## " + link["clause_heading"]
            not in (root / link["clause_path"]).read_text().splitlines()
        ):
            raise ValueError("Clause heading missing")
        if not link["population"] or not link["interpretation"]:
            raise ValueError("Population or interpretation missing")
        if link["evidence_kind"] == "NO_ENTRY" and not link["selector"].get("whole_document"):
            raise ValueError("No-entry link cannot masquerade as posting rows")
        value = select(root, link)
        if value is None or value == []:
            raise ValueError("Empty native selector")
        resolved.append({**link, "native_value": value})
    for gap in data["coverage"]:
        if (
            gap["new_posting_authorized"]
            or set(gap["link_ids"]) != {x["id"] for x in resolved if x["gap_id"] == gap["gap_id"]}
            or not gap["link_ids"]
        ):
            raise ValueError("Incomplete or inconsistent link coverage")
    # Population assignment is explicit and cannot be weakened by changing the source hash.
    for link in resolved:
        path, pop = link["source_path"], link["population"]
        if "SH-FIN-HUMAN-001" in path and pop != "FF-2027-BASE":
            raise ValueError("Mixed invoice population")
        if "/tax-transaction/" in path and pop != "ARU-2026-MODEL":
            raise ValueError("Mixed ARU population")
        if path == "red_wash/source/core_operating_data.json" and pop != "RW-2025-ACQUISITION":
            raise ValueError("Mixed RW population")
    return resolved


def reconciliations(root):
    result = []

    def check(name, population, left, right, meaning):
        left, right = D(str(left)), D(str(right))
        if left != right:
            raise ValueError(f"{name}: {left} != {right}")
        result.append(
            dict(
                check=name,
                population=population,
                left=str(left),
                right=str(right),
                result="PASS",
                meaning=meaning,
            )
        )

    packet = json.loads((root / "docs/finance/evidence/SH-FIN-HUMAN-001/source.json").read_text())[
        "rows"
    ]
    inv = packet["invoices"][0]
    journal = packet["journal"]
    for jid in sorted({r["journal_id"] for r in journal}):
        rows = [r for r in journal if r["journal_id"] == jid]
        check(
            jid,
            "FF-2027-BASE",
            sum(D(r["debit_usd"]) for r in rows),
            sum(D(r["credit_usd"]) for r in rows),
            "Selected forecast journal event balances; not actual payment evidence.",
        )
    cash = sum(D(r["signed_usd"]) for r in journal if r["account"] == "1000")
    check(
        "FF cash including recovery",
        "FF-2027-BASE",
        cash,
        inv["collected_usd"],
        "Recovery is already included in total collections.",
    )
    claim = D(inv["writtenoff_usd"]) - D(inv["writtenoff_credit_usd"]) - D(inv["recovered_usd"])
    check(
        "FF economic claim bridge",
        "FF-2027-BASE",
        D(inv["amount_usd"]) - cash - D(inv["credit_usd"]),
        claim,
        "971,500 surviving written-off claim is not a ledger receivable.",
    )
    check(
        "FF ending AR",
        "FF-2027-BASE",
        sum(D(r["signed_usd"]) for r in journal if r["account"] == "BIZ_AR"),
        inv["remaining_usd"],
        "Selected invoice AR ends at zero.",
    )
    finance = json.loads((root / "industrial/source/finance.json").read_text())["transaction"]
    check(
        "Eight retention allocations",
        "ARU-2026-INPUT",
        sum(finance["retention_allocations"].values()),
        finance["retention_pool"],
        "500,000 award pool; excludes 225,000 consultancy and does not establish payment.",
    )
    base = root / "docs/finance/evidence/tax-transaction"
    ppa = {
        r["measure"]: D(r["amount_usd"])
        for r in csv.DictReader((base / "acquisition_ppa.csv").open())
    }
    check(
        "ARU consideration allocation",
        "ARU-2026-MODEL",
        ppa["identifiable_net_assets_before_refinancing_usd"] + ppa["goodwill_usd"],
        ppa["stock_consideration_usd"],
        "Conditional acquisition model, not filed election.",
    )
    check(
        "ARU sources and uses before fees",
        "ARU-2026-MODEL",
        ppa["close_sources_before_fees_usd"],
        ppa["close_uses_before_fees_usd"],
        "61,500,000 excludes fees; not added to stock consideration.",
    )
    tb = list(csv.DictReader((base / "aru_acquisition_opening_trial_balance.csv").open()))
    check(
        "ARU opening balance",
        "ARU-2026-MODEL",
        sum(D(r["debit_balance_usd"]) for r in tb),
        sum(D(r["credit_balance_usd"]) for r in tb),
        "ARU_GROUP aggregate opening model; no legal-entity consolidation claim.",
    )
    debt = list(csv.DictReader((base / "aru_2026_debt.csv").open()))
    check(
        "ARU term principal bridge",
        "ARU-2026-MODEL",
        D(debt[0]["opening_term_usd"]) - sum(D(r["term_principal_usd"]) for r in debt),
        debt[-1]["closing_term_usd"],
        "Monthly rounded debt schedule remains its own population.",
    )
    rw = json.loads((root / "red_wash/source/core_operating_data.json").read_text())["transaction"]
    check(
        "RW acquired net assets",
        "RW-2025-ACQUISITION",
        rw["operating_assets_usd"]
        + rw["current_assets_usd"]
        - rw["aro_assumed_usd"]
        - rw["other_liabilities_usd"],
        rw["cash_consideration_usd"],
        "28,000,000 historical consideration; title cure does not create a second acquisition.",
    )
    return result


def canonical_json(value):
    return json.dumps(value, indent=2, ensure_ascii=False) + "\n"


def build(root=ROOT, output=None):
    output = output or root / HERE
    output.mkdir(parents=True, exist_ok=True)
    data = json.loads((root / HERE / "source.json").read_text())
    links = validate(root, data)
    checks = reconciliations(root)
    payload = {**data, "links": links, "checks": checks}
    (output / "links.json").write_text(canonical_json(payload))
    lines = [
        "# Contract clauses and accounting evidence",
        "",
        data["boundary"],
        "",
        "Generated from [source.json](source.json). [Workbook](links.xlsx)"
        " · [SQLite](links.sqlite3).",
        "",
        "## Reconciliation",
        "",
        "| Check | Population | Verified USD | Result |",
        "| --- | --- | ---: | --- |",
    ]
    for c in checks:
        lines.append(f"| {c['check']} | {c['population']} | {D(c['left']):,.2f} | {c['result']} |")
    lines += ["", "Each amount belongs to its stated population. Do not total this table.", ""]
    for g in data["coverage"]:
        lines += [
            f"## {g['gap_id']}",
            "",
            f"Preserved disposition: `{g['preserved_gap_status']}`. "
            "New posting authorized: **No**.",
            "",
        ]
        for link in [x for x in links if x["gap_id"] == g["gap_id"]]:
            anchor = re.sub(r"[^\w\- ]", "", link["clause_heading"].lower()).replace(" ", "-")
            lines += [
                f"### {link['id']}",
                "",
                f"[{link['clause_heading']}](../../../../{link['clause_path']}#{anchor}) → "
                f"[{link['source_path']}](../../../../{link['source_path']})",
                "",
                f"`{link['evidence_kind']}` · `{link['population']}`",
                "",
                link["interpretation"],
                "",
                "Exact native selector: `"
                + json.dumps(link["selector"], ensure_ascii=False)
                + "`.",
                "",
            ]
            value = link["native_value"]
            if isinstance(value, list) and value and isinstance(value[0], dict):
                keys = list(value[0])
                lines += [
                    "| " + " | ".join(keys) + " |",
                    "| " + " | ".join("---" for _ in keys) + " |",
                ]
                lines += [
                    "| " + " | ".join(str(r.get(k, "")).replace("|", "\\|") for k in keys) + " |"
                    for r in value
                ]
            else:
                lines += [
                    "```json" if not isinstance(value, str) else "```python",
                    json.dumps(value, indent=2, ensure_ascii=False)
                    if not isinstance(value, str)
                    else value.rstrip(),
                    "```",
                ]
            lines += [""]
    (output / "LINKS.md").write_text("\n".join(lines))
    workbook(output / "links.xlsx", data, links, checks)
    db = output / "links.sqlite3"
    db.unlink(missing_ok=True)
    with sqlite3.connect(db) as con:
        con.execute("CREATE TABLE source(path TEXT PRIMARY KEY, sha256 TEXT NOT NULL)")
        con.executemany("INSERT INTO source VALUES (?,?)", sorted(data["sources"].items()))
        con.execute(
            "CREATE TABLE coverage(gap_id TEXT PRIMARY KEY, disposition TEXT, "
            "new_posting_authorized INTEGER CHECK(new_posting_authorized=0))"
        )
        con.executemany(
            "INSERT INTO coverage VALUES (?,?,0)",
            [(g["gap_id"], g["preserved_gap_status"]) for g in data["coverage"]],
        )
        con.execute(
            "CREATE TABLE clause_link(link_id TEXT PRIMARY KEY, gap_id TEXT RE"
            "FERENCES coverage, population TEXT, evidence_kind TEXT, clause_pa"
            "th TEXT, clause_heading TEXT, source_path TEXT REFERENCES source,"
            " selector_json TEXT, interpretation TEXT, native_value_json TEXT)"
        )
        for link in links:
            con.execute(
                "INSERT INTO clause_link VALUES (?,?,?,?,?,?,?,?,?,?)",
                (
                    link["id"],
                    link["gap_id"],
                    link["population"],
                    link["evidence_kind"],
                    link["clause_path"],
                    link["clause_heading"],
                    link["source_path"],
                    json.dumps(link["selector"]),
                    link["interpretation"],
                    json.dumps(link["native_value"]),
                ),
            )
        con.execute(
            "CREATE TABLE reconciliation(name TEXT PRIMARY KEY, population TEX"
            "T, left_usd TEXT, right_usd TEXT, result TEXT, meaning TEXT)"
        )
        con.executemany(
            "INSERT INTO reconciliation VALUES (?,?,?,?,?,?)",
            [
                (c["check"], c["population"], c["left"], c["right"], c["result"], c["meaning"])
                for c in checks
            ],
        )
    return payload


def workbook(path, data, links, checks):
    wb = Workbook()
    wb.remove(wb.active)

    def sheet(name, headers, rows, widths):
        ws = wb.create_sheet(name)
        ws.append(headers)
        for row in rows:
            ws.append(
                [
                    json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v
                    for v in row
                ]
            )
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False
        for row in ws:
            for cell in row:
                cell.font = Font(
                    name="Arial",
                    size=11,
                    bold=cell.row == 1,
                    color="FFFFFF" if cell.row == 1 else "243238",
                )
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if cell.row == 1:
                    cell.fill = PatternFill("solid", fgColor="243238")
            height = max(
                len(str(c.value or "")) / max(widths[min(c.column - 1, len(widths) - 1)] - 3, 1)
                for c in row
            )
            ws.row_dimensions[row[0].row].height = max(32, (int(height) + 2) * 15)
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = width
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_title_rows = "1:1"
        return ws

    sheet(
        "Read first",
        ["Topic", "Reading instruction"],
        [
            ("Status", data["boundary"]),
            (
                "Use",
                "Find a gap in Coverage, then locate its numbered Evidence sheet. "
                "Exact paths and hashes are in Sources. No totals across populatio"
                "ns.",
            ),
            (
                "Types",
                "SOURCE_LEDGER means existing synthetic forecast journal rows; MOD"
                "ELED_SCHEDULE means model amounts; MODEL_RULE is generator logic,"
                " not emitted evidence; NO_ENTRY supplies no posting.",
            ),
            (
                "Review",
                "This new workbook design is held for exact-file review. No legal "
                "draft is accepted through these links.",
            ),
        ],
        [22, 112],
    )
    sheet(
        "Coverage",
        ["Gap ID", "Preserved disposition", "Link IDs"],
        [
            (g["gap_id"], g["preserved_gap_status"], "\n".join(g["link_ids"]))
            for g in data["coverage"]
        ],
        [40, 54, 65],
    )
    sheet(
        "Reconciliation",
        ["Check", "Population", "Verified USD", "Result", "Meaning"],
        [
            (c["check"], c["population"], float(c["left"]), c["result"], c["meaning"])
            for c in checks
        ],
        [38, 27, 22, 12, 85],
    )
    for cell in wb["Reconciliation"]["D"][1:]:
        cell.alignment = Alignment(horizontal="center", vertical="top")
    for cell in wb["Reconciliation"]["C"][1:]:
        cell.number_format = "#,##0.00;[Red](#,##0.00)"
    sheet(
        "Links",
        ["ID / evidence sheet", "Clause", "Type / population", "Interpretation"],
        [
            (
                f"{i:02}: {x['id']}",
                x["clause_heading"],
                x["evidence_kind"] + " / " + x["population"],
                x["interpretation"],
            )
            for i, x in enumerate(links, 1)
        ],
        [38, 38, 28, 62],
    )
    sheet("Sources", ["Path", "SHA-256"], sorted(data["sources"].items()), [110, 72])
    for i, link in enumerate(links, 1):
        value = link["native_value"]
        # Tall key/value layout keeps every native field readable without a 20-column printout.
        rows = [
            ("Link", link["id"]),
            ("Source", link["source_path"]),
            ("Selector", json.dumps(link["selector"])),
            ("Population", link["population"]),
        ]
        if isinstance(value, list):
            for n, record in enumerate(value, 1):
                rows.extend((f"Record {n} / {k}", str(v)) for k, v in record.items())
        elif isinstance(value, dict):
            rows.extend((k, str(v)) for k, v in value.items())
        elif isinstance(value, str):
            rows.extend((f"Rule line {n}", line) for n, line in enumerate(value.splitlines(), 1))
        else:
            rows.append(("Native value", str(value)))
        sheet(f"Evidence {i:02}", ["Native locator", "Native value"], rows, [48, 120])
    wb.properties.created = wb.properties.modified = datetime(2026, 9, 12)
    buf = io.BytesIO()
    wb.save(buf)
    with zipfile.ZipFile(buf) as src, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as dst:
        for name in sorted(src.namelist()):
            content = src.read(name)
            if name == "docProps/core.xml":
                content = re.sub(
                    rb"<dcterms:modified[^>]*>.*?</dcterms:modified>",
                    b'<dcterms:modified xmlns:dcterms="http://purl.org/dc/terms/" xmlns'
                    b':xsi="http://www.w3.org/2001/XMLSchema-instance" xsi:type="dcterm'
                    b's:W3CDTF">2026-09-12T00:00:00Z</dcterms:modified>',
                    content,
                )
            info = zipfile.ZipInfo(name, (2026, 9, 12, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(info, content)


def database_snapshot(path):
    """Compare logical databases without depending on SQLite storage-version bytes."""
    with sqlite3.connect(path.resolve().as_uri() + "?mode=ro", uri=True) as con:
        if con.execute("PRAGMA integrity_check").fetchall() != [("ok",)]:
            raise ValueError(f"Invalid accounting database integrity: {path}")
        if con.execute("PRAGMA foreign_key_check").fetchall():
            raise ValueError(f"Invalid accounting database foreign keys: {path}")
        schema = con.execute(
            "SELECT type,name,tbl_name,sql FROM sqlite_schema ORDER BY type,name"
        ).fetchall()
        tables = {}
        for kind, name, _, _ in schema:
            if kind != "table":
                continue
            quoted = '"' + name.replace('"', '""') + '"'
            rows = con.execute("SELECT * FROM " + quoted).fetchall()
            # repr distinguishes NULL, numeric, text and BLOB values; preserve duplicate rows.
            tables[name] = sorted(rows, key=repr)
        return {
            "schema": schema,
            "tables": tables,
            "user_version": con.execute("PRAGMA user_version").fetchone()[0],
            "application_id": con.execute("PRAGMA application_id").fetchone()[0],
        }


def compare_databases(expected, actual):
    if database_snapshot(expected) != database_snapshot(actual):
        raise ValueError("Stale accounting derivative: links.sqlite3 schema or contents")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--check", action="store_true")
    args = parser.parse_args()
    if args.check:
        with tempfile.TemporaryDirectory() as tmp:
            payload = build(output=Path(tmp))
            for name in ("LINKS.md", "links.json", "links.xlsx"):
                if digest(Path(tmp) / name) != digest(ROOT / HERE / name):
                    raise ValueError("Stale accounting derivative: " + name)
            compare_databases(Path(tmp) / "links.sqlite3", ROOT / HERE / "links.sqlite3")
    else:
        payload = build()
    print(
        f"PASS: {len(payload['coverage'])} gaps; "
        f"{len(payload['links'])} clause links; "
        f"{len(payload['checks'])} reconciliations; no new postings"
    )


if __name__ == "__main__":
    main()
