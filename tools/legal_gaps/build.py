"""Reproduce complete draft instruments; never alter accepted evidence dispositions."""
from __future__ import annotations

import base64
import hashlib
import html
import json
import re
import sqlite3
from datetime import UTC, datetime
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

import fitz
from markdown_it import MarkdownIt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from playwright.sync_api import sync_playwright

ROOT = Path(__file__).resolve().parents[2]
HERE = ROOT / 'docs/legal/gap-instruments'
BASE = '659a56747fe76522d18645ff115888c13fa8d2b0'
STATUS = 'DRAFT_FOR_REVIEW'
TOOLS = Path(__file__).resolve().parent


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write_json(path, data):
    path.write_text(json.dumps(data, indent=2, ensure_ascii=False) + '\n')


def parser():
    return MarkdownIt('commonmark', {'html': False}).enable('table').enable('strikethrough')


def workbook(record, output):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Read first'
    ws.append([record['title']])
    ws.append([record['document_id'], STATUS])
    ws.append(['Prepared', '2026-09-12'])
    ws.append(['Purpose', 'Draft schedules. No signature, payment, filing or approved new terms implied.'])
    ws.append(['Source', 'Companion Markdown and JSON control; amounts keep their stated source or proposal status.'])
    ws.append(['Review', 'Exact files and proposed terms require acceptance.'])
    for n, schedule in enumerate(record.get('financial_schedules', []), 1):
        name = re.sub(r'[\\/*?:\[\]]', '', schedule['name'])[:27]
        sheet = wb.create_sheet(f'{n:02d} {name}'[:31])
        sheet.append([schedule['name']])
        sheet.append([STATUS, 'Not a payment or filing record'])
        sheet.append(schedule['columns'])
        for row in schedule['rows']:
            sheet.append(row)
        sheet.freeze_panes = 'A4'
        sheet.auto_filter.ref = f'A3:{get_column_letter(len(schedule["columns"]))}{sheet.max_row}'
    for sheet in wb:
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = sheet.freeze_panes or 'A2'
        for row in sheet:
            for cell in row:
                cell.font = Font(name='Arial', size=11, color='101214')
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if cell.row in (1, 3) and sheet != ws or cell.row == 1:
                    cell.fill = PatternFill('solid', fgColor='243238')
                    cell.font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
                elif cell.row % 2 == 0:
                    cell.fill = PatternFill('solid', fgColor='F4F1EA')
                if isinstance(cell.value, (float, int)):
                    cell.number_format = '#,##0.00;[Red](#,##0.00);0.00'
        for col in sheet.columns:
            max_chars = max(len(str(c.value or '')) for c in col)
            sheet.column_dimensions[col[0].column_letter].width = min(52, max(20, max_chars + 2))
        for row in sheet:
            lines = max((len(str(c.value or '')) // max(1, int(sheet.column_dimensions[c.column_letter].width) - 3) + 1) for c in row)
            sheet.row_dimensions[row[0].row].height = max(28, lines * 16 + 10)
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = '1:3'
        sheet.print_options.horizontalCentered = True
        sheet.oddFooter.center.text = 'DRAFT FOR REVIEW | Page &P of &N'
    wb.properties.created = datetime(2026, 9, 12, tzinfo=UTC)
    wb.properties.modified = datetime(2026, 9, 12, tzinfo=UTC)
    wb.save(output)
    # Normalize ZIP headers as well as core properties for reproducible bytes.
    with ZipFile(output) as z:
        files = [(n, z.read(n)) for n in sorted(z.namelist())]
    with ZipFile(output, 'w', ZIP_DEFLATED) as z:
        for name, data in files:
            info = ZipInfo(name, (2026, 9, 12, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            z.writestr(info, data)
    return [s.title for s in wb]


def build():
    HERE.mkdir(exist_ok=True)
    out = HERE / 'editions'
    out.mkdir(exist_ok=True)
    logo = ROOT / 'assets/brand/logos/sable-harbor__primary-horizontal.svg'
    logo_uri = 'data:image/svg+xml;base64,' + base64.b64encode(logo.read_bytes()).decode()
    css = (TOOLS / 'style.css').read_text()
    records = []
    for jp in sorted((HERE / 'source').glob('*.json')):
        record = json.loads(jp.read_text())
        record['slug'] = jp.stem
        record['source'] = str(jp.with_suffix('.md').relative_to(ROOT))
        record['structured'] = str(jp.relative_to(ROOT))
        record['source_sha256'] = sha(jp.with_suffix('.md'))
        record['structured_sha256'] = sha(jp)
        record['dependencies'] = [{'path': p, 'sha256': sha(ROOT / p)} for p in record['source_paths']]
        records.append(record)
    if len(records) != 17:
        raise ValueError(f'Expected17 packages, received{len(records)}')
    with sync_playwright() as pw:
        browser = pw.chromium.launch(executable_path='/usr/bin/chromium', args=['--no-sandbox'])
        for record in records:
            body = parser().render((ROOT / record['source']).read_text())
            # Editions live beside source/, so relative source links retain depth.
            hp = out / f'{record["slug"]}.html'
            hp.write_text(f'''<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{html.escape(record['title'])}</title><style>{css}</style></head><body>
<header class="edition-header"><img src="{logo_uri}" alt="Sable Harbor"><span>LEGAL &amp; CORPORATE RECORDS<br>DRAFT INSTRUMENT · 12 SEPTEMBER 2026</span></header>
<aside class="edition-note"><strong>Draft for exact-file review</strong><br>Proposed language. Accepted source facts retain their original status. No execution, filing, payment or clearance is established by this edition.</aside>
<article id="source-content">{body}</article><aside class="source-identity"><a href="../review.html">All 17 packages</a> · <a href="../source/{record['slug']}.md">Complete Markdown</a><br>{record['source_sha256']}</aside></body></html>''')
            page = browser.new_page()
            page.goto(hp.as_uri(), wait_until='networkidle')
            raw = page.pdf(format='Letter', print_background=True, display_header_footer=True,
                header_template='<div style="font-family:Arial;font-size:8px;color:#747A80;width:100%;margin:0 48px;text-align:right">SABLE HARBOR · DRAFT INSTRUMENT</div>',
                footer_template='<div style="font-family:Arial;font-size:8px;color:#747A80;width:100%;margin:0 48px;display:flex;justify-content:space-between"><span>' + html.escape(record['document_id']) + '</span><span><span class="pageNumber"></span> / <span class="totalPages"></span></span></div>')
            page.close()
            pdf = fitz.open(stream=raw, filetype='pdf')
            pdf.set_metadata({})
            pp = hp.with_suffix('.pdf')
            pdf.save(pp, garbage=4, deflate=True, no_new_id=True)
            record['pages'] = len(pdf)
            pdf.close()
            record['artifacts'] = []
            for path in (hp, pp):
                record['artifacts'].append({'path': str(path.relative_to(ROOT)), 'sha256': sha(path)})
            if record.get('financial_schedules'):
                xp = hp.with_suffix('.xlsx')
                record['sheets'] = workbook(record, xp)
                record['artifacts'].append({'path': str(xp.relative_to(ROOT)), 'sha256': sha(xp)})
            print(f'{record["gap_id"]}: {record["pages"]} pages', flush=True)
        browser.close()
    db = HERE / 'instruments.sqlite3'
    if db.exists():
        db.unlink()
    with sqlite3.connect(db) as con:
        con.execute('CREATE TABLE instruments (gap_id TEXT PRIMARY KEY, document_id TEXT UNIQUE, title TEXT, status TEXT, source_path TEXT, body TEXT, structured_json TEXT)')
        con.execute('CREATE VIRTUAL TABLE instrument_search USING fts5(gap_id UNINDEXED,title,body)')
        for r in records:
            text = (ROOT / r['source']).read_text()
            con.execute('INSERT INTO instruments VALUES (?,?,?,?,?,?,?)', (r['gap_id'],r['document_id'],r['title'],STATUS,r['source'],text,(ROOT / r['structured']).read_text()))
            con.execute('INSERT INTO instrument_search VALUES (?,?,?)',(r['gap_id'],r['title'],text))
        con.commit()
        con.execute('VACUUM')
    write_json(HERE / 'index.json', {'status': STATUS, 'base_revision': BASE, 'packages': records})
    rows = []
    for r in records:
        links = [f'<a href="source/{r["slug"]}.md">Markdown</a>', f'<a href="source/{r["slug"]}.json">Data</a>']
        for a in r['artifacts']:
            p = Path(a['path'])
            links.append(f'<a href="editions/{p.name}">{p.suffix[1:].upper()}</a>')
        rows.append(f'<tr><td><strong>{html.escape(r["title"])}</strong><br><small>{r["gap_id"]}</small></td><td>{r["pages"]}</td><td>{" · ".join(links)}</td></tr>')
    (HERE / 'review.html').write_text(f'''<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>17 legal gap packages — review</title><style>{css}</style></head><body><header class="edition-header"><img src="{logo_uri}" alt="Sable Harbor"><span>LEGAL RECORDS<br>REVIEW ROOM</span></header><h1>17 legal gap packages</h1><p>Complete draft instruments and schedules for review. These files develop the missing language; they do not establish execution or close external evidence gaps.</p><aside class="edition-note"><strong>Drafts awaiting exact-file review.</strong> Proposed terms remain unaccepted. Read each package’s source and proposal notes before treating a provision as a company fact.</aside><p><a href="README.md">Scope and reproduction</a> · <a href="index.json">Structured index</a> · <a href="instruments.sqlite3">Searchable database</a></p><table><thead><tr><th>Package</th><th>PDF pages</th><th>Open the complete files</th></tr></thead><tbody>{''.join(rows)}</tbody></table></body></html>''')
    files = [HERE / 'index.json', HERE / 'review.html', db]
    write_json(HERE / 'manifest.json', {'status': STATUS, 'approved': False, 'base_revision': BASE,
        'inputs': [{'path':str(p.relative_to(ROOT)), 'sha256':sha(p)} for p in (Path(__file__), TOOLS/'style.css',TOOLS/'requirements.txt',logo)],
        'packages': records, 'other_artifacts':[{'path':str(p.relative_to(ROOT)),'sha256':sha(p)} for p in files],
        'counts':{'packages':len(records),'pdf_pages':sum(r['pages'] for r in records),'workbooks':sum(bool(r.get('sheets')) for r in records),'sheets':sum(len(r.get('sheets',[])) for r in records)}})


if __name__ == '__main__':
    build()
