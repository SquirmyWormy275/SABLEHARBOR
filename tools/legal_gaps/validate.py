"""Validate source coverage, complete publication text, schedules and lineage."""
from __future__ import annotations

import hashlib
import json
import sqlite3
import sys
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import unquote, urlsplit

import fitz
from markdown_it import MarkdownIt
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
HERE = ROOT / 'docs/legal/gap-instruments'


def norm(value):
    return ''.join(c.lower() for c in value if c.isalnum())


class HTML(HTMLParser):
    def __init__(self):
        super().__init__()
        self.text = []
        self.links = []
        self.article = False

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == 'article':
            self.article = True
        if tag == 'a' and 'href' in attrs:
            self.links.append(attrs['href'])

    def handle_endtag(self, tag):
        if tag == 'article':
            self.article = False

    def handle_data(self, data):
        if self.article:
            self.text.append(data)


def source_blocks(text):
    result = []
    for token in MarkdownIt('commonmark').enable('table').enable('strikethrough').parse(text):
        if token.type == 'inline':
            result.append(''.join(t.content for t in token.children or [] if t.type in ('text','code_inline','softbreak','hardbreak')))
        elif token.type in ('fence', 'code_block'):
            result.append(token.content)
    return [norm(x) for x in result if norm(x)]


def validate():
    errors = []
    def check(condition, message):
        if not condition:
            errors.append(message)
    manifest = json.loads((HERE / 'manifest.json').read_text())
    gaps = json.loads((ROOT / 'docs/reader/transactions/gaps.json').read_text())
    if isinstance(gaps, dict):
        gaps = gaps.get('gaps', gaps.get('records'))
    expected = {g.get('gap_id', g.get('id')) for g in gaps}
    records = manifest['packages']
    check(len(records) == 17, 'Expected17 packages')
    check({r['gap_id'] for r in records} == expected, 'Gap coverage mismatch')
    check(len({r['document_id'] for r in records}) == 17, 'Duplicate document IDs')
    check(manifest['approved'] is False, 'Draft promoted to approved')
    total_blocks = 0
    with sqlite3.connect(HERE / 'instruments.sqlite3') as con:
        check(con.execute('PRAGMA integrity_check').fetchone()[0] == 'ok', 'Database corrupt')
        check(con.execute('SELECT count(*) FROM instruments').fetchone()[0] == 17, 'Database missing packages')
        for r in records:
            text = (ROOT / r['source']).read_text()
            data = json.loads((ROOT / r['structured']).read_text())
            check(data['status'] == 'DRAFT_FOR_REVIEW', f'{r["gap_id"]}: invalid state')
            check(data['proposed_terms'] or data['unresolved_fields'], f'{r["gap_id"]}: no proposal boundary')
            for field in ('source', 'structured'):
                digest = hashlib.sha256((ROOT/r[field]).read_bytes()).hexdigest()
                check(digest == r[field+'_sha256'], f'Stale {r[field]}')
            row = con.execute('SELECT body,structured_json,status FROM instruments WHERE gap_id=?',(r['gap_id'],)).fetchone()
            check(row == (text,(ROOT/r['structured']).read_text(),'DRAFT_FOR_REVIEW'), f'{r["gap_id"]}: database divergence')
            search = con.execute('SELECT body FROM instrument_search WHERE gap_id=?',(r['gap_id'],)).fetchone()
            check(search == (text,), f'{r["gap_id"]}: FTS divergence')
            blocks = source_blocks(text)
            total_blocks += len(blocks)
            for artifact in r['artifacts']:
                path = ROOT / artifact['path']
                if path.suffix == '.html':
                    parser = HTML()
                    parser.feed(path.read_text())
                    complete = norm(' '.join(parser.text))
                    for block in blocks:
                        check(block in complete, f'{path.name}: missing HTML block {block[:70]}')
                elif path.suffix == '.pdf':
                    with fitz.open(path) as doc:
                        complete = norm(' '.join(p.get_text(sort=True) for p in doc))
                        check(len(doc) == r['pages'], f'{path.name}: page count')
                        for block in blocks:
                            check(block in complete, f'{path.name}: missing PDF block {block[:70]}')
                        for page in doc:
                            for x0,y0,x1,y1,word,*_ in page.get_text('words'):
                                check(x0 >= 0 and y0 >= 0 and x1 <= page.rect.width+1 and y1 <= page.rect.height+1, f'{path.name}: off-page text {word}')
                elif path.suffix == '.xlsx':
                    wb = load_workbook(path, data_only=False)
                    check(wb.sheetnames == r['sheets'], f'{path.name}: sheet mismatch')
                    for sheet, schedule in zip(list(wb)[1:],data['financial_schedules'],strict=True):
                        actual = [list(row) for row in sheet.iter_rows(min_row=4,max_col=len(schedule['columns']),values_only=True)]
                        wanted = [[None if x == '' else x for x in row] for row in schedule['rows']]
                        check(actual == wanted, f'{path.name}/{sheet.title}: changed schedule cells')
    all_hashed = manifest['inputs'] + manifest['other_artifacts']
    for r in records:
        all_hashed += r['dependencies'] + r['artifacts']
    for item in all_hashed:
        path = ROOT/item['path']
        check(path.is_file(), f'Missing {path}')
        if path.is_file():
            check(hashlib.sha256(path.read_bytes()).hexdigest() == item['sha256'], f'Stale {item["path"]}')
    for path in [HERE/'review.html', *sorted((HERE/'editions').glob('*.html'))]:
        parser = HTML()
        parser.feed(path.read_text())
        for href in parser.links:
            parts = urlsplit(href)
            if parts.scheme or not parts.path:
                continue
            target = (path.parent/unquote(parts.path)).resolve()
            check(target.is_file(), f'{path.name}: broken link {href}')
    for path in (HERE/'source').glob('*.md'):
        for token in MarkdownIt('commonmark').enable('table').parse(path.read_text()):
            for child in token.children or []:
                if child.type == 'link_open':
                    href = child.attrGet('href')
                    parts = urlsplit(href)
                    if not parts.scheme and parts.path:
                        check((path.parent/unquote(parts.path)).resolve().is_file(), f'{path.name}: broken source link {href}')
    if errors:
        print('\n'.join(errors))
        return 1
    print(f'PASS:17 packages; {total_blocks} complete source blocks; PDF/HTML/SQLite/workbook cells, hashes and links reconcile')
    return 0


if __name__ == '__main__':
    sys.exit(validate())
