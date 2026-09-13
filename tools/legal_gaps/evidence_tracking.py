"""Source-pinned evidence requests and safe append-only review proposals."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sqlite3
import tempfile
from datetime import UTC, datetime
from pathlib import Path

import xlsxwriter
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
DEST = Path("docs/legal/gap-instruments/evidence-tracking")
SOURCES = [
    "docs/legal/gap-instruments/walkthroughs/source/evidence_requests.csv",
    "docs/legal/gap-instruments/reconciliations/source.json",
]
STATES = ("unresolved", "incomplete", "received", "disputed")
HEADERS = [
    "Request ID",
    "Current status",
    "New status",
    "Evidence path",
    "Evidence SHA-256",
    "Exact text citation",
    "Review note",
    "Event ID",
    "Observed date (YYYY-MM-DD)",
]
EDITABLE = set(range(3, 10))
STATUS = "DRAFT_REVIEW_PROPOSAL_NOT_APPLIED"


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def require(condition, message):
    if not condition:
        raise ValueError(message)


def dump(path, value):
    path.write_text(json.dumps(value, indent=2, ensure_ascii=False) + "\n")


def digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True).encode()).hexdigest()


def model(root=ROOT):
    with (root / SOURCES[0]).open() as stream:
        requests = [
            dict(
                id=r["id"],
                scope=r["scope"],
                request=r["missing_record"],
                purpose=r["purpose"],
                blocked_conclusion=r["blocked_conclusion"],
                source_path=SOURCES[0],
                source_locator="CSV id=" + r["id"],
                context_path=r["source_path"],
            )
            for r in csv.DictReader(stream)
        ]
    recon = json.loads((root / SOURCES[1]).read_text())
    for n, p in enumerate(recon["workpapers"]):
        requests.append(
            dict(
                id=p["id"] + "-REQ",
                scope=p["scope"],
                request=p["evidence_request"],
                purpose="Corroborate the selected model rows with the requested originals.",
                blocked_conclusion=p["limitation"],
                source_path=SOURCES[1],
                source_locator=f"/workpapers/{n}/evidence_request",
                context_path=SOURCES[1],
            )
        )
    close_path = "docs/legal/gap-instruments/period-close/case.json"
    close = json.loads((root / close_path).read_text())
    close_requests = [
        (
            "AR",
            2,
            "Obtain ARU_GROUP January 31, 2027 gross external receivables aging, allowance calculation and subsequent receipts supporting account 1100 in the base scenario.",  # noqa: E501 -- reader-facing prose
        ),
        (
            "AP",
            3,
            "Obtain ARU_GROUP January 31, 2027 supplier open items, invoices and subsequent disbursements supporting trade payable account 2000 in the base scenario.",  # noqa: E501 -- reader-facing prose
        ),
        (
            "PPE",
            4,
            "Obtain ARU_GROUP January 2027 PPE invoices, title, commissioning records, physical inspection and useful-life support for accounts 1400/1490 in the base scenario; keep finance-lease ROU separate.",  # noqa: E501 -- reader-facing prose
        ),
    ]
    for suffix, index, request in close_requests:
        requests.append(
            dict(
                id=f"SH-CLOSE-ARU-2027-01-{suffix}-REQ",
                scope=close["scope"],
                request=request,
                purpose="Corroborate this reporting book and period with the requested originals.",
                blocked_conclusion=close["limitations"][index],
                source_path=close_path,
                source_locator=f"/limitations/{index}",
                context_path="docs/legal/gap-instruments/period-close/README.md",
            )
        )
    require(
        len(requests) == 14 and len({r["id"] for r in requests}) == 14, "Request population drift"
    )
    pins = {
        p: sha(root / p)
        for p in sorted(
            set(
                SOURCES
                + [r["context_path"] for r in requests]
                + [r["source_path"] for r in requests]
            )
        )
    }
    return dict(
        id="SH-EVIDENCE-FOLLOWTHROUGH-001",
        status=STATUS,
        source_pins=pins,
        requests=requests,
        history=[],
        current={r["id"]: "unresolved" for r in requests},
    )


def workbook(data, path):
    wb = xlsxwriter.Workbook(path, {"strings_to_formulas": False, "strings_to_urls": False})
    wb.set_properties(
        {
            "title": "Evidence request follow-through",
            "company": "SABLE HARBOR",
            "created": datetime(2026, 9, 13, tzinfo=UTC),
        }
    )
    dark = wb.add_format(
        {
            "font_name": "Aptos",
            "font_size": 11,
            "bold": True,
            "bg_color": "#172C35",
            "font_color": "#FFFFFF",
            "text_wrap": True,
            "valign": "vcenter",
        }
    )
    body = wb.add_format(
        {
            "font_name": "Aptos",
            "font_size": 11,
            "text_wrap": True,
            "valign": "top",
            "bottom": 1,
            "bottom_color": "#D2D9DA",
        }
    )
    edit = wb.add_format(
        {
            "font_name": "Aptos",
            "font_size": 11,
            "text_wrap": True,
            "valign": "top",
            "bg_color": "#FFF4D8",
            "bottom": 1,
            "bottom_color": "#D2D9DA",
        }
    )
    for name in ["Read first", "Requests", "Updates", "History", "Context"]:
        ws = wb.add_worksheet(name)
        ws.set_landscape()
        ws.set_paper(9)
        ws.fit_to_pages(1, 0)
        ws.set_margins(0.3, 0.3, 0.4, 0.4)
        ws.set_footer(
            "&LSABLE HARBOR | DRAFT REVIEW&C" + name + "&RPage &P of &N", {"margin": 0.15}
        )
        ws.hide_gridlines(2)
        ws.set_row(0, 28)
        if name == "Read first":
            ws.set_column(0, 0, 28)
            ws.set_column(1, 1, 115)
            rows = [
                ("SABLE HARBOR", "Evidence request follow-through"),
                (
                    "Purpose",
                    "Track 14 evidence requests: 11 existing and three for the ARU January 2027 close. All begin unresolved. A supplied record does not establish reliability, execution, approval or closure.",  # noqa: E501 -- reader-facing prose
                ),
                (
                    "Use a copy",
                    "Save a copy of this workbook outside the repository. On Updates, fill the yellow cells for an event. Leave unused rows blank. Do not change IDs, current status, context or other sheets.",  # noqa: E501 -- reader-facing prose
                ),
                (
                    "States",
                    "unresolved: no usable support supplied; incomplete: support supplied but requested coverage remains missing; received: a record supplied, not validated; disputed: an identified objection remains unresolved.",  # noqa: E501 -- reader-facing prose
                ),
                (
                    "Evidence",
                    "Use an existing public repository UTF-8 text record, its full SHA-256 and an exact text excerpt of at least 12 characters. Imported proposals retain the matched line range. Binary evidence needs a reviewed text transcript first.",  # noqa: E501 -- reader-facing prose
                ),
                (
                    "Event details",
                    "Use a unique event ID, observed date YYYY-MM-DD and a substantive note. Explain remaining limitations for incomplete/disputed events. No status means no event.",  # noqa: E501 -- reader-facing prose
                ),
                (
                    "Import",
                    "python tools/legal_gaps/evidence_tracking.py import --workbook /tmp/completed.xlsx --output /tmp/new-evidence-review",  # noqa: E501 -- reader-facing prose
                ),
                (
                    "Continue history",
                    "The import creates tracker.json, tracker.xlsx, tracker.sqlite3 and PROPOSAL.md in a new external folder. Use that tracker.xlsx and --state /tmp/new-evidence-review/tracker.json for the next event. Earlier history is retained and hash chained.",  # noqa: E501 -- reader-facing prose
                ),
                (
                    "Review boundary",
                    "There is no apply or close command. Review the proposed events and source files. Received never closes the original request or authorizes a journal entry. Unresolved source limitations remain in Requests.",  # noqa: E501 -- reader-facing prose
                ),
            ]
            for n, row in enumerate(rows):
                ws.write_row(n, 0, row, dark if n == 0 else body)
                ws.set_row(n, 48 if n else 28)
            ws.print_area(0, 0, len(rows) - 1, 1)
        elif name == "Requests":
            ws.set_column(0, 0, 25)
            ws.set_column(1, 1, 27)
            ws.set_column(2, 2, 43)
            ws.set_column(3, 3, 43)
            ws.write_row(
                0,
                0,
                ["Request ID", "Scope", "Requested evidence", "Conclusion still unsupported"],
                dark,
            )
            for n, r in enumerate(data["requests"], 1):
                ws.write_row(
                    n, 0, [r["id"], r["scope"], r["request"], r["blocked_conclusion"]], body
                )
                ws.set_row(n, 80)
            ws.autofilter(0, 0, len(data["requests"]), 3)
            ws.repeat_rows(0)
        elif name == "Updates":
            widths = [25, 14, 14, 42, 28, 45, 45, 24, 22]
            for n, width in enumerate(widths):
                ws.set_column(n, n, width)
            ws.write_row(0, 0, HEADERS, dark)
            for n, r in enumerate(data["requests"], 1):
                ws.write_row(n, 0, [r["id"], data["current"][r["id"]]], body)
                for col in range(2, 9):
                    ws.write_blank(n, col, None, edit)
                ws.set_row(n, 74)
            ws.data_validation(
                1, 2, len(data["requests"]), 2, {"validate": "list", "source": list(STATES)}
            )
            ws.autofilter(0, 0, len(data["requests"]), 8)
            ws.freeze_panes(1, 2)
            ws.repeat_rows(0)
            # Wide entry surface prints across three readable panels; each repeats request ID.
            ws.fit_to_pages(0, 0)
            ws.set_print_scale(80)
            ws.set_v_pagebreaks([3, 6])
            ws.repeat_columns(0, 0)
        elif name == "History":
            ws.set_column(0, 0, 28)
            ws.set_column(1, 1, 30)
            ws.set_column(2, 2, 100)
            ws.write_row(
                0, 0, ["Event / date", "Request / transition", "Evidence / review note"], dark
            )
            for n, e in enumerate(data["history"], 1):
                ws.write_row(
                    n,
                    0,
                    [
                        e["id"] + "\n" + e["date"],
                        e["request_id"] + "\n" + e["from"] + " → " + e["to"],
                        e["evidence_path"] + "\n" + e["citation"] + "\n" + e["note"],
                    ],
                    body,
                )
                ws.set_row(n, 80)
            if not data["history"]:
                ws.write(1, 2, "No events. No evidence receipt has been asserted.", body)
            ws.repeat_rows(0)
        else:
            ws.set_column(0, 0, 90)
            ws.set_column(1, 1, 80)
            ws.write_row(0, 0, ["Immutable context", "Value"], dark)
            ws.write_row(1, 0, ["State SHA-256", digest(data)], body)
            for n, (p, h) in enumerate(data["source_pins"].items(), 2):
                ws.write_row(n, 0, [p, h], body)
                ws.set_row(n, 38)
            ws.repeat_rows(0)
    wb.close()


def evidence(root, relative, expected, citation):
    p = Path(relative)
    require(
        relative
        and not p.is_absolute()
        and ".." not in p.parts
        and p.parts[0] in {"docs", "industrial", "red_wash", "enterprise", "business"},
        "Unsafe evidence path",
    )
    require(
        not any(s in {"internal", "private", "qa"} or s.startswith(".") for s in p.parts),
        "Evidence must be public",
    )
    target = root / p
    require(
        target.resolve().is_relative_to(root.resolve())
        and not any(x.is_symlink() for x in [target, *target.parents]),
        "Symlink evidence forbidden",
    )
    require(
        target.is_file() and target.suffix.lower() in {".md", ".json", ".csv", ".txt"},
        "Missing evidence or unsupported text type",
    )
    require(sha(target) == expected, "Stale evidence hash")
    text = target.read_text()
    require(
        len(citation.strip()) >= 12 and citation in text, "Citation missing from exact evidence"
    )
    pos = text.index(citation)
    return [text[:pos].count("\n") + 1, text[: pos + len(citation)].count("\n") + 1]


def validate_state(data, root=ROOT):
    base = model(root)
    require(set(data) == set(base), "Unknown state fields")
    for field in ["id", "status", "source_pins", "requests"]:
        require(data[field] == base[field], "Changed source context: " + field)
    current = base["current"].copy()
    seen = set()
    previous = digest(base)
    for e in data["history"]:
        require(
            set(e)
            == {
                "id",
                "request_id",
                "from",
                "to",
                "date",
                "evidence_path",
                "evidence_sha256",
                "citation",
                "lines",
                "note",
                "previous_sha256",
                "sha256",
            },
            "Unknown event fields",
        )
        require(
            e["id"] not in seen and re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_.-]{2,79}", e["id"]),
            "Duplicate or invalid event ID",
        )
        require(e["request_id"] in current, "Unknown request ID")
        require(
            e["from"] == current[e["request_id"]] and e["to"] in STATES and e["to"] != e["from"],
            "Invalid transition",
        )
        require(re.fullmatch(r"\d{4}-\d{2}-\d{2}", e["date"]), "Invalid date")
        datetime.strptime(e["date"], "%Y-%m-%d")
        require(len(e["note"].strip()) >= 12, "Substantive review note required")
        for value in e.values():
            if isinstance(value, str):
                require(
                    not value.lstrip().startswith(("=", "+", "-", "@")), "Formula-like event text"
                )  # noqa: E501 -- reader-facing prose
        if e["to"] in {"incomplete", "received", "disputed"} or e["evidence_path"]:
            require(
                e["lines"]
                == evidence(root, e["evidence_path"], e["evidence_sha256"], e["citation"]),
                "Citation line drift",
            )
        else:
            require(
                not e["evidence_sha256"] and not e["citation"] and e["lines"] == [],
                "Partial evidence fields",
            )
        require(
            e["previous_sha256"] == previous
            and e["sha256"] == digest({k: v for k, v in e.items() if k != "sha256"}),
            "History chain changed",
        )
        previous = e["sha256"]
        seen.add(e["id"])
        current[e["request_id"]] = e["to"]
    require(data["current"] == current, "Current status does not reconcile to history")


def db_write(path, data):
    if path.exists():
        path.unlink()
    with sqlite3.connect(path) as db:
        db.execute("CREATE TABLE records(kind TEXT, id TEXT PRIMARY KEY, payload TEXT NOT NULL)")
        db.execute(
            "INSERT INTO records VALUES (?,?,?)",
            (
                "context",
                data["id"],
                json.dumps(
                    {k: v for k, v in data.items() if k not in {"requests", "history"}},
                    sort_keys=True,
                ),
            ),
        )
        for kind in ["requests", "history"]:
            db.executemany(
                "INSERT INTO records VALUES (?,?,?)",
                [(kind, r["id"], json.dumps(r, sort_keys=True)) for r in data[kind]],
            )


def write_package(out, data):
    out.mkdir(parents=True, exist_ok=True)
    dump(out / "tracker.json", data)
    workbook(data, out / "tracker.xlsx")
    db_write(out / "tracker.sqlite3", data)


def import_workbook(path, output, root=ROOT, state_path=None):
    root = root.resolve()
    output = output.resolve()
    require(
        not output.is_relative_to(root) and not output.exists(),
        "Output must be a new external review folder",
    )
    require(path.resolve() != (root / DEST / "tracker.xlsx").resolve(), "Use a completed copy")
    data = json.loads(state_path.read_text()) if state_path else model(root)
    validate_state(data, root)
    wb = load_workbook(path, data_only=False)
    with tempfile.TemporaryDirectory() as tmp:
        baseline_path = Path(tmp) / "baseline.xlsx"
        workbook(data, baseline_path)
        baseline = load_workbook(baseline_path, data_only=False)
        require(
            wb.sheetnames == baseline.sheetnames and not wb._external_links,
            "Workbook sheets or external links changed",
        )
        for actual, expected in zip(wb.worksheets, baseline.worksheets, strict=True):
            require(
                (actual.max_row, actual.max_column) == (expected.max_row, expected.max_column),
                "Workbook shape changed",
            )
            for row in actual:
                for cell in row:
                    require(
                        cell.data_type != "f" and not cell.hyperlink,
                        "Formula or hyperlink is forbidden",
                    )
                    editable = (
                        actual.title == "Updates" and cell.row > 1 and cell.column in EDITABLE
                    )
                    if not editable:
                        require(
                            cell.value == expected.cell(cell.row, cell.column).value,
                            "Changed workbook context",
                        )
    old_count = len(data["history"])
    for row in wb["Updates"].iter_rows(min_row=2, values_only=True):
        ident, before, after, p, h, citation, note, event_id, date = [
            x if x is not None else "" for x in row
        ]
        if not any(row[2:]):
            continue
        require(
            all(isinstance(x, str) for x in [after, p, h, citation, note, event_id, date]),
            "Enter plain text only",
        )
        e = dict(
            id=event_id,
            request_id=ident,
            **{"from": before, "to": after},
            date=date,
            evidence_path=p,
            evidence_sha256=h,
            citation=citation,
            note=note,
            lines=evidence(root, p, h, citation) if p else [],
            previous_sha256=data["history"][-1]["sha256"]
            if data["history"]
            else digest(model(root)),
        )
        e["sha256"] = digest(e)
        data["history"].append(e)
        data["current"][ident] = after
        validate_state(data, root)
    require(len(data["history"]) > old_count, "No proposed events")
    write_package(output, data)
    dump(
        output / "import-receipt.json",
        dict(
            status=STATUS,
            workbook_sha256=sha(path),
            parent_state_sha256=digest(
                json.loads(state_path.read_text()) if state_path else model(root)
            ),
            added_events=len(data["history"]) - old_count,
        ),
    )
    lines = [
        "# Proposed evidence events",
        "",
        "Review only. No source file, request closure, approval or accounting entry has been applied.",  # noqa: E501 -- reader-facing prose
        "",
    ]
    for e in data["history"][old_count:]:
        lines += [
            f"## {e['id']} — {e['request_id']}",
            "",
            f"{e['date']}: **{e['from']} → {e['to']}**.",
            "",
            e["note"],
            "",
            f"Evidence: `{e['evidence_path']}`; SHA-256 `{e['evidence_sha256']}`; lines {e['lines']}.",  # noqa: E501 -- reader-facing prose
            "",
            f"[Open cited record](<{root / e['evidence_path']}>)"
            if e["evidence_path"]
            else "No file supplied.",
            "",
            "> " + e["citation"].replace("\n", "\n> "),
            "",
        ]
    (output / "PROPOSAL.md").write_text("\n".join(lines))
    return data


def validate(root=ROOT):
    data = json.loads((root / DEST / "tracker.json").read_text())
    validate_state(data, root)
    require(
        data == model(root),
        "Published tracker must remain the unresolved baseline; import is external only",
    )
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp)
        write_package(out, data)
        for name in ["tracker.json", "tracker.xlsx"]:
            require(
                (out / name).read_bytes() == (root / DEST / name).read_bytes(),
                "Stale derivative: " + name,
            )

        def db_rows(path):
            with sqlite3.connect(f"file:{path}?mode=ro", uri=True) as db:
                require(
                    db.execute("PRAGMA integrity_check").fetchall() == [("ok",)], "Corrupt database"
                )
                return (
                    db.execute(
                        "SELECT type,name,tbl_name,sql FROM sqlite_schema ORDER BY name"
                    ).fetchall(),
                    db.execute("SELECT * FROM records ORDER BY id").fetchall(),
                )

        require(
            db_rows(out / "tracker.sqlite3") == db_rows(root / DEST / "tracker.sqlite3"),
            "SQLite contents drift",
        )
    qa = json.loads((root / DEST / "qa/REVIEW.json").read_text())
    require(qa["status"] == "MANUAL_REVIEW_PASS_DRAFT_DESIGN", "Manual workbook review pending")
    require(qa["workbook_sha256"] == sha(root / DEST / "tracker.xlsx"), "Stale workbook QA")
    require(len(qa["pages"]) == 11 and len(qa["sheets"]) == 5, "Incomplete native QA coverage")
    for page in qa["pages"]:
        require(sha(root / DEST / "qa" / page["path"]) == page["sha256"], "QA image hash drift")
    print("PASS: 14 source-pinned unresolved requests; complete workbook/SQLite reconciliation")


def render(root=ROOT):
    """Render every native worksheet; manual inspection must follow this capture."""
    import subprocess

    import fitz

    qa = root / DEST / "qa"
    qa.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        temp = Path(tmp)
        subprocess.run(
            [
                "libreoffice",
                "-env:UserInstallation=" + (temp / "profile").as_uri(),
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(temp),
                str(root / DEST / "tracker.xlsx"),
            ],
            check=True,
            capture_output=True,
            text=True,
        )
        pdf = fitz.open(temp / "tracker.pdf")
        all_text = "".join(page.get_text() for page in pdf)

        def normalize(value):
            return "".join(str(value).split())

        native = load_workbook(root / DEST / "tracker.xlsx", data_only=False)
        cells = [
            cell.value
            for sheet in native
            for row in sheet
            for cell in row
            if cell.value is not None
        ]
        missing = [v for v in cells if normalize(v) not in normalize(all_text)]
        require(not missing, "Native PDF lost cell text: " + repr(missing[:3]))
        pages = []
        for n, page in enumerate(pdf, 1):
            path = qa / f"page-{n:02}.png"
            page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5)).save(path)
            pages.append(dict(page=n, path=path.name, sha256=sha(path)))
        dump(
            qa / "REVIEW.json",
            dict(
                status="RENDERED_PENDING_MANUAL_INSPECTION",
                workbook_sha256=sha(root / DEST / "tracker.xlsx"),
                sheets=native.sheetnames,
                pages=pages,
                populated_cells=len(cells),
                all_cell_text_present=True,
                renderer="LibreOffice native PDF; PyMuPDF 1.5x PNG",
                corrections=[],
            ),
        )
    print(
        f"Captured {len(pages)} pages across {len(native.sheetnames)} sheets; manual review needed"
    )


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("command", choices=["build", "validate", "import", "render"])
    parser.add_argument("--workbook", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--state", type=Path)
    args = parser.parse_args()
    if args.command == "build":
        write_package(ROOT / DEST, model())
    elif args.command == "validate":
        validate()
    elif args.command == "render":
        render()
    else:
        require(args.workbook and args.output, "Import requires --workbook and --output")
        data = import_workbook(args.workbook, args.output, state_path=args.state)
        print(f"PASS: {len(data['history'])} retained events; proposal only at {args.output}")


if __name__ == "__main__":
    main()
