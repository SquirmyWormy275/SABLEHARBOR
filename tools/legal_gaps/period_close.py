"""One pinned ARU_GROUP forecast month; source-derived public close exercise."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import sqlite3
import subprocess
import tempfile
from collections import defaultdict
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import xlsxwriter

ROOT = Path(__file__).resolve().parents[2]
DEST = Path("docs/legal/gap-instruments/period-close")
STATUS = "DRAFT_FOR_EXACT_FILE_REVIEW"
D = Decimal


def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def dump(p, value):
    p.write_text(json.dumps(value, indent=2) + "\n")


def model(root=ROOT):
    pins = json.loads((root / DEST / "inputs.json").read_text())
    tables = {}
    for s in pins["sources"]:
        p = root / s["path"]
        if sha(p) != s["sha256"]:
            raise ValueError("Source hash changed: " + str(p))
        rows = list(csv.DictReader(p.open()))
        if len(rows) != s["row_count"] or any(
            any(r[k] != v for k, v in s["selector"].items()) for r in rows
        ):
            raise ValueError("Source scope/coverage changed")
        if "upstream_path" in s and sha(root / s["upstream_path"]) != s["upstream_sha256"]:
            raise ValueError("Upstream source changed")
        tables[s["id"]] = rows
    opening = {r["account"]: D(r["signed_usd"]) for r in tables["opening_balances"]}
    closing = {r["account"]: D(r["signed_usd"]) for r in tables["trial_balances"]}
    movements = defaultdict(Decimal)
    journals = defaultdict(Decimal)
    for r in tables["journal"]:
        if D(r["signed_usd"]) != D(r["debit_usd"]) - D(r["credit_usd"]):
            raise ValueError("Journal signed amount differs")
        movements[r["account"]] += D(r["signed_usd"])
        journals[r["journal_id"]] += D(r["signed_usd"])
    if any(journals.values()) or sum(opening.values()) or sum(closing.values()):
        raise ValueError("Unbalanced opening, journal or closing book")
    if (set(opening) | set(movements)) - set(closing):
        raise ValueError("Closing account coverage incomplete")
    tb = []
    for r in tables["trial_balances"]:
        a = r["account"]
        begin = opening.get(a, D(0))
        move = movements[a]
        if begin + move != closing[a]:
            raise ValueError("Account rollforward differs: " + a)
        tb.append(
            dict(
                account=a,
                name=r["account_name"],
                account_type=r["account_type"],
                opening=str(begin),
                activity=str(move),
                adjustment="0",
                closing=str(closing[a]),
                difference="0",
            )
        )
    if len(opening) != len(tables["opening_balances"]) or len(closing) != len(
        tables["trial_balances"]
    ):
        raise ValueError("Duplicate balance account")
    native_cash = sorted(
        (r["journal_id"], r["source_id"], r["signed_usd"])
        for r in tables["journal"]
        if r["account"] == "1000"
    )
    bank_cash = sorted(
        (r["journal_id"], r["source_id"], r["signed_cash_usd"])
        for r in tables["industrial_bank_transactions"]
    )
    if native_cash != bank_cash:
        raise ValueError("Bank event identity or amount differs from complete cash journal")
    checks = []

    def check(name, left, right, basis):
        diff = left - right
        if abs(diff) > D(".02"):
            raise ValueError(f"{name}: {diff}")
        checks.append(
            dict(name=name, left=str(left), right=str(right), difference=str(diff), basis=basis)
        )

    bank = tables["industrial_bank_reconciliations"][0]
    check(
        "Bank adjusted cash to ledger",
        D(bank["adjusted_bank_cash_usd"]),
        closing["1000"],
        "Modeled clearing schedule; not independent bank evidence",
    )
    check(
        "Bank clearing timing",
        D(bank["closing_bank_cash_usd"])
        + D(bank["deposits_in_transit_usd"])
        - D(bank["outstanding_payments_usd"]),
        closing["1000"],
        "Opening bank cash + cleared receipts - cleared payments + deposits - outstanding payments",
    )
    check(
        "Complete bank transaction population",
        sum(D(r["signed_cash_usd"]) for r in tables["industrial_bank_transactions"]),
        movements["1000"],
        "All 33 selected modeled cash events; journal_id joins activity",
    )
    for a, label in [("1100", "Receivables"), ("2000", "Payables")]:
        check(
            label + " opening plus activity",
            opening[a] + movements[a],
            closing[a],
            "Complete account journal rollforward; independent January open-item schedule unavailable in this case",  # noqa: E501 -- reader-facing prose
        )
    debt = tables["debt"][0]
    for suffix, account, paid, draw in [
        ("legacy_term", "2400", "375000", "0"),
        ("replacement_term", "2410", "0", debt["replacement_debt_draw_usd"]),
        ("lease", "2600", debt["scheduled_lease_principal_usd"], "0"),
    ]:
        paid = debt["principal_cash_paid_" + suffix + "_usd"]
        check(
            "Debt " + suffix,
            D(debt["opening_" + suffix + "_usd"]) + D(draw) - D(paid),
            -closing[account],
            "Native debt row; no contractual compliance conclusion",
        )
    assets = tables["assets"]
    check(
        "PPE gross to ledger",
        sum(D(r["gross_usd"]) for r in assets),
        closing["1400"],
        "All eight native PPE rows; finance-lease ROU is separate",
    )
    check(
        "PPE accumulated depreciation",
        sum(D(r["accumulated_depreciation_usd"]) for r in assets),
        -closing["1490"],
        "All eight native PPE rows",
    )
    check(
        "PPE net rollforward",
        sum(
            D(r["opening_net_book_usd"]) + D(r["additions_usd"]) - D(r["depreciation_usd"])
            for r in assets
        ),
        closing["1400"] + closing["1490"],
        "Opening NBV plus additions less depreciation",
    )
    stmt = tables["monthly_statements"][0]
    for kind, field, sign in [("asset", "assets_usd", 1), ("liability", "liabilities_usd", -1)]:
        check(
            field,
            sum(D(t["closing"]) for t in tb if t["account_type"] == kind) * sign,
            D(stmt[field]),
            "Complete closing account classification",
        )
    income = -sum(D(t["closing"]) for t in tb if t["account_type"] in ["revenue", "expense"])
    check(
        "Net income",
        income,
        D(stmt["net_income_usd"]),
        "January has no prior-year-to-date forecast income",
    )
    equity = -sum(D(t["closing"]) for t in tb if t["account_type"] == "equity") + income
    check(
        "Equity including current income",
        equity,
        D(stmt["equity_including_current_income_usd"]),
        "Opening equity and January income",
    )
    for flow, field in [
        ("OPERATING", "operating_cash_flow_usd"),
        ("INVESTING", "investing_cash_flow_usd"),
        ("FINANCING", "financing_cash_flow_usd"),
    ]:
        check(
            field,
            sum(
                D(r["signed_usd"])
                for r in tables["journal"]
                if r["account"] == "1000" and r["cash_flow"] == flow
            ),
            D(stmt[field]),
            "Complete cash-account journal classification",
        )
    return dict(
        id="SH-CLOSE-ARU-2027-01",
        status=STATUS,
        scope="ARU_GROUP industrial reporting book; base scenario; January 2027",
        fact_state="CONDITIONAL_FORECAST",
        source_pins=pins,
        tables=tables,
        trial_balance=tb,
        checks=checks,
        proposed_adjustments=[],
        adjustment_basis="No unexplained numerical difference in the selected records. Zero new entries; missing independent evidence does not justify a plug.",  # noqa: E501 -- reader-facing prose
        limitations=[
            "ARU_GROUP is the native industrial reporting scope, not a claim of a single statutory legal entity. Do not add Core, enterprise replacement or allocated legal books.",  # noqa: E501 -- reader-facing prose
            "Bank records are modeled clearing events. Obtain a bank-issued statement and subsequent clearing evidence before concluding independent cash existence.",  # noqa: E501 -- reader-facing prose
            "Account 1100 is external receivables net of allowance. This case supplies complete ledger movement, not independent gross aging or allowance adequacy evidence.",  # noqa: E501 -- reader-facing prose
            "Account 2000 is trade payables including forecast settlement flows. Obtain a dated supplier open-item schedule and subsequent disbursements; terminal invoice masters are not month-end balances.",  # noqa: E501 -- reader-facing prose
            "PPE schedule covers accounts 1400/1490; finance-lease ROU accounts 1500/1590 remain separately visible in the complete trial balance. Physical existence, title, useful lives and executed debt terms require additional evidence.",  # noqa: E501 -- reader-facing prose
            "These are public source-derived learning examples, not private assessment keys, audited results or a completed corporate close.",  # noqa: E501 -- reader-facing prose
        ],
    )


def workbook(data, path, worked):
    wb = xlsxwriter.Workbook(path)
    wb.set_properties(
        {
            "title": "ARU Group | January 2027 close",
            "author": "Sable Harbor",
            "created": datetime(2026, 9, 13),
        }
    )
    title = wb.add_format(
        {"bold": True, "font_size": 18, "font_color": "#FFFFFF", "bg_color": "#192D35"}
    )
    text = wb.add_format(
        {
            "font_name": "Calibri",
            "font_size": 10,
            "text_wrap": True,
            "valign": "vcenter",
            "bottom": 1,
            "bottom_color": "#DDE2E4",
        }
    )
    head = wb.add_format(
        {
            "bold": True,
            "font_size": 10,
            "bg_color": "#E3E9EA",
            "text_wrap": True,
            "valign": "vcenter",
        }
    )
    num = wb.add_format(
        {
            "font_size": 10,
            "num_format": "#,##0.00;[Red](#,##0.00);–",
            "bottom": 1,
            "bottom_color": "#DDE2E4",
        }
    )
    blank = wb.add_format(
        {
            "bg_color": "#FFF4D7",
            "border": 1,
            "border_color": "#D9C690",
            "num_format": "#,##0.00;[Red](#,##0.00);–",
        }
    )

    def sheet(name, headers, widths, rows):
        ws = wb.add_worksheet(name)
        ws.hide_gridlines(2)
        ws.set_landscape()
        ws.set_paper(9)
        ws.fit_to_pages(1, 0)
        ws.set_margins(0.3, 0.3, 0.45, 0.45)
        ws.set_header("&LARU GROUP | JANUARY 2027&R" + ("WORKED" if worked else "BLANK"))
        ws.set_footer("&LDRAFT • Public forecast exercise&RPage &P of &N")
        ws.merge_range(0, 0, 1, len(headers) - 1, name, title)
        for c, width in enumerate(widths):
            ws.set_column(c, c, width)
        ws.write_row(3, 0, headers, head)
        ws.set_row(3, 32)
        ws.repeat_rows(0, 3)
        ws.freeze_panes(4, 0)
        for ri, row in enumerate(rows, 4):
            ws.set_row(ri, 32)
            for ci, value in enumerate(row):
                if isinstance(value, (int, float, Decimal)):
                    ws.write_number(ri, ci, float(value), num)
                else:
                    ws.write(ri, ci, value, text)
        ws.print_area(0, 0, max(4, len(rows) + 3), len(headers) - 1)
        return ws

    intro = [
        ["Scope", data["scope"]],
        [
            "Task",
            "Roll opening balances through all January activity; prepare five reconciliations, assess adjustments, then complete closing TB and statements.",  # noqa: E501 -- reader-facing prose
        ],
        [
            "Start",
            "Opening and Activity are source inputs. Yellow cells are learner responses; worked version contains formulas.",  # noqa: E501 -- reader-facing prose
        ],
        [
            "Evidence",
            "All original columns and rows are preserved in source/*.csv and close.sqlite3. Workbook rows are a readable projection, not additional transactions.",  # noqa: E501 -- reader-facing prose
        ],
        [
            "Sign",
            "Positive balances are debits. Liabilities and equity are credit-negative in the TB; statements show positive liability/equity balances.",  # noqa: E501 -- reader-facing prose
        ],
        ["Adjustment rule", data["adjustment_basis"]],
    ] + [["Limit " + str(i + 1), v] for i, v in enumerate(data["limitations"])]
    ws = sheet("Read first", ["Item", "Instruction / boundary"], [22, 115], intro)
    for r in range(4, 4 + len(intro)):
        ws.set_row(r, 38)
    sheet(
        "Opening",
        ["Account", "Name", "Opening USD"],
        [14, 75, 22],
        [
            [r["account"], r["account_name"], D(r["signed_usd"])]
            for r in data["tables"]["opening_balances"]
        ],
    )
    jr = data["tables"]["journal"]
    n = len(jr) + 4
    sheet(
        "Activity",
        [
            "Journal ID",
            "Account",
            "Source ID",
            "Debit USD",
            "Credit USD",
            "Signed USD",
            "Cash flow",
        ],
        [36, 10, 48, 17, 17, 17, 23],
        [
            [
                r["journal_id"],
                r["account"],
                r["source_id"],
                D(r["debit_usd"]),
                D(r["credit_usd"]),
                D(r["signed_usd"]),
                r["cash_flow"],
            ]
            for r in jr
        ],
    )
    sheet(
        "Bank events",
        ["Journal ID", "Posting date", "Clearing date", "Bank reference", "Signed cash USD"],
        [40, 18, 18, 38, 22],
        [
            [
                r["journal_id"],
                r["posting_date"],
                r["clearing_date"],
                r["bank_reference"],
                D(r["signed_cash_usd"]),
            ]
            for r in data["tables"]["industrial_bank_transactions"]
        ],
    )
    bank = data["tables"]["industrial_bank_reconciliations"][0]
    bf = [
        "opening_bank_cash_usd",
        "cleared_receipts_usd",
        "cleared_payments_usd",
        "closing_bank_cash_usd",
        "deposits_in_transit_usd",
        "outstanding_payments_usd",
        "adjusted_bank_cash_usd",
        "ledger_cash_usd",
    ]
    sheet(
        "Bank inputs",
        ["Native field", "Source USD"],
        [82, 28],
        [[f.replace("_", " ").capitalize(), D(bank[f])] for f in bf],
    )
    for account, name in [("1100", "Receivables"), ("2000", "Payables")]:
        lines = [r for r in jr if r["account"] == account]
        sheet(
            name,
            ["Journal ID", "Source ID", "Debit USD", "Credit USD", "Signed USD"],
            [42, 65, 22, 22, 22],
            [
                [
                    r["journal_id"],
                    r["source_id"],
                    D(r["debit_usd"]),
                    D(r["credit_usd"]),
                    D(r["signed_usd"]),
                ]
                for r in lines
            ],
        )
    summaries = []
    for c in data["checks"]:
        summaries.append(
            [c["name"].replace("_", " "), D(c["left"]), D(c["right"]), None, c["basis"]]
        )
    ws = sheet(
        "Five reconciliations",
        ["Check", "Calculated USD", "Comparison USD", "Difference USD", "Evidence boundary"],
        [36, 22, 22, 22, 63],
        summaries,
    )
    for i, c in enumerate(data["checks"], 4):
        if worked:
            # Recompute each left side from the reader-visible native input sheets.
            formulas = [
                "='Bank inputs'!B11",
                "='Bank inputs'!B8+'Bank inputs'!B9-'Bank inputs'!B10",
                "=SUM('Bank events'!E5:E37)",
                '=SUMIF(Opening!A5:A26,"1100",Opening!C5:C26)+SUM(Receivables!E5:E34)',
                '=SUMIF(Opening!A5:A26,"2000",Opening!C5:C26)+SUM(Payables!E5:E77)',
                "=Debt!B5+Debt!C5-Debt!D5",
                "=Debt!B6+Debt!C6-Debt!D6",
                "=Debt!B7+Debt!C7-Debt!D7",
                "=SUM('PPE detail'!B5:B12)",
                "=SUM('PPE detail'!C5:C12)",
                "=SUM(PPE!B5:C12)-SUM(PPE!D5:D12)",
                "=Statements!B5",
                "=Statements!B6",
                "=Statements!B7",
                "=Statements!B8",
                "=Statements!B11",
                "=Statements!B12",
                "=Statements!B13",
            ]
            ws.write_formula(i, 1, formulas[i - 4], num, float(c["left"]))
            ws.write_formula(i, 3, f"=B{i + 1}-C{i + 1}", num, float(c["difference"]))
        else:
            ws.write_blank(i, 1, None, blank)
            ws.write_blank(i, 3, None, blank)
    ws = sheet(
        "Debt",
        [
            "Instrument",
            "Opening USD",
            "Draw USD",
            "Principal paid USD",
            "Closing source USD",
            "Calculated closing USD",
        ],
        [30, 22, 20, 24, 24, 26],
        [],
    )
    d = data["tables"]["debt"][0]
    for i, s in enumerate(["legacy_term", "replacement_term", "lease"], 4):
        ws.set_row(i, 32)
        vals = [
            s.replace("_", " ").title(),
            D(d["opening_" + s + "_usd"]),
            D(d["replacement_debt_draw_usd"]) if s == "replacement_term" else D(0),
            D(d["principal_cash_paid_" + s + "_usd"]),
            D(d["closing_" + s + "_usd"]),
        ]
        for j, v in enumerate(vals):
            ws.write(i, j, float(v) if isinstance(v, Decimal) else v, num if j else text)
        if worked:
            ws.write_formula(i, 5, f"=B{i + 1}+C{i + 1}-D{i + 1}", num, float(vals[4]))
        else:
            ws.write_blank(i, 5, None, blank)
    ws.print_area(0, 0, 7, 5)
    ars = data["tables"]["assets"]
    sheet(
        "PPE detail",
        ["Asset ID", "Gross USD", "Accumulated depreciation USD", "Ledger account", "Status"],
        [48, 24, 28, 18, 38],
        [
            [
                r["asset_id"],
                D(r["gross_usd"]),
                D(r["accumulated_depreciation_usd"]),
                r["ledger_account"],
                r["asset_status"],
            ]
            for r in ars
        ],
    )
    ws = sheet(
        "PPE",
        [
            "Asset ID",
            "Opening NBV USD",
            "Additions USD",
            "Depreciation USD",
            "Closing source USD",
            "Calculated closing USD",
        ],
        [46, 22, 22, 22, 22, 26],
        [
            [
                r["asset_id"],
                D(r["opening_net_book_usd"]),
                D(r["additions_usd"]),
                D(r["depreciation_usd"]),
                D(r["closing_net_book_usd"]),
                None,
            ]
            for r in ars
        ],
    )
    for i, r in enumerate(ars, 4):
        if worked:
            ws.write_formula(
                i, 5, f"=B{i + 1}+C{i + 1}-D{i + 1}", num, float(r["closing_net_book_usd"])
            )
        else:
            ws.write_blank(i, 5, None, blank)
    ws = sheet(
        "Adjustments",
        ["Account", "Proposed debit USD", "Proposed credit USD", "Evidence / reason"],
        [20, 28, 28, 75],
        [
            [
                "None supported",
                0 if worked else None,
                0 if worked else None,
                data["adjustment_basis"],
            ]
        ],
    )
    if not worked:
        ws.write_blank(4, 1, None, blank)
        ws.write_blank(4, 2, None, blank)
    tb = data["trial_balance"]
    ws = sheet(
        "Closing TB",
        [
            "Account",
            "Name",
            "Opening USD",
            "Activity USD",
            "Adjustment USD",
            "Closing USD",
            "Released closing USD",
            "Difference USD",
        ],
        [12, 43, 19, 19, 19, 19, 22, 19],
        [[r["account"], r["name"], None, None, None, None, D(r["closing"]), None] for r in tb],
    )
    for i, r in enumerate(tb, 4):
        row = i + 1
        formulas = {
            2: (f"=SUMIF(Opening!A5:A26,A{row},Opening!C5:C26)", r["opening"]),
            3: (f"=SUMIF(Activity!B5:B{n},A{row},Activity!F5:F{n})", r["activity"]),
            4: ("=0", "0"),
            5: (f"=SUM(C{row}:E{row})", r["closing"]),
            7: (f"=F{row}-G{row}", "0"),
        }
        for c, (f, v) in formulas.items():
            if worked:
                ws.write_formula(i, c, f, num, float(v))
            else:
                ws.write_blank(i, c, None, blank)
    for sheet_name, kinds in [
        ("Income statement", {"revenue", "expense"}),
        ("Balance sheet", {"asset", "liability", "equity"}),
    ]:
        selected = [(j, r) for j, r in enumerate(tb) if r["account_type"] in kinds]
        ws = sheet(
            sheet_name,
            ["Account", "Description", "Class", "Amount USD"],
            [14, 70, 22, 26],
            [[r["account"], r["name"], r["account_type"], None] for j, r in selected],
        )
        for i, (j, r) in enumerate(selected, 4):
            sign = -1 if r["account_type"] in {"liability", "equity", "revenue"} else 1
            if worked:
                ws.write_formula(
                    i, 3, f"='Closing TB'!F{j + 5}*{sign}", num, float(D(r["closing"]) * sign)
                )
            else:
                ws.write_blank(i, 3, None, blank)
        end = len(selected) + 4
        ws.set_row(end, 32)
        ws.write(
            end,
            1,
            "Net income" if sheet_name == "Income statement" else "Current-period income in equity",
            head,
        )
        if worked:
            ws.write_formula(
                end,
                3,
                "=Statements!B7",
                num,
                float(data["tables"]["monthly_statements"][0]["net_income_usd"]),
            )
        else:
            ws.write_blank(end, 3, None, blank)
        ws.print_area(0, 0, end, 3)
    ws = sheet(
        "Statements",
        ["Measure", "Calculated USD", "Released USD", "Difference USD"],
        [62, 26, 26, 26],
        [],
    )
    measures = [
        ("Assets", "assets_usd", "asset"),
        ("Liabilities", "liabilities_usd", "liability"),
        ("Net income", "net_income_usd", "income"),
        ("Equity including income", "equity_including_current_income_usd", "equity"),
        ("Opening cash", "opening_year_cash_usd", "opening_cash"),
        ("Ending cash", "ending_cash_usd", "cash"),
        ("Operating cash flow", "operating_cash_flow_usd", "OPERATING"),
        ("Investing cash flow", "investing_cash_flow_usd", "INVESTING"),
        ("Financing cash flow", "financing_cash_flow_usd", "FINANCING"),
    ]
    stmt = data["tables"]["monthly_statements"][0]
    for i, (label, field, kind) in enumerate(measures, 4):
        ws.set_row(i, 32)
        ws.write(i, 0, label, text)
        ws.write_number(i, 2, float(stmt[field]), num)
        indices = [
            j + 5
            for j, r in enumerate(tb)
            if r["account_type"] == kind
            or (kind == "income" and r["account_type"] in ["revenue", "expense"])
        ]
        f = (
            "="
            + ("-" if kind in ["liability", "income", "equity"] else "")
            + "SUM("
            + ",".join(f"'Closing TB'!F{j}" for j in indices)
            + ")"
        )
        if kind == "equity":
            f += " + B7"
        if kind in ["opening_cash", "cash"]:
            f = "=Opening!C5" if kind == "opening_cash" else "='Closing TB'!F5"
        if kind in ["OPERATING", "INVESTING", "FINANCING"]:
            f = f'=SUMIFS(Activity!F5:F{n},Activity!B5:B{n},"1000",Activity!G5:G{n},"{kind}")'
        if worked:
            ws.write_formula(i, 1, f, num, float(stmt[field]))
            ws.write_formula(i, 3, f"=B{i + 1}-C{i + 1}", num, 0)
        else:
            ws.write_blank(i, 1, None, blank)
            ws.write_blank(i, 3, None, blank)
    ws.print_area(0, 0, 12, 3)
    sheet(
        "Source index",
        ["Table", "Rows", "Source member / upstream"],
        [35, 12, 95],
        [
            [s["id"], s["row_count"], s.get("release_member", s.get("upstream_path"))]
            for s in data["source_pins"]["sources"]
        ],
    )
    wb.close()


def generate(root=ROOT, out=None):
    out = out or root / DEST
    out.mkdir(parents=True, exist_ok=True)
    data = model(root)
    dump(out / "case.json", data)
    for mode in ["blank", "worked"]:
        workbook(data, out / (mode + ".xlsx"), mode == "worked")
    p = out / "close.sqlite3"
    if p.exists():
        p.unlink()
    with sqlite3.connect(p) as db:
        db.execute("CREATE TABLE case_record(id TEXT PRIMARY KEY, payload TEXT NOT NULL)")
        db.execute(
            "INSERT INTO case_record VALUES(?,?)",
            (
                data["id"],
                json.dumps({k: v for k, v in data.items() if k != "tables"}, sort_keys=True),
            ),
        )
        for name, rows in data["tables"].items():
            columns = list(rows[0])
            db.execute(
                'CREATE TABLE "'
                + name
                + '" ('
                + ",".join('"' + c + '" TEXT NOT NULL' for c in columns)
                + ")"
            )
            db.executemany(
                'INSERT INTO "' + name + '" VALUES (' + ",".join("?" for _ in columns) + ")",
                [[r[c] for c in columns] for r in rows],
            )
    with (out / "closing_trial_balance.csv").open("w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(data["trial_balance"][0]))
        w.writeheader()
        w.writerows(data["trial_balance"])
    intro = "# ARU Group — January 2027 period close\n\n**SH-CLOSE-ARU-2027-01 · DRAFT FOR EXACT-FILE REVIEW · Conditional forecast, base scenario.**\n\n"  # noqa: E501 -- reader-facing prose
    intro += "Reconstruct one complete native industrial reporting book: 22 opening balances, all 274 January activity rows, five reconciliation areas, no new adjustment entries, and all 34 closing accounts. ARU_GROUP is a reporting scope, not an assertion of one statutory company.\n\nStart with [blank.xlsx](blank.xlsx). Work through Opening → Activity → Bank / receivables / payables / Debt / PPE → Adjustments → Closing TB → Statements. Compare [worked.xlsx](worked.xlsx) and [WORKED.md](WORKED.md) afterward. Yellow cells are responses; existing source values remain visible.\n\n"  # noqa: E501 -- reader-facing prose
    intro += "## Evidence and limits\n\nThe [input register](inputs.json) pins the immutable business-operations-v1.0.0 release and exact member hashes. All selected native columns survive in the eight source CSVs listed below and [SQLite tables](close.sqlite3). [case.json](case.json) preserves checks and source rows. The [closing TB CSV](closing_trial_balance.csv) is a derivative. Month 0 is used once as opening balances; it is excluded from activity.\n\n"  # noqa: E501 -- reader-facing prose
    intro += "\n".join("- " + x for x in data["limitations"]) + "\n\n"
    intro += "## Reader deliverable\n\nSubmit the completed workbook, an explanation for each difference, any evidence-supported proposed journal, and requests for missing independent support. A zero source tie-out is not a signed close or audit conclusion. The public worked example leaves all proposed adjustments at zero because the supplied records contain no unexplained numerical discrepancy.\n\n"  # noqa: E501 -- reader-facing prose
    intro += "## Reproduce\n\n`python tools/legal_gaps/period_close.py build` regenerates the derivatives. `validate` checks source hashes, complete row populations, per-journal balance, account rollforwards, 18 reconciliations, workbook bytes and full SQLite content. `render` creates native LibreOffice review pages; manual review must be recorded separately.\n"  # noqa: E501 -- reader-facing prose
    intro += "\n## Individual native files\n\n"
    for source in data["source_pins"]["sources"]:
        name = source["id"]
        intro += (
            f"- [{name.replace('_', ' ').capitalize()}](source/{name}.csv): "
            f"{source['row_count']} rows.\n"
        )
    intro += "\nVerify original extraction with `python tools/legal_gaps/period_close.py verify-archive --archive /path/to/sable-harbor-business-operations-v1.0.0.zip`. The command compares every selected column and row with the hash-pinned archive. `verify-formulas` independently recalculates all 270 worked spreadsheet formulas in LibreOffice. `validate-qa` checks both exact workbook hashes and all retained manual-review pages.\n"  # noqa: E501 -- reader-facing prose
    (out / "README.md").write_text(intro)
    worked = "# Public worked close — ARU Group, January 2027\n\nNo new economic facts or correcting entries are introduced. All amounts are USD; liabilities and equity are credit-negative in the trial balance.\n\n| Account | Opening | Activity | Adjustment | Closing |\n|---|---:|---:|---:|---:|\n"  # noqa: E501 -- reader-facing prose
    for r in data["trial_balance"]:
        worked += (
            "| "
            + r["account"]
            + " "
            + r["name"]
            + " | "
            + " | ".join(
                f"{D(r[k]):,.2f}" for k in ["opening", "activity", "adjustment", "closing"]
            )
            + " |\n"
        )
    worked += "\n## Reconciliation results\n\n"
    for c in data["checks"]:
        worked += f"- {c['name']}: {D(c['left']):,.2f} versus {D(c['right']):,.2f}; difference {D(c['difference']):,.2f}. {c['basis']}.\n"  # noqa: E501 -- reader-facing prose
    worked += (
        "\n## Adjustment and conclusion\n\n"
        + data["adjustment_basis"]
        + " Closing cash is $3,494,528; assets $86,966,303; liabilities $27,952,890; equity including current income $59,013,413; January net loss $82,694. Independent cash, AR/AP completeness, asset existence and executed debt support remain unproven.\n"  # noqa: E501 -- reader-facing prose
    )
    (out / "WORKED.md").write_text(worked)
    names = [
        "README.md",
        "WORKED.md",
        "case.json",
        "blank.xlsx",
        "worked.xlsx",
        "close.sqlite3",
        "closing_trial_balance.csv",
    ]
    dump(
        out / "manifest.json",
        dict(
            status=STATUS,
            inputs_sha256=sha(root / DEST / "inputs.json"),
            files={n: sha(out / n) for n in names},
        ),
    )
    return data


def verify_archive(archive, root=ROOT):
    """Prove the committed selection equals the immutable release, all columns and rows."""
    import io
    import zipfile

    pins = json.loads((root / DEST / "inputs.json").read_text())
    if sha(archive) != pins["release_sha256"]:
        raise ValueError("Release archive hash differs")
    with zipfile.ZipFile(archive) as z:
        for source in pins["sources"]:
            if "release_member" not in source:
                continue
            raw = z.read(source["release_member"])
            if hashlib.sha256(raw).hexdigest() != source["release_member_sha256"]:
                raise ValueError("Release member hash differs")
            complete = list(csv.DictReader(io.StringIO(raw.decode())))
            selected = [
                r for r in complete if all(r[k] == v for k, v in source["selector"].items())
            ]
            saved = list(csv.DictReader((root / source["path"]).open()))
            if selected != saved or len(complete) != source["full_member_rows"]:
                raise ValueError("Release selected row/column population differs")
    print("PASS: exact immutable archive selection, all six forecast members")


def db_content(p):
    with sqlite3.connect(f"file:{p}?mode=ro", uri=True) as db:
        if db.execute("PRAGMA integrity_check").fetchall() != [("ok",)]:
            raise ValueError("SQLite integrity")
        schema = db.execute(
            "SELECT type,name,tbl_name,sql FROM sqlite_schema ORDER BY type,name"
        ).fetchall()
        rows = {
            n: sorted(db.execute('SELECT * FROM "' + n + '"').fetchall())
            for (n,) in db.execute("SELECT name FROM sqlite_schema WHERE type='table'").fetchall()
        }
        return (
            schema,
            rows,
            db.execute("PRAGMA user_version").fetchall(),
            db.execute("PRAGMA application_id").fetchall(),
        )


def validate(root=ROOT):
    saved = root / DEST
    manifest = json.loads((saved / "manifest.json").read_text())
    with tempfile.TemporaryDirectory() as t:
        tmp = Path(t)
        generate(root, tmp)
        expected = json.loads((tmp / "manifest.json").read_text())
        if (
            manifest["status"] != STATUS
            or manifest["inputs_sha256"] != expected["inputs_sha256"]
            or set(manifest["files"]) != set(expected["files"])
        ):
            raise ValueError("Manifest coverage drift")
        for n, h in manifest["files"].items():
            if sha(saved / n) != h:
                raise ValueError("Saved artifact drift " + n)
            if n.endswith(".sqlite3"):
                if db_content(saved / n) != db_content(tmp / n):
                    raise ValueError("Database contents drift")
            elif (saved / n).read_bytes() != (tmp / n).read_bytes():
                raise ValueError("Generated artifact drift " + n)
    print(
        "PASS: complete ARU January book, 18 checks, pinned rows, deterministic workbooks and SQLite contents"  # noqa: E501 -- reader-facing prose
    )


def validate_qa(root=ROOT):
    import openpyxl

    target = root / DEST
    receipt = json.loads((target / "qa/REVIEW.json").read_text())
    if receipt["status"] != "MANUAL_REVIEW_PASS_DRAFT_DESIGN":
        raise ValueError("Manual review pending")
    if {r["mode"] for r in receipt["records"]} != {"blank", "worked"}:
        raise ValueError("Workbook QA coverage differs")
    for record in receipt["records"]:
        path = target / (record["mode"] + ".xlsx")
        if sha(path) != record["workbook_sha256"]:
            raise ValueError("Reviewed workbook changed")
        if record["worksheets"] != openpyxl.load_workbook(path).sheetnames:
            raise ValueError("Worksheet QA coverage differs")
        if len(record["pages"]) != 39:
            raise ValueError("Page QA coverage differs")
        for page in record["pages"]:
            if sha(root / page["path"]) != page["sha256"]:
                raise ValueError("Reviewed image changed")
    print("PASS: two exact workbooks, 34 worksheets, 78 manually reviewed native pages")


def verify_formulas(root=ROOT):
    """Force a native spreadsheet recalculation and compare every formula result."""
    import openpyxl

    path = root / DEST / "worked.xlsx"
    formulas = openpyxl.load_workbook(path, data_only=False)
    cached = openpyxl.load_workbook(path, data_only=True)
    with tempfile.TemporaryDirectory() as t:
        subprocess.run(
            [
                "libreoffice",
                "-env:UserInstallation=" + Path(t, "profile").as_uri(),
                "--headless",
                "--convert-to",
                "xlsx",
                "--outdir",
                t,
                str(path),
            ],
            check=True,
            capture_output=True,
        )
        actual = openpyxl.load_workbook(Path(t, "worked.xlsx"), data_only=True)
        checked = 0
        for ws in formulas:
            for row in ws:
                for cell in row:
                    if cell.data_type != "f":
                        continue
                    expected = cached[ws.title][cell.coordinate].value
                    value = actual[ws.title][cell.coordinate].value
                    if not isinstance(value, (int, float)) or abs(
                        D(str(value)) - D(str(expected))
                    ) > D(".02"):
                        raise ValueError(
                            f"Native formula result differs: {ws.title}!{cell.coordinate}: "
                            f"{value} / {expected}"
                        )
                    checked += 1
    print(f"PASS: {checked} formulas independently recalculated by LibreOffice")
    return checked


def render(root=ROOT):
    import fitz
    from PIL import Image

    out = root / DEST / "qa"
    out.mkdir(exist_ok=True)
    records = []
    with tempfile.TemporaryDirectory() as t:
        for mode in ["blank", "worked"]:
            path = root / DEST / (mode + ".xlsx")
            subprocess.run(
                [
                    "libreoffice",
                    "-env:UserInstallation=" + Path(t, "profile").as_uri(),
                    "--headless",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    t,
                    str(path),
                ],
                check=True,
                capture_output=True,
            )
            doc = fitz.open(Path(t, mode + ".pdf"))
            pages = []
            for i, page in enumerate(doc):
                pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5))
                p = out / f"{mode}-{i + 1:02}.png"
                pix.save(p)
                pages.append(
                    dict(
                        path=str(p.relative_to(root)),
                        sha256=sha(p),
                        text_characters=len(page.get_text()),
                    )
                )
            records.append(dict(mode=mode, workbook_sha256=sha(path), pages=pages))
            for start in range(0, len(pages), 4):
                images = [
                    Image.open(root / r["path"]).convert("RGB") for r in pages[start : start + 4]
                ]
                thumb = []
                for im in images:
                    im.thumbnail((1000, 710))
                    thumb.append(im.copy())
                contact = Image.new("RGB", (2000, 1420), "#cbd1d4")
                for j, im in enumerate(thumb):
                    contact.paste(im, ((j % 2) * 1000, (j // 2) * 710))
                contact.save(out / f"{mode}-contact-{start // 4 + 1:02}.png")
    dump(out / "REVIEW.json", dict(status="RENDERED_PENDING_MANUAL_REVIEW", records=records))
    print([(r["mode"], len(r["pages"])) for r in records])


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "command",
        choices=["build", "validate", "render", "verify-archive", "validate-qa", "verify-formulas"],
    )
    parser.add_argument("--archive", type=Path)
    args = parser.parse_args()
    if args.command == "verify-archive":
        if args.archive is None:
            parser.error("verify-archive requires --archive")
        verify_archive(args.archive)
        raise SystemExit(0)
    {
        "build": generate,
        "validate": validate,
        "render": render,
        "validate-qa": validate_qa,
        "verify-formulas": verify_formulas,
    }[args.command]()
