#!/usr/bin/env python3
"""Validate fixed reader routes against the evidence bytes they actually cite."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path
from urllib.parse import unquote, urlsplit

from markdown_it import MarkdownIt
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
REGISTER = Path("docs/reader/exercises/routes.json")
EXPECTED = {
    "SH-EX-ACQ-001": "docs/reader/exercises/ACQUISITION.md",
    "SH-EX-INV-001": "docs/reader/exercises/INVOICE.md",
    "SH-EX-CON-001": "docs/reader/exercises/CONTRACTS.md",
}


def local_target(root: Path, value: object) -> Path | None:
    """Accept repository-relative file paths only, including on a copied fixture."""
    if not isinstance(value, str) or not value or Path(value).is_absolute():
        return None
    target = (root / value).resolve()
    if not target.is_relative_to(root.resolve()):
        return None
    return target


def guide_links(root: Path, guide: Path) -> list[str]:
    errors = []
    for token in MarkdownIt().parse(guide.read_text()):
        for child in token.children or []:
            if child.type not in {"link_open", "image"}:
                continue
            href = child.attrGet("href") if child.type == "link_open" else child.attrGet("src")
            if not href:
                continue
            url = urlsplit(href)
            if url.scheme or url.netloc:
                continue
            target = (guide.parent / unquote(url.path)).resolve() if url.path else guide
            if not target.is_relative_to(root.resolve()) or not target.exists():
                errors.append(f"{guide.name}: broken local link {href}")
    return errors


def validate(root: Path = ROOT) -> list[str]:
    root = root.resolve()
    errors: list[str] = []
    try:
        register = json.loads((root / REGISTER).read_text())
    except (OSError, ValueError) as exc:
        return [f"route register unreadable: {exc}"]
    if not isinstance(register, dict):
        return ["route register must be an object"]
    if register.get("schema_version") != "1.0.0":
        errors.append("unsupported schema_version")
    if register.get("status") != "SOURCE_BOUND_LEARNING_GUIDES":
        errors.append("invalid route register status")
    if not re.fullmatch(r"[0-9a-f]{40}", str(register.get("source_revision", ""))):
        errors.append("invalid register source_revision")
    routes = register.get("routes")
    if not isinstance(routes, list) or len(routes) != 3:
        return errors + ["exactly three routes required"]
    if any(not isinstance(route, dict) for route in routes):
        return errors + ["route entries must be objects"]
    ids = [route.get("id") for route in routes]
    if sorted(str(value) for value in ids) != sorted(EXPECTED):
        errors.append("route IDs must match the three stable IDs without duplicates")
    for route in routes:
        label = str(route.get("id", "unknown"))
        for field in ("title", "period", "release_or_source", "source_revision", "guide"):
            if not isinstance(route.get(field), str) or not route[field].strip():
                errors.append(f"{label}: missing/invalid {field}")
        if not re.fullmatch(r"[0-9a-f]{40}", str(route.get("source_revision", ""))):
            errors.append(f"{label}: invalid source_revision")
        if "scenario" not in route or route["scenario"] not in (None, "base"):
            errors.append(f"{label}: invalid/missing scenario")
        if route.get("classification") != "PUBLIC_SYNTHETIC_LEARNING_GUIDE":
            errors.append(f"{label}: invalid classification")
        if route.get("new_design_required") is not False:
            errors.append(f"{label}: new_design_required must be false")
        native_ids = route.get("native_ids")
        if (
            not isinstance(native_ids, list)
            or not native_ids
            or any(not isinstance(value, str) or not value for value in native_ids)
        ):
            errors.append(f"{label}: native_ids must be nonempty strings")
        if route.get("guide") != EXPECTED.get(label):
            errors.append(f"{label}: stable guide path differs")
        guide = local_target(root, route.get("guide"))
        steps = route.get("steps")
        if (
            not isinstance(steps, list)
            or not steps
            or any(not isinstance(step, str) or not step for step in steps)
        ):
            errors.append(f"{label}: steps must be nonempty strings")
        if guide is None or not guide.is_file():
            errors.append(f"{label}: missing guide")
        else:
            numbered = re.findall(r"^(\d+)\. (.+)$", guide.read_text(), re.M)
            if [int(number) for number, _ in numbered] != list(range(1, len(numbered) + 1)):
                errors.append(f"{label}: guide step numbering is not sequential")
            if [step for _, step in numbered] != steps:
                errors.append(f"{label}: steps differ from numbered guide text")
            errors.extend(guide_links(root, guide))
        evidence = route.get("evidence")
        if not isinstance(evidence, list) or not evidence:
            errors.append(f"{label}: evidence must be a nonempty list")
            continue
        seen: set[str] = set()
        for item in evidence:
            if not isinstance(item, dict):
                errors.append(f"{label}: evidence must be objects")
                continue
            path = item.get("path")
            target = local_target(root, path)
            if target is None or not target.is_file():
                errors.append(f"{label}: missing evidence {path}")
                continue
            if path in seen:
                errors.append(f"{label}: duplicate evidence path {path}")
            seen.add(path)
            digest = hashlib.sha256(target.read_bytes()).hexdigest()
            if item.get("sha256") != digest:
                errors.append(f"{label}: evidence hash mismatch {path}")
            sheets = item.get("sheets")
            if target.suffix.lower() == ".xlsx" or sheets is not None:
                if not isinstance(sheets, list) or not sheets:
                    errors.append(f"{label}: workbook requires sheet names {path}")
                    continue
                try:
                    workbook = load_workbook(target, read_only=True)
                    actual = workbook.sheetnames
                    workbook.close()
                    if sheets != actual:
                        errors.append(f"{label}: workbook sheets differ {path}")
                except Exception as exc:
                    errors.append(f"{label}: unreadable workbook {path}: {exc}")
    index = root / "docs/reader/exercises/README.md"
    if not index.is_file():
        errors.append("missing exercise index")
    else:
        errors.extend(guide_links(root, index))
    return errors


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=ROOT)
    args = parser.parse_args()
    errors = validate(args.root)
    if errors:
        print("FAIL: reader exercise routes")
        for error in errors:
            print(f"- {error}")
        return 1
    print("PASS: 3 reader routes; source hashes, workbook sheets, steps and local links")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
