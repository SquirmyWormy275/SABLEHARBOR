"""Replay explicit reader routes and check that their promised evidence opens."""

import argparse
import hashlib
import json
from pathlib import Path
from urllib.parse import unquote, urlsplit

import fitz
import openpyxl
from markdown_it import MarkdownIt

ROOT = Path(__file__).resolve().parents[2]
HERE = Path("docs/reader/usability")


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def local_links(root, source):
    found = set()
    for token in MarkdownIt("commonmark").enable("table").parse((root / source).read_text()):
        for child in token.children or []:
            if child.type != "link_open":
                continue
            parts = urlsplit(child.attrGet("href") or "")
            if parts.scheme or parts.netloc or not parts.path:
                continue
            target = ((root / source).parent / unquote(parts.path)).resolve()
            if target.is_relative_to(root.resolve()):
                found.add(str(target.relative_to(root.resolve())))
    return found


def inspect(root=ROOT):
    spec = json.loads((root / HERE / "tasks.json").read_text())
    rows = []
    seen = set()
    for task in spec["tasks"]:
        assert task["id"] not in seen, "Duplicate route ID"
        seen.add(task["id"])
        route = task["route"]
        assert route[0] in ("README.md", "docs/wiki/Home.md", "docs/wiki/Start-Here.md")
        for source, target in zip(route, route[1:], strict=False):
            assert target in local_links(root, source), f"Missing reader link: {source} -> {target}"
            path = (root / target).resolve()
            assert path.is_relative_to(root.resolve()) and path.is_file(), target
        endpoint = root / route[-1]
        result = {}
        if "expected_sheets" in task:
            book = openpyxl.load_workbook(endpoint, read_only=True, data_only=False)
            assert book.sheetnames == task["expected_sheets"], "Promised workbook sheets differ"
            result["worksheets"] = book.sheetnames
            book.close()
        if "expected_pdf_text" in task:
            with fitz.open(endpoint) as pdf:
                text = "\n".join(page.get_text() for page in pdf)
                assert all(t in text for t in task["expected_pdf_text"]), (
                    "Promised PDF text missing"
                )
                result["pdf_pages"] = len(pdf)
        if "expected_text" in task:
            assert all(t in endpoint.read_text() for t in task["expected_text"]), (
                "Promised reading content missing"
            )
        rows.append(
            {
                **task,
                "result": "PASS",
                "links_followed": len(route) - 1,
                "input_hashes": {p: sha(root / p) for p in route},
                **result,
            }
        )
    return {
        "method": spec["status"],
        "accepted_main": spec["accepted_main"],
        "task_source_sha256": sha(root / HERE / "tasks.json"),
        "tasks": rows,
        "corrections": spec["corrections"],
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check", action="store_true")
    args = parser.parse_args()
    result = inspect()
    path = ROOT / HERE / "RESULTS.json"
    if args.check:
        assert json.loads(path.read_text()) == result, (
            "Reader results stale; replay and review changed routes"
        )
    else:
        path.write_text(json.dumps(result, indent=2) + "\n")
    print(f"PASS: {len(result['tasks'])} complete reader routes and promised PDF/workbook contents")


if __name__ == "__main__":
    main()
